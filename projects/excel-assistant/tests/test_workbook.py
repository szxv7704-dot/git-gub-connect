import tempfile
import unittest
from pathlib import Path
import csv

from openpyxl import Workbook, load_workbook

from excel_assistant.core import TaskPlan
from excel_assistant.workbook import execute_plan, inspect_workbook


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.source = Path(self.temp.name) / "sample.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "원본데이터"
        ws.append(["이름", "부서", "상태", "금액"])
        ws.append(["김민지", "총무", "완료", 1000])
        ws.append(["박서준", "연구", "진행", 2000])
        ws.append(["김민지", "총무", "완료", 3000])
        wb.save(self.source)

    def tearDown(self):
        self.temp.cleanup()

    def test_inspect(self):
        info = inspect_workbook(self.source)
        self.assertEqual(info["원본데이터"]["headers"], ["이름", "부서", "상태", "금액"])

    def test_group_sum_output(self):
        output = Path(self.temp.name) / "result.xlsx"
        plan = TaskPlan(operation="group_sum", group_column="이름", value_column="금액", descending=True)
        headers, rows, stats = execute_plan(self.source, "원본데이터", 1, plan, output)
        self.assertEqual(headers, ["이름", "건수", "금액_합계"])
        self.assertEqual(rows[0], ["김민지", 2, 4000.0])
        self.assertEqual(stats["result_rows"], 2)
        wb = load_workbook(output)
        self.assertIn("작업결과", wb.sheetnames)
        self.assertIn("작업기록", wb.sheetnames)

    def test_group_statistics_output(self):
        output = Path(self.temp.name) / "stats.xlsx"
        plan = TaskPlan(
            operation="group_sum", group_column="이름", value_column="금액",
            descending=True, include_statistics=True,
        )
        _, _, stats = execute_plan(self.source, "원본데이터", 1, plan, output)
        self.assertEqual(stats["statistics"]["maximum_group"], "김민지")
        self.assertEqual(stats["statistics"]["maximum"], 4000.0)
        self.assertEqual(stats["statistics"]["minimum"], 2000.0)
        self.assertEqual(stats["statistics"]["average"], 3000.0)
        wb = load_workbook(output)
        self.assertIn("통계요약", wb.sheetnames)

    def test_numeric_filter_and_mixed_sort(self):
        output = Path(self.temp.name) / "filtered.xlsx"
        plan = TaskPlan(operation="filter", filter_column="금액", filter_value=2000.0, filter_operator="gte")
        _, rows, _ = execute_plan(self.source, "원본데이터", 1, plan, output)
        self.assertEqual(len(rows), 2)
        plan = TaskPlan(operation="sort", sort_column="금액", descending=True)
        _, rows, _ = execute_plan(self.source, "원본데이터", 1, plan, output)
        self.assertEqual([row[3] for row in rows], [3000, 2000, 1000])

    def test_title_row_footer_and_currency_text(self):
        source = Path(self.temp.name) / "irregular.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["2026년 지급 현황", None, None])
        ws.append([None, None, None])
        ws.append(["담당자", "상태", "지급액"])
        ws.append(["김", "완료", "1,200원"])
        ws.append(["박", "진행", "800원"])
        ws.append([None, None, None])
        ws.append(["※ 테스트용 안내 문구입니다.", None, None])
        wb.save(source)
        info = inspect_workbook(source)["Sheet"]
        self.assertEqual(info["header_row"], 3)
        self.assertEqual(info["profiles"]["지급액"]["kind"], "number")
        output = Path(self.temp.name) / "irregular_result.xlsx"
        plan = TaskPlan(operation="group_sum", group_column="담당자", value_column="지급액")
        _, rows, stats = execute_plan(source, "Sheet", 3, plan, output)
        self.assertEqual(len(rows), 2)
        self.assertEqual(stats["source_rows"], 2)

    def test_csv_cp949(self):
        source = Path(self.temp.name) / "sample.csv"
        source.write_text("부서,금액,상태\n총무,1000,완료\n연구,2000,진행\n", encoding="cp949")
        info = inspect_workbook(source)["CSV 데이터"]
        self.assertEqual(info["profiles"]["금액"]["kind"], "number")
        output = Path(self.temp.name) / "csv_result.xlsx"
        plan = TaskPlan(operation="sort", sort_column="금액", descending=True)
        _, rows, _ = execute_plan(source, "CSV 데이터", 1, plan, output)
        self.assertEqual(rows[0][0], "연구")

    def test_multi_group_and_sum_check(self):
        source = Path(self.temp.name) / "hierarchy.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["부서", "세부사업", "중앙", "광역", "기초", "기타", "합계"])
        ws.append(["복지", "청년지원", 10, 20, 30, 40, 100])
        ws.append(["복지", "노인지원", 5, 5, 5, 5, 21])
        ws.append(["문화", "축제", 1, 2, 3, 4, 10])
        wb.save(source)
        output = Path(self.temp.name) / "check.xlsx"
        plan = TaskPlan(operation="sum_check", component_columns=("중앙", "광역", "기초", "기타"), total_column="합계")
        headers, rows, stats = execute_plan(source, "Sheet", 1, plan, output)
        self.assertEqual(stats["statistics"]["mismatch_count"], 1)
        self.assertEqual(rows[1][-1], "불일치")
        plan = TaskPlan(operation="group_sum", group_column="부서", group_columns=("부서", "세부사업"), value_column="합계")
        headers, rows, _ = execute_plan(source, "Sheet", 1, plan, output)
        self.assertEqual(headers[:2], ["부서", "세부사업"])
        self.assertEqual(len(rows), 3)

    def test_multiple_tables_are_exposed_as_separate_choices(self):
        source = Path(self.temp.name) / "two_tables.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["부서", "금액"])
        ws.append(["총무", 100])
        ws.append([])
        ws.append(["이름", "점수"])
        ws.append(["김", 90])
        wb.save(source)
        info = inspect_workbook(source)
        self.assertEqual(len(info), 2)
        self.assertEqual(info["Sheet · 표 1"]["headers"], ["부서", "금액"])
        self.assertEqual(info["Sheet · 표 2"]["headers"], ["이름", "점수"])

    def test_uncached_formula_is_reported(self):
        source = Path(self.temp.name) / "formula.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["항목", "금액", "합계"])
        ws.append(["A", 10, "=B2*2"])
        wb.save(source)
        info = inspect_workbook(source)["Sheet"]
        self.assertEqual(info["formula_count"], 1)
        self.assertEqual(info["uncached_formula_count"], 1)


if __name__ == "__main__":
    unittest.main()
