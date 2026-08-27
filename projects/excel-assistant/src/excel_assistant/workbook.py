from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime
import csv
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .core import RequestError, TaskPlan, coerce_number, validate_task_plan


HEADER_FILL = PatternFill("solid", fgColor="17643F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
DUPLICATE_FILL = PatternFill("solid", fgColor="FFE2E0")


def inspect_workbook(path: str | Path) -> dict[str, dict[str, Any]]:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        headers, rows = _read_csv(path)
        profiles = _profiles_from_rows(headers, rows)
        return {"CSV 데이터": {"header_row": 1, "headers": headers, "rows": len(rows), "columns": len(headers), "profiles": profiles}}
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise RequestError("현재는 .xlsx, .xlsm, .csv 파일을 지원합니다. 구형 .xls 파일은 Excel에서 .xlsx로 저장한 뒤 열어 주세요.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    formula_workbook = load_workbook(path, read_only=True, data_only=False)
    result: dict[str, dict[str, Any]] = {}
    try:
        for sheet in workbook.worksheets:
            _ensure_dimensions(sheet)
            tables = detect_table_regions(sheet)
            formula_sheet = formula_workbook[sheet.title]
            formula_count, uncached_count = _formula_status(formula_sheet, sheet)
            for index, (header_row, end_row) in enumerate(tables, start=1):
                headers = read_headers(sheet, header_row)
                profiles = profile_columns(sheet, header_row, end_row=end_row)
                key = sheet.title if len(tables) == 1 else f"{sheet.title} · 표 {index}"
                result[key] = {
                    "source_sheet": sheet.title,
                    "header_row": header_row,
                    "end_row": end_row,
                    "headers": headers,
                    "rows": max(end_row - header_row, 0),
                    "columns": len(headers),
                    "profiles": profiles,
                    "table_count": len(tables),
                    "formula_count": formula_count,
                    "uncached_formula_count": uncached_count,
                }
    finally:
        workbook.close()
        formula_workbook.close()
    return result


def _ensure_dimensions(sheet) -> None:
    """Populate dimensions for read-only sheets whose XML omits a dimension ref."""
    if sheet.max_row is None or sheet.max_column is None:
        sheet.calculate_dimension(force=True)


def detect_header_row(sheet, scan_rows: int = 20) -> int:
    _ensure_dimensions(sheet)
    best_row, best_score = 1, -1
    scanned = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, scan_rows), values_only=True))
    for index, values in enumerate(scanned):
        row = index + 1
        nonempty = [value for value in values if value not in (None, "")]
        unique_text = len({str(value).strip() for value in nonempty})
        text_count = sum(isinstance(value, str) for value in nonempty)
        long_text = sum(len(str(value)) > 40 for value in nonempty)
        next_nonempty = 0
        if index + 1 < len(scanned):
            next_nonempty = sum(value not in (None, "") for value in scanned[index + 1])
        continuity = min(next_nonempty, len(nonempty)) * 0.2
        score = len(nonempty) + unique_text * 0.25 + text_count * 0.25 + continuity - long_text * 1.5
        if len(nonempty) >= 2 and score > best_score:
            best_row, best_score = row, score
    return best_row


def _header_score_values(values, next_values=()) -> float:
    nonempty = [value for value in values if value not in (None, "")]
    if len(nonempty) < 2:
        return -1
    unique_text = len({str(value).strip() for value in nonempty})
    text_count = sum(isinstance(value, str) for value in nonempty)
    next_nonempty = sum(value not in (None, "") for value in next_values)
    return len(nonempty) + unique_text * 0.25 + text_count * 0.25 + min(next_nonempty, len(nonempty)) * 0.2


def detect_table_regions(sheet, scan_limit: int = 200) -> list[tuple[int, int]]:
    """Find distinct table headers separated by blank rows or repeated header blocks."""
    _ensure_dimensions(sheet)
    candidates: list[int] = []
    blank_before = True
    scanned = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, scan_limit), values_only=True))
    for index, values in enumerate(scanned):
        row = index + 1
        populated = sum(value not in (None, "") for value in values)
        next_values = scanned[index + 1] if index + 1 < len(scanned) else ()
        score = _header_score_values(values, next_values)
        if score >= 3.0 and (blank_before or not candidates):
            candidates.append(row)
        blank_before = populated == 0
    primary = detect_header_row(sheet)
    if primary not in candidates:
        candidates.append(primary)
    candidates = sorted(row for index, row in enumerate(sorted(set(candidates))) if index == 0 or row - sorted(set(candidates))[index - 1] > 2)
    regions = []
    for index, header in enumerate(candidates):
        end = candidates[index + 1] - 1 if index + 1 < len(candidates) else sheet.max_row
        regions.append((header, end))
    return regions or [(primary, sheet.max_row)]


def _formula_status(formula_sheet, value_sheet, sample_limit: int = 50000) -> tuple[int, int]:
    formula_count = uncached = checked = 0
    for row in formula_sheet.iter_rows():
        for cell in row:
            if checked >= sample_limit:
                return formula_count, uncached
            checked += 1
            if cell.data_type == "f":
                formula_count += 1
                if value_sheet[cell.coordinate].value is None:
                    uncached += 1
    return formula_count, uncached


def read_headers(sheet, header_row: int) -> list[str]:
    headers: list[str] = []
    seen: Counter[str] = Counter()
    values = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
    for raw in values:
        if raw in (None, ""):
            continue
        name = str(raw).strip()
        seen[name] += 1
        headers.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return headers


def profile_columns(sheet, header_row: int, sample_limit: int = 1000, end_row: int | None = None) -> dict[str, dict[str, Any]]:
    """Infer basic column types without treating identifiers stored as text as numbers."""
    _ensure_dimensions(sheet)
    names: list[str] = []
    positions: list[int] = []
    seen: Counter[str] = Counter()
    header_values = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
    for col, raw in enumerate(header_values, start=1):
        if raw in (None, ""):
            continue
        name = str(raw).strip()
        seen[name] += 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
        positions.append(col)
    stats = {name: {"nonempty": 0, "numbers": 0, "dates": 0, "unique": set()} for name in names}
    last_row = min(end_row or sheet.max_row, header_row + sample_limit)
    for values in sheet.iter_rows(min_row=header_row + 1, max_row=last_row, values_only=True):
        for name, col in zip(names, positions):
            value = values[col - 1] if col <= len(values) else None
            if value in (None, ""):
                continue
            item = stats[name]
            item["nonempty"] += 1
            if isinstance(value, (datetime, date)):
                item["dates"] += 1
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                item["numbers"] += 1
            elif _numeric_text(value, name):
                item["numbers"] += 1
            if len(item["unique"]) < 200:
                item["unique"].add(str(value))
    result = {}
    for name, item in stats.items():
        nonempty = item["nonempty"]
        if nonempty and item["dates"] / nonempty >= 0.8:
            kind = "date"
        elif nonempty and item["numbers"] / nonempty >= 0.8:
            kind = "number"
        else:
            kind = "text"
        result[name] = {"kind": kind, "nonempty": nonempty, "unique": len(item["unique"])}
    return result


def _numeric_text(value: Any, header: str) -> bool:
    text = str(value).strip()
    identifier_words = ("번호", "코드", "전화", "휴대폰", "계좌", "주민", "우편")
    if any(word in header for word in identifier_words):
        return False
    if not re.fullmatch(r"\(?[-+]?\d[\d,]*(?:\.\d+)?\)?(?:원|%|명|건|개)?", text):
        return False
    return coerce_number(text) is not None


def _profiles_from_rows(headers: list[str], rows: list[list[Any]]) -> dict[str, dict[str, Any]]:
    result = {}
    for index, header in enumerate(headers):
        values = [row[index] for row in rows if index < len(row) and row[index] not in (None, "")]
        dates = sum(isinstance(value, (date, datetime)) for value in values)
        numbers = sum((isinstance(value, (int, float)) and not isinstance(value, bool)) or _numeric_text(value, header) for value in values)
        kind = "date" if values and dates / len(values) >= 0.8 else "number" if values and numbers / len(values) >= 0.8 else "text"
        result[header] = {"kind": kind, "nonempty": len(values), "unique": len({str(value) for value in values[:200]})}
    return result


def _read_csv(path: Path) -> tuple[list[str], list[list[Any]]]:
    raw = path.read_bytes()
    encoding = "utf-8-sig"
    for candidate in ("utf-8-sig", "cp949", "utf-8"):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    else:
        raise RequestError("CSV 문자 인코딩을 확인하지 못했습니다.")
    reader = csv.reader(text.splitlines())
    all_rows = list(reader)
    if not all_rows:
        raise RequestError("CSV 파일이 비어 있습니다.")
    headers = _unique_headers(all_rows[0])
    positions = [index for index, value in enumerate(all_rows[0]) if str(value).strip()]
    rows = [[row[index] if index < len(row) else None for index in positions] for row in all_rows[1:] if any(str(value).strip() for value in row)]
    return headers, rows


def _unique_headers(values: list[Any]) -> list[str]:
    headers, seen = [], Counter()
    for value in values:
        if value in (None, ""):
            continue
        name = str(value).strip()
        seen[name] += 1
        headers.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return headers


def _table_data(path: str | Path, sheet_name: str, header_row: int, end_row: int | None = None) -> tuple[list[str], list[list[Any]]]:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        return _read_csv(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        _ensure_dimensions(sheet)
        raw_headers = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        headers: list[str] = []
        positions: list[int] = []
        seen: Counter[str] = Counter()
        for col, raw in enumerate(raw_headers, start=1):
            if raw in (None, ""):
                continue
            name = str(raw).strip()
            seen[name] += 1
            headers.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
            positions.append(col)
        rows = []
        blank_streak = 0
        for raw_row in sheet.iter_rows(min_row=header_row + 1, max_row=min(end_row or sheet.max_row, sheet.max_row), values_only=True):
            values = [raw_row[col - 1] if col <= len(raw_row) else None for col in positions]
            populated = sum(value not in (None, "") for value in values)
            if not populated:
                blank_streak += 1
                continue
            blank_streak = 0
            if len(values) >= 3 and populated == 1 and isinstance(next(value for value in values if value not in (None, "")), str):
                lone = str(next(value for value in values if value not in (None, "")))
                if len(lone) > 25 or any(word in lone for word in ("안내", "주의", "테스트용", "합계")):
                    continue
            rows.append(values)
        return headers, rows
    finally:
        workbook.close()


def _style_result_sheet(sheet, table_name: str) -> None:
    for cell in sheet[1]:
        cell.fill = copy(HEADER_FILL)
        cell.font = copy(HEADER_FONT)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        width = min(max((len(str(cell.value)) if cell.value is not None else 0) for cell in column) + 3, 38)
        sheet.column_dimensions[get_column_letter(index)].width = max(width, 10)
    if sheet.max_row >= 2 and sheet.max_column >= 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def execute_plan(
    source_path: str | Path,
    sheet_name: str,
    header_row: int,
    plan: TaskPlan,
    output_path: str | Path,
    end_row: int | None = None,
) -> tuple[list[str], list[list[Any]], dict[str, Any]]:
    headers, rows = _table_data(source_path, sheet_name, header_row, end_row)
    validate_task_plan(plan, headers)
    header_index = {header: index for index, header in enumerate(headers)}
    output_headers: list[str]
    output_rows: list[list[Any]]
    excluded = 0
    statistics: dict[str, Any] = {}

    if plan.operation == "group_sum":
        group_columns = list(plan.group_columns or ((plan.group_column,) if plan.group_column else ()))
        group_indices, value_i = [header_index[column] for column in group_columns], header_index[plan.value_column]
        totals: dict[tuple[str, ...], float] = defaultdict(float)
        counts: Counter[tuple[str, ...]] = Counter()
        for row in rows:
            group_values = [row[index] for index in group_indices]
            number = coerce_number(row[value_i])
            if any(value in (None, "") for value in group_values) or (number is None and plan.aggregation != "count"):
                excluded += 1
                continue
            key = tuple(str(value).strip() for value in group_values)
            totals[key] += 1 if plan.aggregation == "count" else number
            counts[key] += 1
        output_rows = []
        for key in totals:
            result = totals[key] / counts[key] if plan.aggregation == "average" and counts[key] else totals[key]
            output_rows.append([*key, counts[key], result])
        result_index = len(group_columns) + 1
        output_rows.sort(key=lambda item: item[result_index] if plan.descending else tuple(item[:len(group_columns)]), reverse=plan.descending)
        suffix = {"sum": "합계", "average": "평균", "count": "개수"}.get(plan.aggregation, "결과")
        output_headers = [*group_columns, "건수", f"{plan.value_column}_{suffix}"]
        if plan.include_statistics and output_rows:
            maximum = max(output_rows, key=lambda item: item[result_index])
            minimum = min(output_rows, key=lambda item: item[result_index])
            average = sum(item[result_index] for item in output_rows) / len(output_rows)
            statistics = {
                "maximum_group": " / ".join(str(value) for value in maximum[:len(group_columns)]), "maximum": maximum[result_index],
                "minimum_group": " / ".join(str(value) for value in minimum[:len(group_columns)]), "minimum": minimum[result_index],
                "average": average,
            }
    elif plan.operation == "filter":
        if plan.filter_column not in header_index:
            raise RequestError("필터 기준 열을 찾을 수 없습니다.")
        col_i = header_index[plan.filter_column]
        target = plan.filter_value
        def matches(value):
            if plan.filter_operator in {"gt", "gte", "lt", "lte"}:
                number = coerce_number(value)
                if number is None:
                    return False
                return {"gt": number > target, "gte": number >= target, "lt": number < target, "lte": number <= target}[plan.filter_operator]
            actual, expected = str(value or "").strip().casefold(), str(target).strip().casefold()
            return {"eq": actual == expected, "ne": actual != expected, "contains": expected in actual}.get(plan.filter_operator, actual == expected)
        output_rows = [row for row in rows if matches(row[col_i])]
        excluded = len(rows) - len(output_rows)
        output_headers = headers
    elif plan.operation == "sort":
        if plan.sort_column not in header_index:
            raise RequestError("정렬 기준 열을 찾을 수 없습니다.")
        col_i = header_index[plan.sort_column]
        def key(row):
            value = row[col_i]
            number = coerce_number(value)
            if value in (None, ""):
                return (2, "")
            if number is not None:
                return (0, number)
            return (1, str(value).casefold())
        nonblank = [row for row in rows if row[col_i] not in (None, "")]
        blanks = [row for row in rows if row[col_i] in (None, "")]
        output_rows = sorted(nonblank, key=key, reverse=plan.descending) + blanks
        output_headers = headers
    elif plan.operation == "duplicates":
        if plan.group_column not in header_index:
            raise RequestError("중복 검사 열을 찾을 수 없습니다.")
        col_i = header_index[plan.group_column]
        counts = Counter(str(row[col_i]).strip() for row in rows if row[col_i] not in (None, ""))
        output_rows = [row for row in rows if row[col_i] not in (None, "") and counts[str(row[col_i]).strip()] > 1]
        excluded = len(rows) - len(output_rows)
        output_headers = headers
    elif plan.operation == "sum_check":
        if plan.total_column not in header_index or any(column not in header_index for column in plan.component_columns):
            raise RequestError("합계 검증에 필요한 열을 찾을 수 없습니다.")
        total_i = header_index[plan.total_column]
        component_indices = [header_index[column] for column in plan.component_columns]
        output_headers = [*headers, "계산합계", "차액", "검증결과"]
        output_rows = []
        mismatch_count = 0
        for row in rows:
            values = [coerce_number(row[index]) for index in component_indices]
            declared = coerce_number(row[total_i])
            if declared is None or any(value is None for value in values):
                calculated, difference, result = None, None, "확인 필요"
            else:
                calculated = sum(values)
                difference = declared - calculated
                result = "일치" if abs(difference) < 0.000001 else "불일치"
                mismatch_count += result == "불일치"
            output_rows.append([*row, calculated, difference, result])
        statistics = {"mismatch_count": mismatch_count}
    else:
        raise RequestError(f"지원하지 않는 작업입니다: {plan.operation}")

    workbook = Workbook()
    result_sheet = workbook.active
    result_sheet.title = "작업결과"
    result_sheet.append(output_headers)
    for row in output_rows:
        result_sheet.append(row)
    _style_result_sheet(result_sheet, "ResultTable")
    for cell in result_sheet[1]:
        if any(word in str(cell.value) for word in ("합계", "평균", "금액", "비용", "가격")):
            for row in range(2, result_sheet.max_row + 1):
                result_sheet.cell(row, cell.column).number_format = '#,##0.00'
    if plan.operation == "duplicates" and plan.group_column in output_headers:
        col = output_headers.index(plan.group_column) + 1
        for row in range(2, result_sheet.max_row + 1):
            result_sheet.cell(row, col).fill = copy(DUPLICATE_FILL)
    if plan.operation == "sum_check":
        result_col = output_headers.index("검증결과") + 1
        difference_col = output_headers.index("차액") + 1
        for row in range(2, result_sheet.max_row + 1):
            if result_sheet.cell(row, result_col).value == "불일치":
                result_sheet.cell(row, result_col).fill = copy(DUPLICATE_FILL)
                result_sheet.cell(row, difference_col).fill = copy(DUPLICATE_FILL)

    if statistics and plan.include_statistics:
        summary = workbook.create_sheet("통계요약")
        summary.append(["통계", "대상", "값"])
        summary.append(["최댓값", statistics["maximum_group"], statistics["maximum"]])
        summary.append(["최솟값", statistics["minimum_group"], statistics["minimum"]])
        summary.append(["평균", "그룹별 합계 평균", statistics["average"]])
        _style_result_sheet(summary, "StatisticsTable")

    log = workbook.create_sheet("작업기록")
    log.append(["항목", "내용"])
    log.append(["원본 파일", Path(source_path).name])
    log.append(["원본 시트", sheet_name])
    log.append(["작업 유형", plan.operation_label])
    log.append(["작업 설명", plan.explanation])
    log.append(["열 매칭 신뢰도", f"{plan.confidence:.0%}"])
    log.append(["원본 데이터 행", len(rows)])
    log.append(["결과 데이터 행", len(output_rows)])
    log.append(["제외 데이터 행", excluded])
    _style_result_sheet(log, "WorkLogTable")
    workbook.save(output_path)
    if not Path(output_path).exists() or Path(output_path).stat().st_size == 0:
        raise RequestError("결과 파일이 정상적으로 생성되지 않았습니다.")
    return output_headers, output_rows, {
        "source_rows": len(rows), "result_rows": len(output_rows),
        "excluded_rows": excluded, "statistics": statistics,
    }
