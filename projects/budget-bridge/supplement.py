"""추경(추가경정예산) 사업별 설명서와 UBIS 추경 검토조서를 읽고 점검한다.

본예산 설명서는 사업 **전체**의 산출근거를 적는다. 추경 설명서는 다르다.
2025·2026년 1회·2회 추경 설명서 4부와 추경 검토조서 4부를 읽어 확인한 모양이다.

  사업 머리 표       2026년 제2회추경예산안(A) · 2026년 기정예산(B) · 2025년 최종예산액 · 증감(A-B)
  사업계획 절        1. 부서기본운영경비 / - 사업비 : △13,597천원(자체비)
  증감 산출내역 표   사업명 | 추가경정예산안(A) | 기정예산액(B)(전년도예산액) | 비교증감내역: 산출내역·요구액 | 비고
                     합  계 | 233,731 | 247,328(180,642) |  | △13,597
                     부서기본운영경비 | 44,120 | 46,710(35,487) | ① 일반운영비 | △2,590
                     ․ 개인당(△740천원×7명×6/12개월) | △2,590천원
                     44,371 | 47,048(35,415) | ② 특근매식비(△9천원×7명×85회×6/12개월) | △2,677천원 | △2,677
  미반영 사업 표     사업명 | 추가경정예산안(A) | 기정예산액(B)(전년도예산액) | 증감(A-B) | 비고
  총괄 표            순 | 사업항목명 | 추가경정예산액 | 기정예산액 | 비교증감 | 쪽수

즉 산출식은 **증감분만** 있다. 그래서 본예산처럼 '산출근거 합계 = 요구액'을 볼 수 없고,
대신 아래가 성립해야 한다. 네 부 모두에서 성립하는 것을 확인한 규칙만 검사한다.

  1. 모든 줄에서 추경안(A) − 기정(B) = 증감
  2. 산출식 계산값 = 적힌 금액 (△ 부호, 6/12개월 같은 분수 포함)
  3. 한 줄의 금액 = 바로 아래 줄들의 합
  4. 표의 합계 줄 = 줄들의 합 (A · B · 증감 각각)
  5. 사업계획 절의 '사업비' = 그 절 표의 증감 합계
  6. 사업의 모든 표(증감 + 미반영)를 더하면 사업 머리 표의 A · B · 증감이 된다
  7. 총괄 표의 사업별 금액 = 각 사업 쪽의 금액
  8. 정책/단위/세부사업별 현황 표에서 위 단계 = 아래 단계의 합

UBIS 추경 검토조서의 열은 본예산과 뜻이 다르다. '추경요구액(B)'는 사업 전체가 아니라
**증감**이고, '기정액(A)'이 설명서의 기정예산(B)이다. 이름이 같은 열을 본예산처럼
읽으면 모든 사업이 틀린 것으로 나온다.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from crosscheck import SKIPPED, Report
from hwpx_native_parser import is_private_use, native_chunks
from plan_parser import (ARROW_BULLETS, DOT_BULLETS, HEADING, _cells, _is_organisation,
                         _looks_like_title, _organisation, _squeeze_name)

TOLERANCE = 0.5          # 천원. 1천원 오타도 잡는다.
TRIANGLE = "△▲"
SIGNED = re.compile(r"^([△▲-])?\s*([\d,]+(?:\.\d+)?)$")
WITH_PREVIOUS = re.compile(r"^([△▲-]?\s*[\d,]+)\s*\(\s*([△▲-]?\s*[\d,]+)\s*\)$")
MARKED = re.compile(r"^([△▲-])?\s*([\d,]+)\s*천원$")
CIRCLED = re.compile(r"^([①-⑳](?:\s*,\s*[①-⑳])*)\s*(.*)$")
FORMULA_TAIL = re.compile(r"^(.*?)\s*\(([^()]*[×xX][^()]*)\)\s*$")
FACTOR = re.compile(r"([△▲-])?([\d,]+(?:\.\d+)?)(?:\s*/\s*([\d,]+(?:\.\d+)?))?\s*(천원|원|%)?")
PLAN_COST = re.compile(r"사업비\s*:\s*([△▲-]?\s*[\d,]+)\s*(천원|원)")
PERIOD = re.compile(r"사업기간\s*:\s*(20\d{2})\s*\.[^~]*~\s*(20\d{2})")
TOTAL_COST = re.compile(r"총사업비\s*:\s*([△▲-]?\s*[\d,]+)\s*천원")
SUMMARY_WORDS = ("합계", "계", "소계", "총계")


# ------------------------------------------------------------------ 숫자

def number(text) -> Optional[float]:
    """'△1,000' → -1000 · '2,376,941' → 2376941 · 아니면 None."""
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text)
    found = SIGNED.match(str(text).strip())
    if not found:
        return None
    value = float(found.group(2).replace(",", ""))
    return -value if found.group(1) else value


def marked(text) -> Optional[float]:
    """'△2,590천원' → -2590. '천원'이 붙은 칸만."""
    found = MARKED.match(str(text or "").strip())
    if not found:
        return None
    value = float(found.group(2).replace(",", ""))
    return -value if found.group(1) else value


def with_previous(text) -> Optional[tuple[float, Optional[float]]]:
    """'2,376,941(0)' → (2376941, 0). 괄호 안은 전년도 예산액이다."""
    found = WITH_PREVIOUS.match(str(text or "").strip())
    if found:
        return number(found.group(1)), number(found.group(2))
    return None


def _product(term: str) -> Optional[float]:
    parts = [part for part in re.split(r"[×xX*]", term) if part.strip()]
    if not parts:
        return None
    product = 1.0
    sign = 1.0
    for index, part in enumerate(parts):
        found = FACTOR.search(part.replace(" ", ""))
        if not found:
            return None
        value = float(found.group(2).replace(",", ""))
        if found.group(3):
            denominator = float(found.group(3).replace(",", ""))
            if not denominator:
                return None
            value /= denominator
        if found.group(1):
            sign = -sign
        unit = found.group(4)
        if unit == "%":
            value /= 100
        elif index == 0:
            if unit == "천원":
                value *= 1000
            elif unit != "원":
                return None
        product *= value
    return sign * product / 1000


def calc(formula: str) -> Optional[float]:
    """증감 산출식을 천원 단위로 계산한다.

    추경 산출식은 본예산에 없는 것이 셋 있다.
      - 감액은 첫 항에 △ 가 붙는다    `△740천원×7명×6/12개월`
      - 개월 수를 분수로 적는다        `6/12개월`
      - 끝전을 빼거나 더한다          `14천원×3매×1시간×2명×3회-2천원`
    본예산용 계산기(crosscheck.calc)는 셋 다 모른다. 6/12 를 6 으로 읽으면 금액이
    12배로 튀어 멀쩡한 줄이 오류가 된다.
    """
    if not formula:
        return None
    body = formula.split("=")[0].strip()
    terms = re.split(r"\s*(?<=[^\s×xX*/(])\s*([+\-])\s*(?=[\d△▲(])", body)
    total = _product(terms[0])
    if total is None:
        return None
    for index in range(1, len(terms) - 1, 2):
        value = _product(terms[index + 1])
        if value is None:
            return None
        total = total + value if terms[index] == "+" else total - value
    return total


def _formula_ok(value: Optional[float], written: Optional[float]) -> bool:
    """산출식 계산값과 적힌 금액이 맞는가.

    천원 아래 끝전은 버리기도(△2,677.5 → △2,677) 올리기도(940,623.25 → 940,624) 한다.
    계산값에 끝전이 있을 때만 1천원 안쪽을 같다고 본다. 끝전이 없으면 1천원도 봐주지 않는다.
    """
    if value is None or written is None:
        return True
    if abs(value - round(value)) < 1e-6:
        return abs(value - written) < TOLERANCE
    return abs(value - written) < 1


def _same(one: Optional[float], other: Optional[float]) -> bool:
    return one is not None and other is not None and abs(one - other) < TOLERANCE


def _fmt(value: Optional[float]) -> str:
    """금액 표기. 설명서처럼 감액은 △ 로 적는다."""
    if value is None:
        return "없음"
    rounded = round(value)
    return f"△{abs(rounded):,}" if rounded < 0 else f"{rounded:,}"


# ------------------------------------------------------------------ 구조

@dataclass
class SuppLine:
    """증감 산출내역의 한 줄. ① 아래의 ․ · 사설영역 글머리 줄."""

    depth: int                     # ① 바로 아래가 1
    name: str
    formula: str = ""
    amount: Optional[float] = None
    row: int = 0
    bullet: str = ""


@dataclass
class SuppGroup:
    """표의 한 줄 묶음. ① 항목 하나, 또는 미반영 사업 한 줄."""

    title: str
    after: Optional[float] = None        # 추경안(A)
    before: Optional[float] = None       # 기정(B)
    previous: Optional[float] = None     # 괄호 안 전년도 예산액
    change: Optional[float] = None       # 요구액 = 증감(A-B)
    own: Optional[float] = None          # 산출내역 칸에 '천원'을 붙여 적은 금액
    formula: str = ""
    business: str = ""                   # 표 왼쪽 '사업명' 칸
    note: str = ""
    row: int = 0
    lines: list[SuppLine] = field(default_factory=list)


@dataclass
class SuppTable:
    kind: str                            # "증감" / "미반영"
    section: str = ""                    # 앞선 사업계획 절 제목
    section_cost: Optional[float] = None # 그 절의 '- 사업비 : …천원'
    total: Optional[SuppGroup] = None    # 합계 줄
    groups: list[SuppGroup] = field(default_factory=list)
    order: int = 0

    def sum_of(self, key: str) -> Optional[float]:
        values = [getattr(group, key) for group in self.groups]
        if any(value is None for value in values):
            return None
        return sum(values)

    def value(self, key: str) -> Optional[float]:
        """합계 줄이 있으면 그 값, 없으면 줄들의 합."""
        if self.total is not None and getattr(self.total, key) is not None:
            return getattr(self.total, key)
        return self.sum_of(key)


@dataclass
class SuppProject:
    name: str
    number: str = ""
    heading: str = ""
    policy: str = ""
    unit: str = ""
    program: str = ""
    organisation: str = ""
    after: Optional[float] = None
    before: Optional[float] = None
    last_final: Optional[float] = None
    change: Optional[float] = None
    header_labels: list = field(default_factory=list)
    total_cost: Optional[float] = None
    funding: dict = field(default_factory=dict)
    period: tuple = ()                     # 사업기간 (시작 연도, 끝 연도)
    tables: list[SuppTable] = field(default_factory=list)

    @property
    def single_year(self) -> bool:
        return len(self.period) == 2 and self.period[0] == self.period[1]

    @property
    def has_unreflected(self) -> bool:
        return any(table.kind == "미반영" for table in self.tables)

    @property
    def label(self) -> str:
        return f"{self.number}. {self.name}" if self.number else self.name


@dataclass
class SummaryRow:
    """총괄 표 한 줄."""

    order: str
    name: str
    after: Optional[float]
    before: Optional[float]
    change: Optional[float]
    page: str = ""


@dataclass
class StatusRow:
    """정책/단위/세부사업별 현황 표 한 줄. level 0 = 기관."""

    level: int
    name: str
    after: Optional[float]
    before: Optional[float]
    change: Optional[float]


@dataclass
class SuppDocument:
    path: str = ""
    round_label: str = ""                # "2026년 제2회추경"
    projects: list[SuppProject] = field(default_factory=list)
    summary: list[SummaryRow] = field(default_factory=list)
    summary_total: Optional[SummaryRow] = None
    status: list[StatusRow] = field(default_factory=list)


# ------------------------------------------------------------------ 판별

def _head_text(rows, count: int = 4) -> str:
    return "".join("".join(row) for row in rows[:count]).replace(" ", "")


def is_change_table(rows) -> bool:
    """증감 산출내역 표 또는 미반영 사업 표인가."""
    head = _head_text(rows)
    return "사업명" in head and ("추가경정" in head or "기정예산" in head)


def is_year_table(rows) -> bool:
    """사업 머리 표. '2026년 제2회추경예산안(A) · 2026년 기정예산(B) · … · 증감(A-B)'"""
    if len(rows) > 4:
        return False
    head = _head_text(rows)
    return "추경예산안(A)" in head and "증감(A-B)" in head


def is_summary_table(rows) -> bool:
    head = _head_text(rows, 3)
    return "사업항목명" in head and "추가경정" in head and "쪽수" in head


def is_status_table(rows) -> bool:
    head = _head_text(rows, 4)
    return "정책/단위/세부사업" in head and "추가경정" in head


def is_supplement_file(path: str) -> bool:
    """추경 설명서인가. 화면에서 '2번 K-에듀파인 파일이 필요한가'를 정할 때 쓴다."""
    try:
        chunks = native_chunks(path)
    except Exception:  # noqa: BLE001 - 판별 실패는 '아니다'로 본다
        return False
    for chunk in chunks:
        if chunk.get("type") != "table":
            continue
        rows = _cells(chunk.get("text") or "")
        if is_year_table(rows) or is_change_table(rows):
            return True
    return False


# ------------------------------------------------------------------ 표 읽기

def _split_formula(text: str) -> tuple[str, str]:
    found = FORMULA_TAIL.match(text)
    if found and found.group(1).strip():
        return found.group(1).strip(), found.group(2).strip()
    return text.strip(), ""


def _group_row(cells: list[str]) -> Optional[tuple[int, tuple[float, Optional[float]]]]:
    """A 칸과 'B(전년)' 칸이 나란히 있는 줄이면 B 칸의 위치를 준다."""
    for index, cell in enumerate(cells):
        pair = with_previous(cell)
        if pair is None:
            continue
        if index >= 1 and number(cells[index - 1]) is not None:
            return index, pair
    return None


def _bullet_of(text: str) -> str:
    head = text[:1]
    if head and (head in DOT_BULLETS or head in ARROW_BULLETS or is_private_use(head)):
        return head
    return ""


def _parse_lines(raw: list[tuple[int, list[str]]]) -> list[SuppLine]:
    """① 아래 줄들의 계층을 정한다.

    ․ 는 1단. 사설영역 글리프는 처음 나온 자리의 한 단 아래로 정하고, 같은 글리프가
    다시 나오면 같은 단이다. 글머리가 아예 없는 줄('기본(160천원×…)')은 앞줄의
    한 단 아래다 — 설명서에서 글머리가 빠진 채 강사수당 아래 붙은 줄이다. 글머리 없는
    줄이 연달아 나오면 형제다.
    """
    lines: list[SuppLine] = []
    levels: dict[str, int] = {}
    previous_depth = 0
    previous_plain = False
    for row, cells in raw:
        text = cells[0].strip()
        bullet = _bullet_of(text)
        body = text[1:].strip() if bullet else text
        if bullet in DOT_BULLETS and bullet:
            depth = 1
            levels = {mark: level for mark, level in levels.items() if level <= depth}
        elif bullet:
            if bullet not in levels:
                levels[bullet] = previous_depth + 1 if previous_depth else 1
            depth = levels[bullet]
            levels = {mark: level for mark, level in levels.items() if level <= depth}
        elif previous_plain:
            depth = previous_depth
        else:
            depth = previous_depth + 1
        name, formula = _split_formula(body)
        amount = None
        for cell in cells[1:]:
            amount = marked(cell)
            if amount is None:
                amount = number(cell)
            if amount is not None:
                break
        lines.append(SuppLine(depth=depth, name=name, formula=formula, amount=amount, row=row,
                              bullet=bullet))
        previous_depth = depth
        previous_plain = not bullet
    return lines


def parse_change_table(rows: list[list[str]]) -> SuppTable:
    head = _head_text(rows)
    kind = "증감" if "산출내역" in head else "미반영"
    table = SuppTable(kind=kind)
    started = False
    current: Optional[SuppGroup] = None
    pending: list[tuple[int, list[str]]] = []

    def close() -> None:
        nonlocal pending
        if current is not None and pending:
            current.lines.extend(_parse_lines(pending))
        pending = []

    for number_, cells in enumerate(rows, start=1):
        cells = [cell.strip() for cell in cells]
        if not started:
            if "사업명" in "".join(cells).replace(" ", ""):
                started = True
            continue
        if not any(cells):
            continue
        joined = "".join(cells).replace(" ", "")
        if joined in ("산출내역요구액",):
            continue
        first = cells[0].replace(" ", "")
        if first in SUMMARY_WORDS:
            close()
            values = [cell for cell in cells[1:] if cell]
            total = SuppGroup(title="합계", row=number_)
            for cell in values:
                pair = with_previous(cell)
                if pair and total.before is None:
                    total.before, total.previous = pair
                elif number(cell) is not None:
                    if total.after is None:
                        total.after = number(cell)
                    elif total.change is None:
                        total.change = number(cell)
            table.total = total
            current = None
            continue
        found = _group_row(cells)
        if found:
            close()
            at, (before, previous) = found
            group = SuppGroup(title="", after=number(cells[at - 1]), before=before, previous=previous,
                              row=number_, business=cells[at - 2] if at >= 2 else "")
            rest = cells[at + 1:]
            if kind == "미반영":
                group.title = group.business
                group.change = number(rest[0]) if rest else None
                group.note = " ".join(rest[1:]).strip()
            else:
                label = rest[0] if rest else ""
                group.title, group.formula = _split_formula(label)
                numbers: list[float] = []
                for cell in rest[1:]:
                    own = marked(cell)
                    if own is not None and group.own is None and not numbers:
                        group.own = own
                        continue
                    value = number(cell)
                    if value is not None:
                        numbers.append(value)
                    elif cell:
                        group.note = (group.note + " " + cell).strip()
                group.change = numbers[0] if numbers else None
            table.groups.append(group)
            current = group
            continue
        if current is not None and kind == "증감":
            pending.append((number_, cells))
    close()
    return table


def _year_values(rows: list[list[str]]) -> tuple[list[str], list[Optional[float]]]:
    for index, cells in enumerate(rows):
        if any("추경예산안(A)" in cell.replace(" ", "") for cell in cells):
            labels = [cell.strip() for cell in cells]
            for below in rows[index + 1:]:
                values = [number(cell) for cell in below]
                if any(value is not None for value in values):
                    return labels, values
            return labels, []
    return [], []


def _summary_rows(rows: list[list[str]]) -> tuple[Optional[SummaryRow], list[SummaryRow]]:
    total = None
    found: list[SummaryRow] = []
    for cells in rows:
        cells = [cell.strip() for cell in cells]
        if cells and cells[0].replace(" ", "") in SUMMARY_WORDS and len(cells) >= 4:
            total = SummaryRow("", "합계", number(cells[1]), number(cells[2]), number(cells[3]))
            continue
        if len(cells) < 5:
            continue
        values = [number(cell) for cell in cells[2:5]]
        if any(value is None for value in values):
            continue
        found.append(SummaryRow(cells[0], cells[1], *values, page=cells[5] if len(cells) > 5 else ""))
    return total, found


def _status_rows(rows: list[list[str]]) -> list[StatusRow]:
    found: list[StatusRow] = []
    for cells in rows:
        level = 0
        while level < len(cells) and not cells[level].strip():
            level += 1
        rest = cells[level:]
        if len(rest) < 4:
            continue
        values = [number(cell) for cell in rest[1:4]]
        if any(value is None for value in values):
            continue
        found.append(StatusRow(level, rest[0].strip(), *values))
    return found


def parse_supplement(path: str) -> SuppDocument:
    """추경 사업별 설명서(HWPX)를 읽는다."""
    document = SuppDocument(path=path)
    chunks = native_chunks(path)
    current: Optional[SuppProject] = None
    heading = heading_number = heading_line = ""
    part_title = ""
    section = ""
    section_cost: Optional[float] = None
    order = 0
    for chunk in chunks:
        text = chunk.get("text") or ""
        if chunk.get("type") == "text":
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            for line in lines:
                found = HEADING.match(line)
                if found:
                    # '1. 부서기본운영경비' 는 사업계획 절 제목이기도 하고, 조직 표가 바로 뒤따르면
                    # 새 사업의 제목이기도 하다. 둘 다 적어 두고 다음 표를 보고 정한다.
                    heading_number, heading = found.group(1), found.group(2)
                    heading_line = line
                    part_title = ""
                    section, section_cost = found.group(2), None
                cost = PLAN_COST.search(line)
                if cost and "총사업비" not in line:
                    value = number(cost.group(1).replace(" ", ""))
                    section_cost = value if cost.group(2) == "천원" or value is None else value / 1000
            if current is not None:
                period = PERIOD.search(text)
                if period and not current.period:
                    current.period = (int(period.group(1)), int(period.group(2)))
                total = TOTAL_COST.search(text)
                if total and current.total_cost is None:
                    current.total_cost = number(total.group(1))
                    current.funding = {label: float(value.replace(",", ""))
                                       for label, value in re.findall(r"([가-힣]+)\s*([\d,]+)\s*천원",
                                                                      text[total.end():].split("\n")[0])}
            if lines and not HEADING.match(lines[-1]) and _looks_like_title(lines[-1]):
                part_title = lines[-1]
            continue
        rows = _cells(text)
        if not rows:
            continue
        if is_summary_table(rows):
            document.summary_total, document.summary = _summary_rows(rows)
            continue
        if is_status_table(rows):
            document.status = _status_rows(rows)
            continue
        if _is_organisation(rows):
            policy, unit, program, organisation = _organisation(rows)
            name = part_title or heading or "사업명 확인 필요"
            current = SuppProject(name=name, number="" if part_title else heading_number,
                                  heading=heading_line, policy=policy, unit=unit, program=program,
                                  organisation=organisation)
            document.projects.append(current)
            part_title = ""
            section, section_cost = "", None
            continue
        if current is None:
            continue
        if is_year_table(rows):
            labels, values = _year_values(rows)
            current.header_labels = labels
            if not document.round_label and labels:
                document.round_label = labels[0].replace("예산안(A)", "").strip()
            padded = (values + [None] * 4)[:4]
            current.after, current.before, current.last_final, current.change = padded
            continue
        if is_change_table(rows):
            order += 1
            table = parse_change_table(rows)
            table.section, table.section_cost, table.order = section, section_cost, order
            current.tables.append(table)
            section_cost = None
    return document


# ------------------------------------------------------------------ 검사 1 · 설명서 안

GOAL_PLAN = "목표1"
GOAL_UBIS = "목표2A"
GOAL_CHAIN = "추경연결"


def _where(table: SuppTable) -> str:
    kind = "미반영 사업 표" if table.kind == "미반영" else "증감 산출내역 표"
    return f"{table.section} · {kind}" if table.section else kind


def _check_line_tree(report: Report, project: SuppProject, table: SuppTable, group: SuppGroup) -> None:
    lines = group.lines
    for index, line in enumerate(lines):
        if line.formula:
            value = calc(line.formula)
            if value is not None and line.amount is not None and not _formula_ok(value, line.amount):
                report.add(GOAL_PLAN, "오류", project.name, line.name,
                           f"[{_where(table)} {line.row}행] '{line.name}' 산출식 {line.formula} 을 계산하면 "
                           f"{_fmt(value)}천원인데 {_fmt(line.amount)}천원이 적혀 있습니다.",
                           line.amount, value, row=line.row, form="추경", block=table.order)
        children = []
        for later in lines[index + 1:]:
            if later.depth <= line.depth:
                break
            if later.depth == line.depth + 1:
                children.append(later)
        if children and line.amount is not None:
            total = sum(child.amount or 0 for child in children)
            if not _same(total, line.amount):
                report.add(GOAL_PLAN, "오류", project.name, line.name,
                           f"[{_where(table)} {line.row}행] '{line.name}' 은 {_fmt(line.amount)}천원인데 "
                           f"바로 아래 {len(children)}줄을 더하면 {_fmt(total)}천원입니다.",
                           line.amount, total, row=line.row, form="추경", block=table.order)
    top = [line for line in lines if line.depth == 1]
    target = group.own if group.own is not None else group.change
    if top and target is not None:
        total = sum(line.amount or 0 for line in top)
        if not _same(total, target):
            report.add(GOAL_PLAN, "오류", project.name, group.title,
                       f"[{_where(table)} {group.row}행] '{group.title}' 증감 {_fmt(target)}천원 · "
                       f"아래 산출 줄 {len(top)}개의 합 {_fmt(total)}천원.",
                       target, total, row=group.row, form="추경", block=table.order)


def _check_group(report: Report, project: SuppProject, table: SuppTable, group: SuppGroup) -> None:
    title = group.title or group.business or "합계"
    if group.after is not None and group.before is not None and group.change is not None:
        if not _same(group.after - group.before, group.change):
            report.add(GOAL_PLAN, "오류", project.name, title,
                       f"[{_where(table)} {group.row}행] '{title}' 추경안 {_fmt(group.after)} − 기정 "
                       f"{_fmt(group.before)} = {_fmt(group.after - group.before)}천원인데 증감이 "
                       f"{_fmt(group.change)}천원으로 적혀 있습니다.",
                       group.change, group.after - group.before, row=group.row, form="추경", block=table.order)
    if group.own is not None and group.change is not None and not _same(group.own, group.change):
        report.add(GOAL_PLAN, "오류", project.name, title,
                   f"[{_where(table)} {group.row}행] '{title}' 산출내역 칸 {_fmt(group.own)}천원과 요구액 칸 "
                   f"{_fmt(group.change)}천원이 다릅니다.", group.change, group.own, row=group.row, form="추경", block=table.order)
    if group.formula:
        value = calc(group.formula)
        target = group.own if group.own is not None else group.change
        if value is not None and target is not None and not _formula_ok(value, target):
            report.add(GOAL_PLAN, "오류", project.name, title,
                       f"[{_where(table)} {group.row}행] '{title}' 산출식 {group.formula} 을 계산하면 "
                       f"{_fmt(value)}천원인데 {_fmt(target)}천원이 적혀 있습니다.",
                       target, value, row=group.row, form="추경", block=table.order)
    if table.kind == "미반영" and group.change is not None and not _same(group.change, 0):
        report.add(GOAL_PLAN, "오류", project.name, title,
                   f"[{_where(table)} {group.row}행] 추경에 반영하지 않은 사업인데 증감이 {_fmt(group.change)}천원 "
                   "적혀 있습니다. 미반영 사업은 증감이 0이어야 합니다.",
                   group.change, 0, row=group.row, form="추경", block=table.order)
    _check_line_tree(report, project, table, group)


def _check_table(report: Report, project: SuppProject, table: SuppTable) -> None:
    for group in table.groups:
        _check_group(report, project, table, group)
    total = table.total
    if total is None:
        return
    _check_group(report, project, table, total)
    for key, word in (("after", "추경안(A)"), ("before", "기정(B)"), ("change", "증감"),
                      ("previous", "전년도 예산액")):
        written, summed = getattr(total, key), table.sum_of(key)
        if written is None or summed is None or _same(written, summed):
            continue
        report.add(GOAL_PLAN, "오류", project.name, "합계",
                   f"[{_where(table)} 합계 줄] {word} 합계 {_fmt(written)}천원 · 아래 {len(table.groups)}줄의 합 "
                   f"{_fmt(summed)}천원.", written, summed, row=total.row, form="추경", block=table.order)


def _check_project(report: Report, project: SuppProject) -> None:
    if project.after is None or project.change is None:
        report.add(GOAL_PLAN, SKIPPED, project.name, "",
                   "사업 머리의 '추경예산안(A) · 기정예산(B) · 증감' 표를 찾지 못해 사업 합계 검산을 건너뜁니다.")
    elif project.before is not None and not _same(project.after - project.before, project.change):
        report.add(GOAL_PLAN, "오류", project.name, "",
                   f"[사업 머리 표] 추경안 {_fmt(project.after)} − 기정 {_fmt(project.before)} = "
                   f"{_fmt(project.after - project.before)}천원인데 증감이 {_fmt(project.change)}천원입니다.",
                   project.change, project.after - project.before, form="추경")
    if (project.total_cost is not None and project.after is not None
            and not _same(project.total_cost, project.after) and project.single_year):
        # 여러 해 사업(계속비 등)은 총사업비가 올해 금액보다 큰 게 정상이다. 한 해 사업인데
        # 다르면 추경 뒤 금액으로 고치지 않은 것이다. 기정과 같으면 그 뜻이 더 분명하다.
        stale = _same(project.total_cost, project.before)
        report.add(GOAL_PLAN, "오류", project.name, "",
                   f"사업개요의 총사업비 {_fmt(project.total_cost)}천원 · 추경예산안(A) {_fmt(project.after)}천원. "
                   f"사업기간이 {project.period[0]}년 한 해이므로 둘은 같아야 합니다."
                   + ("  총사업비가 기정예산과 같습니다 — 추경 뒤 금액으로 고치지 않은 것 같습니다." if stale else ""),
                   project.total_cost, project.after, form="추경")
    if project.funding and project.total_cost is not None:
        funded = sum(project.funding.values())
        if len(project.funding) > 1 and not _same(funded, project.total_cost):
            report.add(GOAL_PLAN, "오류", project.name, "",
                       "총사업비 재원 내역 " + " + ".join(f"{k} {_fmt(v)}" for k, v in project.funding.items())
                       + f" = {_fmt(funded)}천원인데 총사업비는 {_fmt(project.total_cost)}천원입니다.",
                       project.total_cost, funded, form="추경")

    for table in project.tables:
        _check_table(report, project, table)
        cost, change = table.section_cost, table.value("change")
        if table.kind == "증감" and cost is not None and change is not None and not _same(cost, change):
            report.add(GOAL_PLAN, "오류", project.name, table.section,
                       f"사업계획 '{table.section}' 의 사업비는 {_fmt(cost)}천원인데 아래 산출내역 표의 증감 합계는 "
                       f"{_fmt(change)}천원입니다.", cost, change, form="추경", block=table.order)

    if not project.tables:
        if project.change is not None and not _same(project.change, 0):
            report.add(GOAL_PLAN, SKIPPED, project.name, "",
                       "증감 산출내역 표가 없어 사업 안쪽 검산을 건너뜁니다. (총액배분 사업은 지역별 표로 적기도 합니다)")
        return
    sums = {}
    for key in ("after", "before", "change"):
        values = [table.value(key) for table in project.tables]
        sums[key] = None if any(value is None for value in values) else sum(values)
    words = {"after": "추경안(A)", "before": "기정(B)", "change": "증감"}
    heads = {"after": project.after, "before": project.before, "change": project.change}
    change_ok = _same(sums["change"], heads["change"])
    for key in ("change", "after", "before"):
        if sums[key] is None or heads[key] is None or _same(sums[key], heads[key]):
            continue
        # 증감이 맞고 미반영 사업 표가 없으면, A·B 차이는 손대지 않은 기존 세부사업을 표에
        # 싣지 않았기 때문이다(그 표가 없으니 실을 자리가 없다). 오류가 아니다. 미반영 표가
        # 있는데도 모자라면 거기서 빠뜨린 것이다.
        if key != "change" and change_ok and not project.has_unreflected:
            report.add(GOAL_PLAN, "확인 필요", project.name, "", "미반영 표 없음 — 기존 세부사업 생략")
            continue
        severity = "오류"
        extra = ("" if key == "change" else
                 "  증감은 맞습니다. 미반영 사업 표에 빠진 기존 세부사업이 있습니다." if change_ok else "")
        report.add(GOAL_PLAN, severity, project.name, "",
                   f"사업 머리 표의 {words[key]} {_fmt(heads[key])}천원 · 사업계획 표 {len(project.tables)}개를 "
                   f"더하면 {_fmt(sums[key])}천원 — {_fmt(abs(heads[key] - sums[key]))}천원 차이.{extra}",
                   heads[key], sums[key], form="추경")


def _bigger_unit(row, project: SuppProject, ubis) -> bool:
    """총괄 표 줄이 사업 쪽보다 큰 단위(UBIS 사업 전체)의 금액을 적은 것인가.

    2026 제2회 추경 '이음교육 운영 지원(총액배분사업비)' — 쪽은 274,000/309,000, 총괄과
    UBIS 는 964,500/999,500, 증감은 셋 다 △35,000. UBIS 가 있으면 UBIS 기정액이 총괄과
    같은지로 가리고, 없으면 총액배분·재원배분 사업인지로 가린다(표본에서 이 모양은 그뿐이었다).
    """
    if ubis is not None:
        found, _alias = _find_row(ubis.rows, project.name)
        return found is not None and _same(found.before, row.before)
    return "총액배분" in project.name or "재원배분" in project.name


def _check_summary(report: Report, document: SuppDocument, ubis=None) -> None:
    rows = document.summary
    if not rows:
        report.add(GOAL_PLAN, SKIPPED, "총괄", "", "총괄 표를 찾지 못해 총괄 대조를 건너뜁니다.")
        return
    for row in rows:
        if row.after is not None and row.before is not None and row.change is not None \
                and not _same(row.after - row.before, row.change):
            report.add(GOAL_PLAN, "오류", row.name, "총괄",
                       f"[총괄 표] '{row.name}' 추경 {_fmt(row.after)} − 기정 {_fmt(row.before)} = "
                       f"{_fmt(row.after - row.before)}천원인데 증감이 {_fmt(row.change)}천원입니다.",
                       row.change, row.after - row.before, form="추경")
    total = document.summary_total
    if total is not None:
        for key, word in (("after", "추가경정 예산액"), ("before", "기정 예산액"), ("change", "비교증감")):
            written = getattr(total, key)
            summed = sum(getattr(row, key) or 0 for row in rows)
            if written is not None and not _same(written, summed):
                report.add(GOAL_PLAN, "오류", "총괄", "합계",
                           f"[총괄 표] {word} 합계 {_fmt(written)}천원 · 사업별 {len(rows)}줄의 합 {_fmt(summed)}천원.",
                           written, summed, form="추경")
    by_number = {row.order: row for row in rows if row.order}
    for project in document.projects:
        row = by_number.get(project.number) if project.number else None
        if row is None or _squeeze_name(row.name) != _squeeze_name(project.name):
            row = next((one for one in rows if _squeeze_name(one.name) == _squeeze_name(project.name)), row)
        if row is None:
            report.add(GOAL_PLAN, "오류", project.name, "총괄",
                       "총괄 표에 이 사업이 없습니다. 사업 쪽은 있는데 총괄 표에서 빠졌습니다.", form="추경")
            continue
        diffs = [(word, getattr(project, key), getattr(row, key))
                 for key, word in (("change", "증감"), ("after", "추경안"), ("before", "기정"))
                 if getattr(project, key) is not None and getattr(row, key) is not None
                 and not _same(getattr(project, key), getattr(row, key))]
        if not diffs:
            continue
        change_ok = diffs[0][0] != "증감"
        if change_ok and _bigger_unit(row, project, ubis):
            # 총괄 표가 UBIS 사업 단위(이 쪽보다 큰 묶음)로 적혔다. 증감이 같으니 틀린 게 아니다.
            report.add(GOAL_PLAN, "확인 필요", project.name, "총괄", "총괄은 UBIS 사업 단위")
            continue
        # 한 군데 잘못이 증감·추경안 두 건으로 불어나지 않게 한 줄로 말한다.
        word, mine, theirs = diffs[0]
        rest = "".join(f" · {other} {_fmt(t)} / {_fmt(m)}" for other, m, t in diffs[1:])
        report.add(GOAL_PLAN, "오류", project.name, "총괄",
                   f"[총괄 표 {row.order}번] {word} {_fmt(theirs)}천원 · 사업 쪽 {_fmt(mine)}천원 — "
                   f"{_fmt(abs(theirs - mine))}천원 차이." + (f" (총괄 / 사업 쪽{rest})" if rest else "")
                   + ("  증감은 같지만 총괄 표와 사업 쪽의 금액이 다릅니다." if change_ok else ""),
                   mine, theirs, form="추경")


def _check_status(report: Report, document: SuppDocument) -> None:
    rows = document.status
    for index, row in enumerate(rows):
        if not _same(row.after - row.before, row.change):
            report.add(GOAL_PLAN, "오류", row.name, "현황",
                       f"[정책/단위/세부사업별 현황] '{row.name}' 추경 {_fmt(row.after)} − 기정 {_fmt(row.before)} = "
                       f"{_fmt(row.after - row.before)}천원인데 증감이 {_fmt(row.change)}천원입니다.",
                       row.change, row.after - row.before, form="추경")
        children = []
        for later in rows[index + 1:]:
            if later.level <= row.level:
                break
            if later.level == row.level + 1:
                children.append(later)
        if not children:
            continue
        for key, word in (("after", "추경"), ("before", "기정"), ("change", "증감")):
            summed = sum(getattr(child, key) for child in children)
            if not _same(getattr(row, key), summed):
                report.add(GOAL_PLAN, "오류", row.name, "현황",
                           f"[정책/단위/세부사업별 현황] '{row.name}' {word} {_fmt(getattr(row, key))}천원 · "
                           f"아래 {len(children)}개의 합 {_fmt(summed)}천원.",
                           getattr(row, key), summed, form="추경")


def _check_duplicates(report: Report, document: SuppDocument) -> None:
    counts = Counter(_squeeze_name(project.name) for project in document.projects)
    for project in document.projects:
        if counts[_squeeze_name(project.name)] > 1:
            report.add(GOAL_PLAN, "오류", project.name, "",
                       "같은 이름의 사업 쪽이 설명서에 두 번 이상 실려 있습니다. 겹친 쪽을 지워 주세요.", form="추경")
            counts[_squeeze_name(project.name)] = 0


def check_supplement(document: SuppDocument, ubis=None) -> Report:
    """설명서 안에서 맞아야 할 것. 값을 고치지 않고 지적만 한다.

    ubis(UbisSupp)를 주면 총괄 표와 사업 쪽이 어긋날 때 그게 오류인지 가리는 데 쓴다.
    """
    report = Report()
    for project in document.projects:
        _check_project(report, project)
    _check_summary(report, document, ubis)
    _check_status(report, document)
    _check_duplicates(report, document)
    return report


# ------------------------------------------------------------------ UBIS 추경 검토조서

@dataclass
class UbisHeadline:
    """◎ 줄 하나와 그 아래 '- 세부' 줄들."""

    label: str
    amount: Optional[float]
    details: list = field(default_factory=list)     # (이름, 금액)


@dataclass
class UbisSuppRow:
    name: str                       # 사업명. 비었으면 사업항목(세세부사업)
    item: str = ""                  # 사업항목1(세세부사업)
    program: str = ""               # 세부사업
    before: float = 0.0             # 기정액(A)
    change: float = 0.0             # 추경요구액(B) = 증감
    adjust: float = 0.0             # 조정액(D). 예산과가 요구와 따로 얹거나 뺀 금액
    row: int = 0
    headlines: list[UbisHeadline] = field(default_factory=list)

    @property
    def after(self) -> float:
        return self.before + self.change

    @property
    def unplanned(self) -> bool:
        """'◎ 미편성사업' — 이번 추경에 손대지 않은 사업."""
        return not self.headlines and _same(self.change, 0)

    def pieces(self) -> list[float]:
        """세부 금액. '- 세부' 줄이 있으면 그 줄들, 없으면 ◎ 줄 자체."""
        found: list[float] = []
        for headline in self.headlines:
            values = [value for _label, value in headline.details if value is not None]
            if values:
                found.extend(values)
            elif headline.amount is not None:
                found.append(headline.amount)
        return found


@dataclass
class UbisSupp:
    path: str = ""
    title: str = ""
    total: Optional[UbisSuppRow] = None
    rows: list[UbisSuppRow] = field(default_factory=list)


def _flat(value) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def is_supplement_review(path: str) -> bool:
    """UBIS 검토조서가 추경용인가. '추경요구액' 열이 있으면 추경이다."""
    try:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for row in book.active.iter_rows(min_row=1, max_row=10, values_only=True):
                if any("추경요구액" in _flat(value) for value in row):
                    return True
        finally:
            book.close()
    except Exception:  # noqa: BLE001
        return False
    return False


AMOUNT_IN_TEXT = re.compile(r"([△▲-]?\s*[\d,]+)\s*천원\s*$")
REVIEW_KEYS = (("사업명", "사업명"), ("사업항목", "세세부사업"), ("세부사업", "세부사업"),
               ("기관명", "기관명"), ("기정액", "기정액"), ("요구액", "추경요구액"),
               ("본예산요구액", "본예산요구액"), ("조정액", "조정액"), ("산출기초", "산출기초"))


def _review_columns(sheet) -> tuple[int, dict[str, int]]:
    """검토조서(본예산·추경 공통) 머리글 행과 열 위치."""
    for row in range(1, min(10, sheet.max_row) + 1):
        labels = {column: _flat(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1)}
        if not any("사업명" in label for label in labels.values()):
            continue
        found: dict[str, int] = {}
        for column, label in labels.items():
            for key, needle in REVIEW_KEYS:
                if key in found or needle not in label:
                    continue
                if key == "세부사업" and "세세부" in label:
                    continue
                found[key] = column
                break
            if label == "재원배분여부":
                found["재원배분"] = column
            elif label == "총액배분사업비":
                found["총액배분"] = column
        return row, found
    return 0, {}


def read_supplement_review(path: str) -> UbisSupp:
    """UBIS 추경 검토조서(서식1). 열은 위치가 아니라 머리글 이름으로 찾는다."""
    from openpyxl import load_workbook

    sheet = load_workbook(path, data_only=True).active
    header, found = _review_columns(sheet)
    if not header or "기정액" not in found or "요구액" not in found:
        raise ValueError("UBIS 추경 검토조서가 아닙니다. '기정액'과 '추경요구액' 열을 찾지 못했습니다.")
    result = UbisSupp(path=path, title=str(sheet.cell(1, 1).value or "").strip())
    current: Optional[UbisSuppRow] = None
    for row in range(header + 1, sheet.max_row + 1):
        def cell(key: str):
            return sheet.cell(row, found[key]).value if key in found else None

        name = str(cell("사업명") or "").strip()
        item = str(cell("사업항목") or "").strip()
        program = str(cell("세부사업") or "").strip()
        before, change, adjust = cell("기정액"), cell("요구액"), cell("조정액")
        text = str(cell("산출기초") or "").strip()
        if isinstance(before, (int, float)) and (program or item or name or str(cell("기관명") or "").strip()):
            display = name or item
            if not name and item:
                # 사업명이 빈 줄은 사업항목으로 부른다. 본사업·재원배분·총액배분이 같은
                # 사업항목을 쓰므로 표시 열의 꼬리표를 붙여 서로 다른 이름이 되게 한다.
                for key, tag in (("총액배분", "총액배분사업비"), ("재원배분", "재원배분")):
                    if tag in _flat(cell(key)):
                        display = f"{item}({tag})"
                        break
            entry = UbisSuppRow(name=display, item=item, program=program, before=float(before),
                                change=float(change) if isinstance(change, (int, float)) else 0.0,
                                adjust=float(adjust) if isinstance(adjust, (int, float)) else 0.0, row=row)
            if not program and not item and not name:
                # 맨 위 기관 합계 줄
                if result.total is None:
                    result.total = entry
                current = None
                continue
            result.rows.append(entry)
            current = entry
        elif isinstance(before, (int, float)) and not text:
            current = None          # 맨 아래 합계 줄
            continue
        if current is None or not text:
            continue
        amount = AMOUNT_IN_TEXT.search(text)
        value = number(amount.group(1).replace(" ", "")) if amount else None
        label = AMOUNT_IN_TEXT.sub("", text).strip()
        if text.startswith("◎"):
            if "미편성" not in text:
                current.headlines.append(UbisHeadline(label.lstrip("◎").strip(), value))
        elif text.lstrip().startswith("-") and current.headlines:
            current.headlines[-1].details.append((label.lstrip(" -").rstrip(": ").strip(), value))
    return result


STATUS_WORDS = ("신규", "성립전", "목변경", "계속", "포함")


def match_key(name: str) -> str:
    """사업명을 맞춰 볼 열쇠. 띄어쓰기와 '(신규)' '(성립전 포함)' '(목변경)' 같은 꼬리표를 뗀다.

    설명서는 '제지출금 등', UBIS 는 '제지출금 등(신규)' 처럼 한쪽에만 꼬리표를 단다.
    '(재원배분)' '(총액배분사업비)' 는 떼지 않는다. 같은 이름의 **다른** 사업을 가르는 말이다.
    """
    def keep(found) -> str:
        parts = [part.strip() for part in found.group(1).split(",")]
        parts = [part for part in parts if part and not any(word in part for word in STATUS_WORDS)]
        return f"({','.join(parts)})" if parts else ""

    return _squeeze_name(re.sub(r"\(([^()]*)\)", keep, name or ""))


def _find_row(rows: list[UbisSuppRow], name: str) -> tuple[Optional[UbisSuppRow], str]:
    exact = [row for row in rows if row.name.strip() == name.strip()]
    if exact:
        return exact[0], ""
    for key in (_squeeze_name, match_key):
        loose = [row for row in rows if key(row.name) == key(name)]
        if len(loose) > 1:
            # 사업명이 빈 줄은 사업항목으로 이름을 대신해서, 본사업·재원배분·총액배분의
            # '미편성' 줄이 같은 이름으로 여럿 생긴다. 이번에 손댄 줄만 남긴다.
            loose = [row for row in loose if not row.unplanned]
        if len(loose) == 1:
            return loose[0], loose[0].name
    return None, ""


def _multiset_gap(left: list[float], right: list[float]) -> tuple[list[float], list[float]]:
    remaining = list(right)
    only_left: list[float] = []
    for value in left:
        match = next((index for index, other in enumerate(remaining) if _same(value, other)), None)
        if match is None:
            only_left.append(value)
        else:
            remaining.pop(match)
    return only_left, remaining


def _pieces(project: SuppProject) -> list[float]:
    """설명서 ① 항목 중 증감이 있는 것."""
    return [group.change for table in project.tables if table.kind == "증감" for group in table.groups
            if group.change is not None and not _same(group.change, 0)]


def compare_supplement(document: SuppDocument, ubis: UbisSupp) -> Report:
    """설명서 ↔ UBIS 추경 검토조서. 오류만 싣는다.

      - UBIS 추경요구액(B) ≠ 설명서 증감                 오류. 어느 ① 항목이 다른지도 적는다.
      - UBIS 기정액(A) ≠ 설명서 기정예산(B)              총괄 표가 UBIS 와 같으면 쪽이 UBIS 사업의
                                                        일부만 적은 것(총액배분)이라 오류 아님, 아니면 오류
      - UBIS 에 추경 요구가 있는데 설명서 쪽도 총괄 줄도 없음   오류
      - 설명서에 증감이 있는데 UBIS 에 그 사업이 없음           오류
      - 기관 합계 ≠ 설명서 총괄 표 합계                  오류
    세부 금액 묶음이 달라도 사업 증감이 같으면 묶는 단위가 다른 것뿐이라 싣지 않는다.
    """
    report = Report()
    used: set[int] = set()
    summary_by_key = {match_key(row.name): row for row in document.summary}
    for project in document.projects:
        row, _alias = _find_row(ubis.rows, project.name)
        if row is None:
            if project.change is not None and not _same(project.change, 0):
                report.add(GOAL_UBIS, "오류", project.name, "",
                           f"설명서에는 증감 {_fmt(project.change)}천원이 있는데 UBIS 추경 검토조서에 이 사업이 "
                           "없습니다. UBIS 에 요구를 넣지 않았거나 사업명이 다릅니다.", project.change, None)
            continue
        used.add(row.row)
        if project.change is not None and not _same(project.change, row.change):
            only_mine, only_theirs = _multiset_gap(_pieces(project),
                                                   [value for value in row.pieces() if not _same(value, 0)])
            where = ""
            if only_mine or only_theirs:
                where = (f"  어긋난 세부: 설명서 {', '.join(_fmt(v) for v in only_mine) or '없음'} · "
                         f"UBIS {', '.join(_fmt(v) for v in only_theirs) or '없음'} (천원).")
            report.add(GOAL_UBIS, "오류", project.name, "",
                       f"한글 설명서 증감 {_fmt(project.change)}천원 · 엑셀 UBIS 추경요구액(B) {_fmt(row.change)}천원 "
                       f"— {_fmt(abs(project.change - row.change))}천원 차이.{where}", project.change, row.change)
        if project.before is not None and not _same(project.before, row.before):
            summary = summary_by_key.get(match_key(project.name))
            if summary is not None and _same(summary.before, row.before):
                continue          # 총괄 = UBIS. 이 쪽은 UBIS 사업의 일부만 적었다(총액배분)
            report.add(GOAL_UBIS, "오류", project.name, "",
                       f"한글 설명서 기정예산(B) {_fmt(project.before)}천원 · 엑셀 UBIS 기정액(A) {_fmt(row.before)}천원 "
                       f"— {_fmt(abs(project.before - row.before))}천원 차이. 기정액은 직전에 확정된 금액이라 "
                       "두 곳이 같아야 합니다.", project.before, row.before)
    for row in ubis.rows:
        if row.row in used or _same(row.change, 0):
            continue
        key = match_key(row.name)
        summary = next((one for one in document.summary
                        if key and key in match_key(one.name) and _same(one.change, row.change)), None)
        if summary is not None:
            continue              # 총괄 표에 같은 증감으로 실려 있다(사업 쪽을 따로 쓰지 않는 사업)
        report.add(GOAL_UBIS, "오류", row.name, "",
                   f"UBIS 에는 추경 요구 {_fmt(row.change)}천원이 있는데 설명서에는 이 사업 쪽도 총괄 표 줄도 "
                   "없습니다.", None, row.change)
    total = document.summary_total
    if ubis.total is not None and total is not None:
        for mine, theirs, word in ((total.change, ubis.total.change, "증감(추경요구액)"),
                                   (total.before, ubis.total.before, "기정")):
            if mine is not None and not _same(mine, theirs):
                report.add(GOAL_UBIS, "오류", "총괄", "합계",
                           f"설명서 총괄 표 {word} 합계 {_fmt(mine)}천원 · UBIS 기관 합계 {_fmt(theirs)}천원 — "
                           f"{_fmt(abs(mine - theirs))}천원 차이.", mine, theirs)
    return report


# ------------------------------------------------------------------ 직전 차수와 이어 보기

def round_of(name: str) -> Optional[tuple[int, int]]:
    """파일 이름에서 (연도, 차수). 본예산 0, 1회 추경 1, 2회 추경 2. 모르면 None.

    '2026추경예산1차검토조서' · '2026추경예산2차검토조서' · '2026본예산검토조서' ·
    '2026추경사업별 설명서'(차수 없음 = 1차) · '2026추경2차사업별 설명서'
    """
    stem = Path(name).stem
    year = re.search(r"(?<!\d)(20\d{2})(?![\d.])", stem)
    if not year:
        return None
    rest = stem[year.end():]
    if "본예산" in rest[:6]:
        return int(year.group(1)), 0
    if "추경" in rest[:8]:
        after = rest[rest.index("추경"):]
        found = re.search(r"(\d)\s*차", after[:10]) or re.search(r"제\s*(\d)\s*회", after[:10])
        return int(year.group(1)), int(found.group(1)) if found else 1
    return None


def round_from_title(title: str) -> Optional[tuple[int, int]]:
    """검토조서 1행 제목 '2026년도 제2회 추경 요구사업 검토조서' → (2026, 2)."""
    year = re.search(r"(20\d{2})", title or "")
    if not year:
        return None
    if "본예산" in title:
        return int(year.group(1)), 0
    found = re.search(r"제\s*(\d)\s*회\s*추경", title)
    return (int(year.group(1)), int(found.group(1))) if found else None


def _title_of(path: Path) -> str:
    try:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for row in book.active.iter_rows(min_row=1, max_row=1, values_only=True):
                return str(row[0] or "")
        finally:
            book.close()
    except Exception:  # noqa: BLE001
        return ""
    return ""


def find_previous_review(path: str) -> str:
    """같은 폴더에서 바로 앞 차수의 UBIS 검토조서를 찾는다. 없으면 빈 문자열.

    1회 추경의 앞은 그해 본예산, 2회 추경의 앞은 1회 추경이다. 차수는 파일 1행 제목
    ('2026년도 제2회 추경 요구사업 검토조서')으로 정하고, 제목이 없으면 파일 이름으로 본다.
    """
    here = round_from_title(_title_of(Path(path))) or round_of(Path(path).name)
    if here is None or here[1] == 0:
        return ""
    wanted = (here[0], here[1] - 1)
    candidates = []
    for other in Path(path).parent.glob("*.xlsx"):
        if other.name.startswith("~$") or other.resolve() == Path(path).resolve():
            continue
        if "검토조서" not in other.name:
            continue
        if (round_from_title(_title_of(other)) or round_of(other.name)) == wanted:
            candidates.append(other)
    if not candidates:
        return ""
    return str(max(candidates, key=lambda one: one.stat().st_mtime))


def review_endings(path: str) -> dict[tuple[str, str], float]:
    """검토조서 → (세부사업, 사업항목) → 그 차수가 끝난 금액.

    본예산은 요구액(B), 추경은 기정액(A) + 추경요구액(B). 조정액(D)은 넣지 않는다 —
    2025·2026년 네 차수 모두 D 를 빼야 다음 차수 기정액과 1천원도 틀리지 않고 이어졌다.

    사업명으로 묶지 않는다. 같은 사업항목이 본사업·재원배분·총액배분 세 줄로 나뉘고,
    사업명은 차수마다 '특수교육 교재교구 지원' / '특수교육교재교구지원' 처럼 바뀐다.
    """
    from openpyxl import load_workbook

    sheet = load_workbook(path, data_only=True).active
    header, found = _review_columns(sheet)
    supplement = "요구액" in found and "기정액" in found
    if not header or "세부사업" not in found or not (supplement or "본예산요구액" in found):
        raise ValueError(f"'{Path(path).name}' 에서 세부사업·요구액 열을 찾지 못했습니다.")
    endings: dict[tuple[str, str], float] = {}
    for row in range(header + 1, sheet.max_row + 1):
        program = _flat(sheet.cell(row, found["세부사업"]).value)
        if not program:
            continue
        item = _flat(sheet.cell(row, found["사업항목"]).value) if "사업항목" in found else ""
        if supplement:
            before = sheet.cell(row, found["기정액"]).value
            change = sheet.cell(row, found["요구액"]).value
            if not isinstance(before, (int, float)):
                continue
            value = float(before) + (float(change) if isinstance(change, (int, float)) else 0.0)
        else:
            value = sheet.cell(row, found["본예산요구액"]).value
            if not isinstance(value, (int, float)):
                continue
        endings[(program, item)] = endings.get((program, item), 0.0) + float(value)
    return endings


def compare_previous(current_path: str, previous_path: str) -> Report:
    """이번 추경의 기정액(A) = 직전 차수가 끝난 금액인지. (세부사업, 사업항목) 단위로 본다."""
    from openpyxl import load_workbook

    report = Report()
    label = Path(previous_path).name

    sheet = load_workbook(current_path, data_only=True).active
    header, found = _review_columns(sheet)
    befores: dict[tuple[str, str], float] = {}
    names: dict[tuple[str, str], str] = {}
    for row in range(header + 1, sheet.max_row + 1):
        program = _flat(sheet.cell(row, found["세부사업"]).value)
        before = sheet.cell(row, found["기정액"]).value
        if not program or not isinstance(before, (int, float)):
            continue
        item = _flat(sheet.cell(row, found["사업항목"]).value) if "사업항목" in found else ""
        befores[(program, item)] = befores.get((program, item), 0.0) + float(before)
        name = str(sheet.cell(row, found["사업명"]).value or "").strip() if "사업명" in found else ""
        names.setdefault((program, item), name or item or program)
    previous = review_endings(previous_path)
    report.add(GOAL_CHAIN, "안내", "총괄", "",
               f"직전 차수 자료로 '{label}' 을 썼습니다. 이번 기정액(A)이 그 차수가 끝난 금액과 같은지 "
               f"세부사업·사업항목 {len(befores)}개 단위로 봤습니다.")
    for key, before in befores.items():
        if key not in previous:
            if not _same(before, 0):
                report.add(GOAL_CHAIN, "오류", names[key], "",
                           f"직전 자료에 이 세부사업·사업항목({key[0]} / {key[1] or '없음'})이 없는데 기정액이 "
                           f"{_fmt(before)}천원입니다.", before, None)
            continue
        if not _same(before, previous[key]):
            report.add(GOAL_CHAIN, "오류", names[key], "",
                       f"UBIS 기정액(A) {_fmt(before)}천원 · 직전 차수가 끝난 금액 {_fmt(previous[key])}천원 — "
                       f"{_fmt(abs(before - previous[key]))}천원 차이. ({key[0]} / {key[1] or '사업항목 없음'})",
                       before, previous[key])
    return report
