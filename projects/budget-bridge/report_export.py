"""지적 목록을 엑셀로 내보낸다. 팀 회람·보고용."""

from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ui_common import goal_label

HEADERS = ["구분", "검사", "사업", "항목", "내용", "설명서(천원)", "계산·비교(천원)", "차액(천원)"]
FILLS = {
    "오류": PatternFill("solid", fgColor="FDECEA"),
    "확인 필요": PatternFill("solid", fgColor="FFF4E0"),
    "안내": PatternFill("solid", fgColor="F0F2F5"),
}
WIDTHS = (10, 10, 34, 24, 56, 16, 16, 16)


def export_issues(output: str, session, reports: dict) -> str:
    """reports: {"목표1": Report, "목표2A": Report, …}"""
    book = Workbook()
    sheet = book.active
    sheet.title = "지적 목록"

    sheet.append([f"예산요구 입력본 검사 결과 · {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    sheet["A1"].font = Font(bold=True, size=13)
    for key, value in session.summary().items():
        sheet.append([key, value])
    sheet.append([])

    head = sheet.max_row + 1
    sheet.append(HEADERS)
    for cell in sheet[head]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDEFF2")
        cell.alignment = Alignment(vertical="center")

    for goal, report in reports.items():
        for issue in getattr(report, "issues", []):
            sheet.append([issue.severity, goal_label(goal), session.label_of(issue.project), issue.item, issue.message,
                          issue.left, issue.right, issue.gap])
            fill = FILLS.get(issue.severity)
            if fill:
                for cell in sheet[sheet.max_row]:
                    cell.fill = fill
            for column in ("F", "G", "H"):
                sheet[f"{column}{sheet.max_row}"].number_format = "#,##0"

    for column, width in zip("ABCDEFGH", WIDTHS):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = f"A{head + 1}"

    # 오류가 아니라서 목록에 없는 것과, 검사를 하지 못해서 목록에 없는 것은 다르다.
    # 뒤쪽을 따로 적어 두지 않으면 '오류 0건'이 '다 맞다'로 읽힌다.
    skipped = [(goal, issue) for goal, report in reports.items() for issue in getattr(report, "skipped", [])]
    if skipped:
        missed = book.create_sheet("검사 못 한 것")
        missed.append(["검사", "사업", "항목", "사유"])
        for cell in missed[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EDEFF2")
        for goal, issue in skipped:
            missed.append([goal_label(goal), session.label_of(issue.project), issue.item, issue.message])
        for column, width in zip("ABCD", (12, 34, 24, 80)):
            missed.column_dimensions[column].width = width

    unsettled = book.create_sheet("비목 미확정")
    unsettled.append(["목코드", "비목명", "항목명", "행 수", "사유"])
    for cell in unsettled[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDEFF2")
    for project, index, item in session.unsettled:
        _, _, source = session.code_for(project, index)
        unsettled.append([item.bimok_code5 or "목코드 없음", item.bimok_name, f"{project.label} · {item.name}", 1, source])
    for column, width in zip("ABCDE", (12, 26, 30, 8, 60)):
        unsettled.column_dimensions[column].width = width
    unsettled.freeze_panes = "A2"

    book.save(output)
    return output
