from __future__ import annotations

import csv
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).parents[1]))

from workflow_core import (  # noqa: E402
    MaskRule, TableData, WorkflowError, apply_clean_plan, apply_masking,
    merge_files, partial_mask, plan_cleaning, read_table, save_clean_workbook, save_table,
    suggest_mask_rules, verify_masking,
)


def make_xlsx(path: Path, headers: list[str], rows: list[list[object]]) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def test_merge_aligns_columns_and_reports_schema_difference(tmp_path: Path) -> None:
    one = make_xlsx(tmp_path / "가학교.xlsx", ["학교명", "성명", "인원"], [["가", "김하늘", 2]])
    two = make_xlsx(tmp_path / "나학교.xlsx", ["인원", "학교명", "연락처"], [[3, "나", "010-1234-5678"]])
    merged = merge_files([one, two])
    assert merged.headers == ["원본 파일명", "학교명", "성명", "인원", "연락처"]
    assert merged.rows == [["가학교.xlsx", "가", "김하늘", 2, None], ["나학교.xlsx", "나", None, 3, "010-1234-5678"]]
    assert "없는 열: 성명" in merged.warnings[0]
    assert "추가 열: 연락처" in merged.warnings[0]


def test_csv_cp949_and_duplicate_headers(tmp_path: Path) -> None:
    path = tmp_path / "자료.csv"
    with path.open("w", encoding="cp949", newline="") as handle:
        csv.writer(handle).writerows([["성명", "성명", ""], ["김민수", "보호자", "값"]])
    assert read_table(path).headers == ["성명", "성명_2", "열3"]


def test_csv_formula_injection_is_escaped_but_negative_numbers_are_kept(tmp_path: Path) -> None:
    path = tmp_path / "외부자료.csv"
    path.write_text("항목,값\n=HYPERLINK(\"https://example.invalid\"),1\n+CMD,2\n-3,3\n", encoding="utf-8-sig")
    table = read_table(path)
    assert table.rows[0][0].startswith("'=")
    assert table.rows[1][0] == "'+CMD"
    assert table.rows[2][0] == "-3"
    output = save_table(table, tmp_path / "safe.xlsx")
    book = load_workbook(output, data_only=False, read_only=True)
    try:
        assert book.active["A2"].data_type == "s"
        assert book.active["A3"].data_type == "s"
    finally:
        book.close()


def test_merge_matches_headers_despite_spacing_and_case(tmp_path: Path) -> None:
    one = make_xlsx(tmp_path / "one.xlsx", ["학교 명", "EMAIL"], [["가교", "a@example.com"]])
    two = make_xlsx(tmp_path / "two.xlsx", ["학교명", "email"], [["나교", "b@example.com"]])
    merged = merge_files([one, two], include_source=False)
    assert merged.headers == ["학교 명", "EMAIL"]
    assert merged.rows == [["가교", "a@example.com"], ["나교", "b@example.com"]]
    assert not merged.warnings


def test_read_table_skips_preamble_and_resolves_ordinal_column(tmp_path: Path) -> None:
    path = tmp_path / "신청자료.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["2026년 신청 현황"])
    sheet.append([])
    sheet.append(["작성기관", "홍은성"])
    sheet.append([])
    sheet.append(["지역", "학교명", "신청금액", "비고", "상태", "대상"])
    sheet.append(["동부", "가학교", 100, "", "", "Y"])
    sheet.append(["서부", "나학교", None, "", "", "Y"])
    book.save(path)
    table = read_table(path)
    assert table.headers == ["지역", "학교명", "신청금액", "비고", "상태", "대상"]
    assert len(table.rows) == 2
    plan = plan_cleaning("지역별 신청금액을 합산하고 3열의 값이 없는 행은 삭제해줘", table)
    assert not plan.warnings
    cleaned = apply_clean_plan(table, plan)
    assert cleaned.rows == [["동부", 100.0]]


@pytest.mark.parametrize("request", [
    "지역별 신청금액 합산",
    "신청금액이 없는 행 삭제",
    "학교명과 성명이 같은 중복 행 제거",
    "신청일 날짜 형식 통일",
    "신청금액 숫자와 콤마 형식 정리",
    "지역을 가나다순으로 정렬",
])
def test_common_natural_language_commands_are_recognized(request: str) -> None:
    table = TableData(
        ["지역", "학교명", "성명", "신청일", "신청금액"],
        [["동부", "가학교", "김민수", "2026.1.2", "1,000"]],
    )
    plan = plan_cleaning(request, table)
    assert plan.actions or plan.warnings
    assert plan.confidence > 0


def test_mask_modes_full_partial_and_drop() -> None:
    original = TableData(["성명", "연락처", "이메일", "주소", "비고"], [["김민수", "010-1234-5678", "teacher@example.com", "서울시 종로구 사직로 1", "유지"]])
    rules = [MaskRule("성명", "partial", "이름"), MaskRule("연락처", "partial", "전화번호"), MaskRule("이메일", "full", "이메일"), MaskRule("주소", "drop", "주소")]
    masked = apply_masking(original, rules)
    verify_masking(original, masked, rules)
    assert masked.headers == ["성명", "연락처", "이메일", "비고"]
    assert masked.rows[0] == ["김*수", "010-****-**78", "*******@*******.***", "유지"]


@pytest.mark.parametrize(("value", "kind", "expected"), [
    ("이순신", "이름", "이*신"), ("01012345678", "전화번호", "010******78"),
    ("2020.03.02", "생년월일", "2020.**.**"), ("a@example.com", "이메일", "a**@example.com"),
    ("서울특별시 강남구 테헤란로 1", "주소", "서울특별시 강남구 ***"),
])
def test_partial_mask_by_kind(value: str, kind: str, expected: str) -> None:
    assert partial_mask(value, kind) == expected


def test_suggest_mask_rules_uses_headers_and_values() -> None:
    table = TableData(["학생 성명", "비상 연락", "전자우편"], [["박하늘", "010-9999-1234", "a@b.kr"]])
    assert [(r.column, r.kind) for r in suggest_mask_rules(table)] == [("학생 성명", "이름"), ("비상 연락", "전화번호"), ("전자우편", "이메일")]


def test_natural_language_multi_action_uses_actual_headers() -> None:
    table = TableData(["지역", "학생 수", "담당자"], [["동부", " 10명 ", "김"], ["서부", "5", "이"], ["동부", "2", "박"]])
    plan = plan_cleaning("학생 수 숫자 형식으로 정리하고, 지역별 학생 수 합계를 내림차순으로 정렬해줘", table)
    assert [action.operation for action in plan.actions] == ["normalize_number", "sum", "sort"]
    result = apply_clean_plan(table, plan)
    assert result.headers == ["지역", "학생 수 합계"]
    assert result.rows == [["동부", 12.0], ["서부", 5.0]]


def test_natural_language_different_wording_for_blank_check() -> None:
    table = TableData(["학교명", "담당자"], [["가교", ""], ["나교", "홍길동"]])
    plan = plan_cleaning("담당자가 누락된 자료를 찾아 표시해 주세요", table)
    result = apply_clean_plan(table, plan)
    assert plan.actions[0].operation == "flag_blanks"
    assert result.rows[0][-1] == "담당자"
    assert result.rows[1][-1] == "정상"


def test_unknown_request_is_not_silently_executed() -> None:
    table = TableData(["학교명", "점수"], [["가교", 10]])
    plan = plan_cleaning("보기 좋게 알아서 완벽히 고쳐줘", table)
    assert plan.confidence == 0
    with pytest.raises(WorkflowError, match="안전하게 실행할 작업"):
        apply_clean_plan(table, plan)


def test_cleaning_catalog_handles_trim_deduplicate_blank_date_and_replace() -> None:
    table = TableData(
        ["학교명", "신청일", "담당자"],
        [[" 가교 ", "2026.3.2", "홍길동"], [" 가교 ", "2026년 3월 2일", ""], ["나교", "2026-04-01", "김하늘"]],
    )
    trim_plan = plan_cleaning("학교명 앞뒤 공백을 정리해줘", table)
    trimmed = apply_clean_plan(table, trim_plan)
    assert trimmed.rows[0][0] == "가교"
    date_plan = plan_cleaning("신청일 날짜 형식을 통일해줘", trimmed)
    dated = apply_clean_plan(trimmed, date_plan)
    assert [row[1] for row in dated.rows] == ["2026-03-02", "2026-03-02", "2026-04-01"]
    replace_plan = plan_cleaning("학교명 열에서 '가교'를 '가학교'로 바꿔줘", dated)
    replaced = apply_clean_plan(dated, replace_plan)
    assert replaced.rows[0][0] == "가학교"
    duplicate_plan = plan_cleaning("학교명과 신청일이 겹치는 중복 행을 제거해줘", replaced)
    deduplicated = apply_clean_plan(replaced, duplicate_plan)
    assert len(deduplicated.rows) == 2
    blank_plan = plan_cleaning("담당자가 빈칸인 행을 삭제해줘", dated)
    assert len(apply_clean_plan(dated, blank_plan).rows) == 2


def test_save_and_reload_verification(tmp_path: Path) -> None:
    table = TableData(["학교", "인원"], [["가", 1], ["나", 2]])
    assert read_table(save_table(table, tmp_path / "result.xlsx")).rows == table.rows


def test_group_sum_is_saved_as_excel_formulas_with_source_sheet(tmp_path: Path) -> None:
    source = TableData(["지역", "학교급", "학생 수"], [["동부", "초", 2], ["동부", "초", 3], ["서부", "중", 4]])
    plan = plan_cleaning("지역과 학교급별 학생 수 합계를 큰 순서로 정렬해줘", source)
    preview = apply_clean_plan(source, plan)
    output = save_clean_workbook(source, plan, preview, tmp_path / "formula.xlsx", include_source=False)
    book = load_workbook(output, data_only=False)
    try:
        assert book.sheetnames == ["정리결과", "원본데이터"]
        assert book["원본데이터"].sheet_state == "hidden"
        formula = book["정리결과"]["C2"].value
        assert formula.startswith("=SUMIFS(")
        assert "'원본데이터'!$C$2:$C$4" in formula
        assert "'원본데이터'!$A$2:$A$4,A2" in formula
        assert "'원본데이터'!$B$2:$B$4,B2" in formula
        assert book.calculation.fullCalcOnLoad
    finally:
        book.close()
