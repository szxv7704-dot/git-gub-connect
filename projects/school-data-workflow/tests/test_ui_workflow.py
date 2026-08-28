from __future__ import annotations

import os
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).parents[1]))

from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox  # noqa: E402
from ui import WorkflowWindow  # noqa: E402


def make_file(path: Path, rows: list[list[object]]) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.append(["지역", "성명", "연락처", "학생 수"])
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def test_ui_runs_the_full_pipeline_and_updates_each_preview(tmp_path: Path, monkeypatch) -> None:
    app = QApplication.instance() or QApplication([])
    window = WorkflowWindow()
    one = make_file(tmp_path / "가교.xlsx", [["동부", "김민수", "010-1111-2222", 2]])
    two = make_file(tmp_path / "나교.xlsx", [["서부", "이하늘", "010-3333-4444", 3], ["동부", "박바다", "010-5555-6666", 4]])
    window._add_paths([str(one), str(two)])

    window.next()
    assert window.stage == 1
    assert len(window.results[0].rows) == 3
    assert window.preview_table.rowCount() == 3
    assert window.mask_table.rowCount() == len(window.results[0].headers)

    window.next()
    assert window.stage == 2
    assert window.results[1].rows[0][2] == "김*수"
    assert window.results[1].rows[0][3] == "010-****-**22"
    assert "처리 전: 김민수" in window.preview_table.item(0, 2).toolTip()

    window.request.setPlainText("지역별 학생 수 합계를 큰 순서로 정렬해줘")
    assert window.clean_plan is not None
    assert [action.operation for action in window.clean_plan.actions] == ["sum", "sort"]
    window.next()
    assert window.stage == 3
    assert window.results[2].headers == ["지역", "학생 수 합계"]
    assert window.results[2].rows == [["동부", 6.0], ["서부", 3.0]]
    assert window.download_button.isEnabled()
    output = tmp_path / "최종결과.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(output), "Excel"))
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: QMessageBox.StandardButton.Ok)
    window.download_stage(final=True)
    book = load_workbook(output, data_only=False, read_only=True)
    try:
        assert str(book["정리결과"]["B2"].value).startswith("=SUMIFS(")
    finally:
        book.close()
    window.close()
    app.processEvents()
