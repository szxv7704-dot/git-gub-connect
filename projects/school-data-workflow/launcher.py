from __future__ import annotations

import sys
import os


def _prepare_bundled_dlls() -> None:
    """Ensure bundled Qt DLLs win over DLLs installed elsewhere on Windows."""
    if sys.platform != "win32":
        return
    bundle_dir = getattr(sys, "_MEIPASS", None)
    if not bundle_dir:
        return
    # PyInstaller normally configures this, but an installed Qt/PySide6 can
    # otherwise satisfy QtWidgets with a mismatched QtCore on some PCs.
    try:
        os.add_dll_directory(bundle_dir)
    except (AttributeError, OSError):
        pass
    os.environ["PATH"] = bundle_dir + os.pathsep + os.environ.get("PATH", "")


_prepare_bundled_dlls()


if "--self-test" in sys.argv:
    # 패키지 자체 점검은 Qt를 불러오기 전에 실행해 GUI 초기화와 분리한다.
    from workflow_core import MaskRule, apply_clean_plan, apply_masking, merge_files, plan_cleaning, save_clean_workbook, verify_masking
    from openpyxl import Workbook, load_workbook
    from pathlib import Path
    from tempfile import TemporaryDirectory

    with TemporaryDirectory(prefix="school-data-workflow-") as folder:
        root = Path(folder)
        inputs = []
        for name, rows in (
            ("가교.xlsx", [["동부", "김민수", "010-1111-2222", 2]]),
            ("나교.xlsx", [["서부", "이하늘", "010-3333-4444", 3], ["동부", "박바다", "010-5555-6666", 4]]),
        ):
            path = root / name
            book = Workbook()
            sheet = book.active
            sheet.append(["지역", "성명", "연락처", "학생 수"])
            for row in rows:
                sheet.append(row)
            book.save(path)
            inputs.append(path)
        merged = merge_files(inputs)
        rules = [MaskRule("성명", "partial", "이름"), MaskRule("연락처", "drop", "전화번호")]
        masked = apply_masking(merged, rules)
        verify_masking(merged, masked, rules)
        plan = plan_cleaning("지역별 학생 수 합계를 큰 순서로 정렬해줘", masked)
        cleaned = apply_clean_plan(masked, plan)
        output = save_clean_workbook(masked, plan, cleaned, root / "결과.xlsx")
        check = load_workbook(output, data_only=False, read_only=True)
        try:
            formula = check["정리결과"]["B2"].value
            if not isinstance(formula, str) or not formula.startswith("=SUMIFS("):
                raise RuntimeError("수식 자체 점검 실패")
        finally:
            check.close()
    raise SystemExit(0)


from PySide6.QtWidgets import QApplication
from ui import WorkflowWindow

app = QApplication(sys.argv)
app.setStyle("Fusion")
window = WorkflowWindow()
window.show()
raise SystemExit(app.exec())
