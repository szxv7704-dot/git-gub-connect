from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from PySide6.QtCore import Qt, Signal
from PySide6.QtGui import QColor
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QButtonGroup,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QRadioButton,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from workflow_core import (
    MaskRule,
    TableData,
    WorkflowError,
    apply_clean_plan,
    apply_masking,
    merge_files,
    plan_cleaning,
    read_table,
    save_clean_workbook,
    save_table,
    suggest_mask_rules,
    verify_masking,
)


STAGES = (
    ("1", "파일 취합", "여러 파일을 하나로 합칩니다"),
    ("2", "개인정보 보호", "필요한 경우 비식별화합니다"),
    ("3", "데이터 정리", "수식과 규칙을 적용합니다"),
    ("4", "완료", "검토하고 결과를 저장합니다"),
)


class StepCard(QFrame):
    clicked = Signal(int)

    def __init__(self, index: int, number: str, title: str, description: str) -> None:
        super().__init__()
        self.index = index
        self.setObjectName("stepCard")
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        row = QHBoxLayout(self)
        row.setContentsMargins(12, 10, 12, 10)
        row.setSpacing(9)
        self.number = QLabel(number)
        self.number.setObjectName("stepNumber")
        self.number.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.number.setFixedSize(28, 28)
        text = QVBoxLayout()
        text.setSpacing(1)
        self.title = QLabel(title)
        self.title.setObjectName("stepTitle")
        self.description = QLabel(description)
        self.description.setObjectName("stepDescription")
        text.addWidget(self.title)
        text.addWidget(self.description)
        row.addWidget(self.number)
        row.addLayout(text, 1)

    def mousePressEvent(self, event) -> None:
        if event.button() == Qt.MouseButton.LeftButton:
            self.clicked.emit(self.index)
        super().mousePressEvent(event)

    def set_state(self, state: str) -> None:
        self.setProperty("state", state)
        self.number.setText("✓" if state == "done" else str(self.index + 1))
        self.style().unpolish(self)
        self.style().polish(self)


class StagePanel(QFrame):
    def __init__(self, eyebrow: str, title: str, description: str) -> None:
        super().__init__()
        self.setObjectName("stagePanel")
        self.layout = QVBoxLayout(self)
        self.layout.setContentsMargins(22, 20, 22, 20)
        self.layout.setSpacing(12)
        label = QLabel(eyebrow)
        label.setObjectName("eyebrow")
        heading = QLabel(title)
        heading.setObjectName("stageTitle")
        copy = QLabel(description)
        copy.setObjectName("stageCopy")
        copy.setWordWrap(True)
        self.layout.addWidget(label)
        self.layout.addWidget(heading)
        self.layout.addWidget(copy)


class WorkflowWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("학교 자료 취합 도우미 · UI 시안")
        self.resize(1420, 880)
        self.setMinimumSize(1120, 720)
        self.stage = 0
        self.completed = set()
        self.files: list[Path] = []
        self.privacy_enabled = True
        self.stage_outputs = [False, False, False, False]
        self.results: dict[int, TableData] = {}
        self.clean_plan = None
        self._build()
        self.go_to(0)

    def _build(self) -> None:
        self.setStyleSheet(STYLES)
        page = QWidget()
        page.setObjectName("page")
        self.setCentralWidget(page)
        root = QVBoxLayout(page)
        root.setContentsMargins(22, 18, 22, 18)
        root.setSpacing(14)

        root.addWidget(self._header())
        self.step_cards: list[StepCard] = []
        steps = QHBoxLayout()
        steps.setSpacing(8)
        for index, data in enumerate(STAGES):
            card = StepCard(index, *data)
            card.clicked.connect(self.try_go_to)
            steps.addWidget(card, 1)
            self.step_cards.append(card)
        root.addLayout(steps)

        body = QHBoxLayout()
        body.setSpacing(14)
        self.stack = QStackedWidget()
        self.stack.addWidget(self._merge_stage())
        self.stack.addWidget(self._privacy_stage())
        self.stack.addWidget(self._clean_stage())
        self.stack.addWidget(self._finish_stage())
        body.addWidget(self.stack, 42)
        body.addWidget(self._preview_panel(), 58)
        root.addLayout(body, 1)
        root.addWidget(self._footer())

    def _header(self) -> QFrame:
        frame = QFrame()
        frame.setObjectName("header")
        row = QHBoxLayout(frame)
        row.setContentsMargins(20, 15, 20, 15)
        text = QVBoxLayout()
        title = QLabel("학교 자료 취합 도우미")
        title.setObjectName("appTitle")
        subtitle = QLabel("취합부터 개인정보 보호, 데이터 정리까지 한 흐름으로 진행합니다.")
        subtitle.setObjectName("appSubtitle")
        text.addWidget(title)
        text.addWidget(subtitle)
        row.addLayout(text)
        row.addStretch()
        badge = QLabel("● 오프라인 처리 · 원본 보존")
        badge.setObjectName("safeBadge")
        row.addWidget(badge)
        return frame

    def _merge_stage(self) -> QWidget:
        panel = StagePanel("1단계 · 필수", "받은 파일을 한곳에 모아주세요", "학교별 엑셀 또는 CSV 파일을 여러 개 선택합니다. 헤더가 같은 표는 아래 방향으로 이어 붙입니다.")
        toolbar = QHBoxLayout()
        add = QPushButton("＋ 파일 추가")
        add.setProperty("primary", True)
        add.clicked.connect(self.add_files)
        folder = QPushButton("＋ 폴더에서 추가")
        folder.clicked.connect(self.add_folder)
        remove = QPushButton("선택 삭제")
        remove.clicked.connect(self.remove_files)
        toolbar.addWidget(add)
        toolbar.addWidget(folder)
        toolbar.addStretch()
        toolbar.addWidget(remove)
        panel.layout.addLayout(toolbar)
        self.file_list = QListWidget()
        self.file_list.setAlternatingRowColors(True)
        self.file_list.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.file_list.setMinimumHeight(220)
        panel.layout.addWidget(self.file_list, 1)
        self.file_summary = QLabel("아직 선택된 파일이 없습니다.")
        self.file_summary.setObjectName("infoLine")
        panel.layout.addWidget(self.file_summary)
        self.source_column = QCheckBox("취합 결과에 ‘원본 파일명’ 열 추가")
        self.source_column.setChecked(True)
        self.keep_sheets = QCheckBox("시트별로 나누지 않고 첫 번째 표를 한 시트에 취합")
        self.keep_sheets.setChecked(True)
        panel.layout.addWidget(self.source_column)
        panel.layout.addWidget(self.keep_sheets)
        note = QLabel("서로 다른 열이 발견되면 자동으로 합치지 않고 다음 화면에서 차이를 알려줍니다.")
        note.setObjectName("note")
        note.setWordWrap(True)
        panel.layout.addWidget(note)
        return panel

    def _privacy_stage(self) -> QWidget:
        panel = StagePanel("2단계 · 선택", "개인정보를 보호할까요?", "취합된 파일에서 이름, 연락처, 생년월일 등 개인정보 후보를 찾아 검토합니다. 필요 없다면 이 단계를 건너뛸 수 있습니다.")
        self.privacy_on = QRadioButton("비식별화를 진행합니다")
        self.privacy_off = QRadioButton("비식별화 없이 다음 단계로 갑니다")
        self.privacy_on.setChecked(True)
        group = QButtonGroup(self)
        group.addButton(self.privacy_on)
        group.addButton(self.privacy_off)
        group.buttonClicked.connect(self.privacy_choice_changed)
        panel.layout.addWidget(self._choice_card(self.privacy_on, "개인정보 후보를 자동 탐지한 뒤 사용자가 적용 여부를 확인합니다.", "권장"))
        panel.layout.addWidget(self._choice_card(self.privacy_off, "개인정보가 없는 공개 자료이거나 원본 값을 유지해야 할 때 선택합니다.", "선택"))
        self.privacy_options = QFrame()
        self.privacy_options.setObjectName("subPanel")
        opts = QVBoxLayout(self.privacy_options)
        opts.addWidget(QLabel("열별 처리 방식 — 자동 제안을 반드시 확인하세요"))
        self.mask_table = QTableWidget(0, 3)
        self.mask_table.setHorizontalHeaderLabels(["열", "탐지 유형", "처리 방식"])
        self.mask_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.mask_table.setMinimumHeight(230)
        opts.addWidget(self.mask_table)
        opts.addWidget(QLabel("완전 마스킹은 전체 값을 가리고, 부분 마스킹은 식별에 필요한 부분만 가립니다. 열 삭제는 결과에서 열 자체를 제거합니다."))
        panel.layout.addWidget(self.privacy_options)
        panel.layout.addStretch()
        return panel

    @staticmethod
    def _choice_card(radio: QRadioButton, description: str, tag: str) -> QFrame:
        card = QFrame()
        card.setObjectName("choiceCard")
        row = QHBoxLayout(card)
        row.setContentsMargins(14, 12, 14, 12)
        text = QVBoxLayout()
        text.addWidget(radio)
        copy = QLabel(description)
        copy.setObjectName("muted")
        copy.setWordWrap(True)
        text.addWidget(copy)
        row.addLayout(text, 1)
        badge = QLabel(tag)
        badge.setObjectName("smallBadge")
        row.addWidget(badge)
        return card

    def _clean_stage(self) -> QWidget:
        panel = StagePanel("3단계 · 선택", "데이터를 어떻게 정리할까요?", "원하는 결과를 평소 말하듯 입력하세요. 적용할 수식과 정리 규칙을 먼저 보여드린 뒤 실행합니다.")
        label = QLabel("정리 요청")
        label.setObjectName("fieldLabel")
        panel.layout.addWidget(label)
        self.request = QTextEdit()
        self.request.setPlaceholderText("예: 지역별로 학생 수를 합산하고, 합계가 큰 순서대로 정렬해줘")
        self.request.setMinimumHeight(120)
        self.request.textChanged.connect(self.analyze_request)
        panel.layout.addWidget(self.request)
        examples = QLabel("빠른 예시")
        examples.setObjectName("fieldLabel")
        panel.layout.addWidget(examples)
        chips = QHBoxLayout()
        for text, prompt in (
            ("지역별 합계", "지역별로 인원수를 합산하고 큰 순서대로 정렬해줘"),
            ("중복 확인", "학교명과 성명이 같은 중복 행을 찾아줘"),
            ("빈칸 점검", "필수 열에 빈칸이 있는 행을 표시해줘"),
        ):
            button = QPushButton(text)
            button.setProperty("chip", True)
            button.clicked.connect(lambda _checked=False, value=prompt: self.request.setPlainText(value))
            chips.addWidget(button)
        chips.addStretch()
        panel.layout.addLayout(chips)
        self.formula_option = QCheckBox("결과에 계산식이 보이도록 엑셀 수식으로 작성")
        self.formula_option.setChecked(True)
        self.summary_option = QCheckBox("원본 데이터 시트와 요약 결과 시트를 함께 저장")
        self.summary_option.setChecked(True)
        panel.layout.addWidget(self.formula_option)
        panel.layout.addWidget(self.summary_option)
        analysis = QFrame()
        analysis.setObjectName("subPanel")
        al = QVBoxLayout(analysis)
        al.addWidget(QLabel("적용 예정 작업"))
        self.plan_label = QLabel("요청을 입력하면 적용할 열, 계산식, 정렬 방식이 여기에 표시됩니다.")
        self.plan_label.setObjectName("muted")
        self.plan_label.setWordWrap(True)
        al.addWidget(self.plan_label)
        panel.layout.addWidget(analysis)
        panel.layout.addStretch()
        return panel

    def _finish_stage(self) -> QWidget:
        panel = StagePanel("4단계 · 완료", "마지막으로 결과를 확인하세요", "각 단계의 처리 여부와 행 수를 확인한 뒤 최종 파일을 저장합니다. 이전 단계 결과도 오른쪽에서 따로 저장할 수 있습니다.")
        self.finish_items = QVBoxLayout()
        panel.layout.addLayout(self.finish_items)
        panel.layout.addStretch()
        warning = QLabel("최종 저장 전 미리보기에서 지역별 합계와 개인정보 잔존 여부를 확인하세요.")
        warning.setObjectName("note")
        warning.setWordWrap(True)
        panel.layout.addWidget(warning)
        return panel

    def _preview_panel(self) -> QFrame:
        frame = QFrame()
        frame.setObjectName("previewPanel")
        layout = QVBoxLayout(frame)
        layout.setContentsMargins(18, 17, 18, 15)
        header = QHBoxLayout()
        text = QVBoxLayout()
        title = QLabel("현재 단계 미리보기")
        title.setObjectName("previewTitle")
        self.preview_caption = QLabel("파일을 추가하면 첫 번째 표의 일부가 표시됩니다.")
        self.preview_caption.setObjectName("muted")
        text.addWidget(title)
        text.addWidget(self.preview_caption)
        header.addLayout(text)
        header.addStretch()
        self.rows_badge = QLabel("0행")
        self.rows_badge.setObjectName("smallBadge")
        header.addWidget(self.rows_badge)
        layout.addLayout(header)
        self.preview_table = QTableWidget()
        self.preview_table.setAlternatingRowColors(True)
        self.preview_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.preview_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.preview_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.preview_table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.preview_table, 1)
        self.preview_empty = QLabel("엑셀·CSV 파일을 추가하면\n여기에서 단계별 결과를 확인할 수 있습니다.")
        self.preview_empty.setObjectName("emptyState")
        self.preview_empty.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview_empty.setMinimumHeight(100)
        layout.addWidget(self.preview_empty)
        actions = QHBoxLayout()
        self.preview_status = QLabel("아직 생성된 결과가 없습니다.")
        self.preview_status.setObjectName("muted")
        actions.addWidget(self.preview_status)
        actions.addStretch()
        self.download_button = QPushButton("이 단계 결과 저장")
        self.download_button.setProperty("download", True)
        self.download_button.setEnabled(False)
        self.download_button.clicked.connect(self.download_stage)
        actions.addWidget(self.download_button)
        layout.addLayout(actions)
        return frame

    def _footer(self) -> QFrame:
        frame = QFrame()
        frame.setObjectName("footer")
        row = QHBoxLayout(frame)
        row.setContentsMargins(14, 9, 14, 9)
        self.status = QLabel("1단계에서 파일을 추가해 시작하세요.")
        self.status.setObjectName("muted")
        self.back_button = QPushButton("← 이전")
        self.back_button.clicked.connect(self.back)
        self.next_button = QPushButton("파일 취합하고 다음 →")
        self.next_button.setProperty("primary", True)
        self.next_button.clicked.connect(self.next)
        row.addWidget(self.status, 1)
        row.addWidget(self.back_button)
        row.addWidget(self.next_button)
        return frame

    def add_files(self) -> None:
        paths, _ = QFileDialog.getOpenFileNames(self, "학교 자료 추가", "", "엑셀·CSV (*.xlsx *.xlsm *.csv)")
        self._add_paths(paths)

    def add_folder(self) -> None:
        folder = QFileDialog.getExistingDirectory(self, "학교 자료가 있는 폴더 선택")
        if not folder:
            return
        paths = [str(path) for path in Path(folder).iterdir() if path.suffix.lower() in {".xlsx", ".xlsm", ".csv"}]
        self._add_paths(paths)

    def _add_paths(self, paths: list[str]) -> None:
        known = {path.resolve() for path in self.files}
        for raw in paths:
            path = Path(raw).resolve()
            if path not in known:
                self.files.append(path)
                known.add(path)
        self.files.sort(key=lambda path: path.name.lower())
        self.file_list.clear()
        for path in self.files:
            item = QListWidgetItem(f"  {path.name}")
            item.setToolTip(str(path))
            self.file_list.addItem(item)
        total = sum(path.stat().st_size for path in self.files if path.exists())
        self.file_summary.setText(f"{len(self.files)}개 파일 · {total / 1024 / 1024:.1f} MB") if self.files else self.file_summary.setText("아직 선택된 파일이 없습니다.")
        self.status.setText(f"{len(self.files)}개 파일을 선택했습니다. 미리보기에서 열 구성을 확인하세요.")
        self._load_preview(self.files[0] if self.files else None)
        self._update_navigation()

    def remove_files(self) -> None:
        rows = sorted({index.row() for index in self.file_list.selectedIndexes()}, reverse=True)
        for row in rows:
            self.files.pop(row)
        self._add_paths([])

    def _load_preview(self, path: Path | None) -> None:
        if path is None:
            self.preview_table.clear()
            self.preview_table.setRowCount(0)
            self.preview_table.setColumnCount(0)
            self.preview_table.hide()
            self.preview_empty.show()
            self.rows_badge.setText("0행")
            return
        try:
            table = read_table(path)
        except Exception as exc:
            self.preview_table.hide()
            self.preview_empty.setText(f"미리보기를 열 수 없습니다.\n{exc}")
            self.preview_empty.show()
            return
        self._show_table(table, f"{path.name} · 원본 상위 행")

    def _show_table(self, table: TableData, caption: str | None = None, baseline: TableData | None = None) -> None:
        self.preview_table.clear()
        rows = table.rows[:100]
        headers = table.headers
        self.preview_empty.hide()
        self.preview_table.show()
        self.preview_table.setColumnCount(len(headers))
        self.preview_table.setHorizontalHeaderLabels(headers)
        self.preview_table.setRowCount(len(rows))
        for row_index, row in enumerate(rows):
            for column, value in enumerate(row):
                item = QTableWidgetItem("" if value is None else str(value))
                if row_index == 0:
                    item.setBackground(QColor("#F8FBFF"))
                if baseline and headers[column] in baseline.headers and row_index < len(baseline.rows):
                    before_column = baseline.headers.index(headers[column])
                    before = baseline.rows[row_index][before_column]
                    if str(before or "") != str(value or ""):
                        item.setBackground(QColor("#FFF1B8"))
                        item.setToolTip(f"처리 전: {before}\n처리 후: {value}")
                self.preview_table.setItem(row_index, column, item)
        self.rows_badge.setText(f"{len(table.rows):,}행")
        self.preview_caption.setText(caption or f"{table.stage} · 상위 {len(rows)}행")
        if table.warnings:
            self.preview_status.setText(f"주의 {len(table.warnings)}건: {table.warnings[0]}")
        else:
            self.preview_status.setText(f"{len(table.rows):,}행 · {len(table.headers):,}열을 확인했습니다.")

    def _populate_mask_rules(self, table: TableData) -> None:
        suggestions = {rule.column: rule for rule in suggest_mask_rules(table)}
        self.mask_table.setRowCount(len(table.headers))
        for row, header in enumerate(table.headers):
            suggestion = suggestions.get(header)
            self.mask_table.setItem(row, 0, QTableWidgetItem(header))
            self.mask_table.setItem(row, 1, QTableWidgetItem(suggestion.kind if suggestion else "해당 없음"))
            combo = QComboBox()
            combo.addItems(["처리 안 함", "부분 마스킹", "완전 마스킹", "열 삭제"])
            combo.setCurrentText("부분 마스킹" if suggestion else "처리 안 함")
            self.mask_table.setCellWidget(row, 2, combo)

    def _mask_rules(self) -> list[MaskRule]:
        modes = {"부분 마스킹": "partial", "완전 마스킹": "full", "열 삭제": "drop"}
        rules = []
        for row in range(self.mask_table.rowCount()):
            header = self.mask_table.item(row, 0).text()
            kind = self.mask_table.item(row, 1).text()
            mode_text = self.mask_table.cellWidget(row, 2).currentText()
            if mode_text in modes:
                rules.append(MaskRule(header, modes[mode_text], kind))
        return rules

    def analyze_request(self) -> None:
        source = self.results.get(1) or self.results.get(0)
        if not source:
            self.clean_plan = None
            return
        self.clean_plan = plan_cleaning(self.request.toPlainText(), source)
        details = self.clean_plan.explanation
        if self.clean_plan.warnings:
            details += "\n확인 필요: " + " ".join(self.clean_plan.warnings)
        details += f"\n해석 신뢰도 {self.clean_plan.confidence:.0%} · 실행 전 미리보기로 검증합니다."
        self.plan_label.setText(details)

    def privacy_choice_changed(self) -> None:
        self.privacy_enabled = self.privacy_on.isChecked()
        self.privacy_options.setVisible(self.privacy_enabled)
        self._update_navigation()

    def try_go_to(self, index: int) -> None:
        if index <= self.stage or index in self.completed or index - 1 in self.completed:
            self.go_to(index)

    def go_to(self, index: int) -> None:
        self.stage = max(0, min(index, len(STAGES) - 1))
        self.stack.setCurrentIndex(self.stage)
        if self.stage == 3:
            self._refresh_finish_summary()
        for i, card in enumerate(self.step_cards):
            card.set_state("current" if i == self.stage else "done" if i in self.completed else "pending")
        captions = (
            "선택한 원본 파일의 상위 행입니다.",
            "취합 결과에서 개인정보 처리 전후를 비교합니다.",
            "수식과 정리 규칙을 적용한 결과를 보여줍니다.",
            "최종 결과를 저장하기 전에 마지막으로 확인합니다.",
        )
        self.preview_caption.setText(captions[self.stage])
        if self.stage in self.results:
            baseline = self.results.get(0) if self.stage == 1 else None
            self._show_table(self.results[self.stage], baseline=baseline)
        elif self.stage in {1, 2} and (self.stage - 1) in self.results:
            source = self.results[self.stage - 1]
            baseline = self.results.get(0) if self.stage == 2 and source is self.results.get(1) else None
            self._show_table(source, f"{source.stage} 결과 · 다음 작업 전", baseline)
        elif self.stage == 3 and self.results:
            self._show_table(self.results[max(self.results)])
        else:
            self.preview_status.setText("이 단계의 처리가 끝나면 결과 파일을 저장할 수 있습니다.")
        self.download_button.setEnabled(self.stage_outputs[self.stage] or self.stage in self.results)
        self._update_navigation()

    def _update_navigation(self) -> None:
        self.back_button.setEnabled(self.stage > 0)
        labels = (
            "파일 취합하고 다음 →",
            "비식별화 적용하고 다음 →" if self.privacy_enabled else "이 단계 건너뛰기 →",
            "데이터 정리하고 결과 보기 →",
            "최종 엑셀 저장",
        )
        self.next_button.setText(labels[self.stage])
        self.next_button.setEnabled(bool(self.files) if self.stage == 0 else True)

    def next(self) -> None:
        if self.stage == 0 and not self.files:
            QMessageBox.information(self, "파일 필요", "먼저 취합할 엑셀 또는 CSV 파일을 추가해 주세요.")
            return
        if self.stage == 2 and not self.request.toPlainText().strip():
            answer = QMessageBox.question(self, "정리 요청 없음", "데이터 정리 없이 취합 결과만 최종 파일로 저장할까요?")
            if answer != QMessageBox.StandardButton.Yes:
                return
        try:
            if self.stage == 0:
                result = merge_files(self.files, include_source=self.source_column.isChecked())
                self.results[0] = result
                self._populate_mask_rules(result)
            elif self.stage == 1:
                source = self.results[0]
                if self.privacy_enabled:
                    rules = self._mask_rules()
                    if not rules:
                        answer = QMessageBox.question(self, "처리 규칙 없음", "선택한 개인정보 처리 규칙이 없습니다. 이 단계를 건너뛸까요?")
                        if answer != QMessageBox.StandardButton.Yes:
                            return
                    result = apply_masking(source, rules)
                    verify_masking(source, result, rules)
                else:
                    result = source
                self.results[1] = result
            elif self.stage == 2:
                source = self.results.get(1, self.results[0])
                self.clean_plan = plan_cleaning(self.request.toPlainText(), source)
                if self.clean_plan.warnings:
                    raise WorkflowError(" ".join(self.clean_plan.warnings))
                result = apply_clean_plan(source, self.clean_plan)
                self.results[2] = result
                self.results[3] = result
            else:
                self.download_stage(final=True)
                return
        except Exception as exc:
            QMessageBox.critical(self, "처리할 수 없음", str(exc))
            return
        self.stage_outputs[self.stage] = True
        self.completed.add(self.stage)
        self._show_table(self.results[self.stage], f"{self.results[self.stage].stage} 결과 · 상위 100행")
        self.download_button.setEnabled(True)
        self.go_to(self.stage + 1)

    def back(self) -> None:
        self.go_to(self.stage - 1)

    def download_stage(self, final: bool = False) -> None:
        default = "학교자료_최종결과.xlsx" if final or self.stage == 3 else f"학교자료_{self.stage + 1}단계_결과.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "결과 저장 위치", default, "Excel 통합문서 (*.xlsx)")
        if not path:
            return
        table = self.results.get(2 if final or self.stage == 3 else self.stage)
        if table is None:
            QMessageBox.warning(self, "결과 없음", "먼저 이 단계의 처리를 완료해 주세요.")
            return
        try:
            if (final or self.stage == 3) and self.clean_plan is not None and 2 in self.results:
                source = self.results.get(1, self.results[0])
                save_clean_workbook(
                    source,
                    self.clean_plan,
                    self.results[2],
                    Path(path),
                    use_formulas=self.formula_option.isChecked(),
                    include_source=self.summary_option.isChecked(),
                )
            else:
                save_table(table, Path(path), table.stage)
        except Exception as exc:
            QMessageBox.critical(self, "저장 실패", str(exc))
            return
        QMessageBox.information(self, "저장 완료", f"결과를 저장하고 다시 열어 행·열 수를 검증했습니다.\n{path}")

    def _refresh_finish_summary(self) -> None:
        while self.finish_items.count():
            item = self.finish_items.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        states = (
            ("파일 취합", f"{len(self.files)}개 파일", True),
            ("개인정보 보호", "비식별화 적용" if self.privacy_enabled else "건너뜀", self.privacy_enabled),
            ("데이터 정리", self.request.toPlainText().strip() or "정리하지 않음", bool(self.request.toPlainText().strip())),
        )
        for title, detail, active in states:
            card = QFrame()
            card.setObjectName("summaryCard")
            row = QHBoxLayout(card)
            mark = QLabel("✓" if active else "—")
            mark.setObjectName("summaryMark")
            texts = QVBoxLayout()
            texts.addWidget(QLabel(title))
            copy = QLabel(detail)
            copy.setObjectName("muted")
            copy.setWordWrap(True)
            texts.addWidget(copy)
            row.addWidget(mark)
            row.addLayout(texts, 1)
            self.finish_items.addWidget(card)


STYLES = """
QMainWindow, QWidget#page { background:#F2F5F8; color:#172235; font-family:'Malgun Gothic'; font-size:13px; }
QFrame#header { background:#173F5F; border-radius:14px; }
QLabel#appTitle { color:white; font-size:24px; font-weight:700; }
QLabel#appSubtitle { color:#D9E7F2; }
QLabel#safeBadge { color:#14532D; background:#DCFCE7; border-radius:12px; padding:6px 11px; font-weight:700; }
QFrame#stepCard { background:#FFFFFF; border:1px solid #D9E2EA; border-radius:11px; }
QFrame#stepCard[state='current'] { background:#EFF7FF; border:2px solid #2B7BBB; }
QFrame#stepCard[state='done'] { background:#F0FDF4; border:1px solid #86C89B; }
QLabel#stepNumber { background:#E8EEF4; color:#536579; border-radius:14px; font-weight:700; }
QFrame#stepCard[state='current'] QLabel#stepNumber { background:#1769AA; color:white; }
QFrame#stepCard[state='done'] QLabel#stepNumber { background:#16834B; color:white; }
QLabel#stepTitle { font-weight:700; color:#24364B; }
QLabel#stepDescription { color:#718096; font-size:11px; }
QFrame#stagePanel, QFrame#previewPanel { background:white; border:1px solid #D9E2EA; border-radius:14px; }
QLabel#eyebrow { color:#1769AA; font-size:11px; font-weight:700; }
QLabel#stageTitle { color:#18324A; font-size:22px; font-weight:700; }
QLabel#stageCopy, QLabel#muted { color:#68798B; }
QLabel#previewTitle { color:#18324A; font-size:17px; font-weight:700; }
QLabel#fieldLabel { font-weight:700; color:#334155; }
QLabel#infoLine { background:#F5F8FB; border-radius:7px; color:#456078; padding:8px; }
QLabel#note { background:#FFF7E5; color:#805B13; border:1px solid #F0D28A; border-radius:8px; padding:9px; }
QLabel#smallBadge { background:#EAF2F8; color:#2A5B7E; border-radius:10px; padding:4px 9px; font-weight:700; }
QLabel#emptyState { color:#8291A2; background:#F8FAFC; border:1px dashed #CBD5E1; border-radius:10px; font-size:15px; }
QLabel#summaryMark { color:#15803D; font-size:18px; font-weight:700; }
QFrame#choiceCard, QFrame#summaryCard { background:#FAFCFE; border:1px solid #DCE5ED; border-radius:10px; }
QFrame#subPanel { background:#F5F8FB; border:1px solid #DFE7EE; border-radius:10px; }
QFrame#footer { background:white; border:1px solid #D9E2EA; border-radius:12px; }
QPushButton { min-height:36px; padding:0 14px; background:white; color:#2D4257; border:1px solid #C7D3DE; border-radius:8px; font-weight:600; }
QPushButton:hover { background:#F1F7FC; border-color:#7FAACB; }
QPushButton[primary='true'] { background:#1769AA; color:white; border:0; min-height:40px; }
QPushButton[primary='true']:hover { background:#10598F; }
QPushButton[primary='true']:disabled { background:#AAB8C4; }
QPushButton[download='true'] { background:#0F766E; color:white; border:0; }
QPushButton[download='true']:disabled { background:#A9B8B5; }
QPushButton[chip='true'] { min-height:30px; border-radius:15px; padding:0 12px; background:#F4F8FB; }
QListWidget, QTableWidget, QTextEdit { background:white; border:1px solid #CED9E3; border-radius:9px; selection-background-color:#DCEEFF; selection-color:#172235; }
QListWidget { padding:6px; }
QTextEdit { padding:8px; }
QHeaderView::section { background:#EEF3F7; color:#3A4C5D; border:0; border-bottom:1px solid #CBD6DF; padding:8px; font-weight:700; }
QTableWidget { gridline-color:#E6ECF1; alternate-background-color:#F8FAFC; }
QCheckBox, QRadioButton { spacing:8px; }
"""


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        from openpyxl import Workbook, load_workbook

        with TemporaryDirectory(prefix="school-data-workflow-") as folder:
            root = Path(folder)
            inputs = []
            for name, rows in (
                ("가교.xlsx", [["동부", "김민수", "010-1111-2222", 2]]),
                ("나교.xlsx", [["서부", "이하늘", "010-3333-4444", 3], ["동부", "박바다", "010-5555-6666", 4]]),
            ):
                path = root / name
                book = Workbook()
                sheet = book.active
                sheet.append(["지역", "성명", "연락처", "학생 수"])
                for row in rows:
                    sheet.append(row)
                book.save(path)
                inputs.append(path)
            merged = merge_files(inputs)
            rules = [MaskRule("성명", "partial", "이름"), MaskRule("연락처", "drop", "전화번호")]
            masked = apply_masking(merged, rules)
            verify_masking(merged, masked, rules)
            clean_plan = plan_cleaning("지역별 학생 수 합계를 큰 순서로 정렬해줘", masked)
            cleaned = apply_clean_plan(masked, clean_plan)
            output = save_clean_workbook(masked, clean_plan, cleaned, root / "결과.xlsx")
            check = load_workbook(output, data_only=False, read_only=True)
            try:
                formula = check["정리결과"]["B2"].value
                if not isinstance(formula, str) or not formula.startswith("=SUMIFS("):
                    raise RuntimeError("수식 자체 점검 실패")
            finally:
                check.close()
        raise SystemExit(0)
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    window = WorkflowWindow()
    window.show()
    sys.exit(app.exec())
