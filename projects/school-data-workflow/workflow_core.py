from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


class WorkflowError(ValueError):
    """사용자 입력 또는 표 구조 때문에 작업을 안전하게 수행할 수 없을 때 발생한다."""


MAX_INPUT_BYTES = 250 * 1024 * 1024
MAX_ROWS = 1_000_000
MAX_COLUMNS = 500


@dataclass
class TableData:
    headers: list[str]
    rows: list[list[Any]]
    stage: str = "원본"
    warnings: list[str] = field(default_factory=list)
    source_files: list[Path] = field(default_factory=list)

    def normalized(self) -> "TableData":
        width = len(self.headers)
        return TableData(
            list(self.headers),
            [(list(row) + [None] * width)[:width] for row in self.rows],
            self.stage,
            list(self.warnings),
            list(self.source_files),
        )


@dataclass(frozen=True)
class MaskRule:
    column: str
    mode: str  # full, partial, drop
    kind: str = "자동"


@dataclass(frozen=True)
class CleanAction:
    operation: str
    columns: tuple[str, ...] = ()
    value: str | None = None
    descending: bool = False


@dataclass
class CleanPlan:
    request: str
    actions: list[CleanAction]
    explanation: str
    confidence: float
    warnings: list[str] = field(default_factory=list)


HEADER_KINDS = {
    "이름": ("이름", "성명", "학생명", "교직원명", "담당자", "신청자"),
    "전화번호": ("전화", "연락처", "휴대폰", "핸드폰", "휴대전화", "전화번호"),
    "이메일": ("이메일", "메일", "email", "e-mail"),
    "생년월일": ("생년월일", "생일", "출생일", "출생"),
    "주민등록번호": ("주민등록번호", "주민번호", "주민등록", "rrn"),
    "주소": ("주소", "거주지", "소재지"),
    "계좌번호": ("계좌번호", "계좌", "통장번호"),
    "학번": ("학번", "학생번호"),
}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _key(value: str) -> str:
    return re.sub(r"[^0-9a-z가-힣]", "", value.casefold())


def _unique_headers(values: Iterable[Any]) -> list[str]:
    result: list[str] = []
    seen: Counter[str] = Counter()
    for index, raw in enumerate(values, 1):
        base = _text(raw) or f"열{index}"
        seen[base] += 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


def _split_header_rows(values: list[list[Any]]) -> tuple[list[str], list[list[Any]]]:
    """Skip title/blank preamble rows and select the most plausible header row."""
    if not values:
        raise WorkflowError("빈 파일입니다.")
    candidates = values[: min(20, len(values))]
    best_index = 0
    best_score = float("-inf")
    for index, row in enumerate(candidates):
        cells = [_text(cell) for cell in row]
        nonempty = [cell for cell in cells if cell]
        if not nonempty:
            continue
        text_like = sum(not re.fullmatch(r"[-+]?\\d+(?:[.,]\\d+)?", cell) for cell in nonempty)
        unique = len({_key(cell) for cell in nonempty})
        # Prefer a wide, text-heavy, unique row; lightly prefer earlier rows.
        score = len(nonempty) * 3 + text_like * 2 + unique - index * 0.15
        if score > best_score:
            best_score = score
            best_index = index
    return _unique_headers(values[best_index]), values[best_index + 1 :]


def read_table(path: Path, sheet_name: str | None = None) -> TableData:
    path = Path(path)
    if not path.exists():
        raise WorkflowError(f"파일을 찾을 수 없습니다: {path.name}")
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise WorkflowError(f"파일이 안전 처리 한도(250MB)를 넘습니다: {path.name}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "cp949", "utf-8"):
            try:
                with path.open("r", encoding=encoding, newline="") as handle:
                    reader = csv.reader(handle)
                    first = next(reader, None)
                    if first is None:
                        values = []
                    else:
                        values = [first]
                        for row_number, row in enumerate(reader, 2):
                            if row_number > MAX_ROWS + 1:
                                raise WorkflowError(f"행 수가 안전 처리 한도({MAX_ROWS:,}행)를 넘습니다: {path.name}")
                            values.append([_safe_csv_value(cell) for cell in row])
                break
            except UnicodeDecodeError as exc:
                last_error = exc
        else:
            raise WorkflowError(f"CSV 문자 인코딩을 읽을 수 없습니다: {last_error}")
        if not values:
            raise WorkflowError(f"빈 파일입니다: {path.name}")
        values = [[_safe_csv_value(value) for value in row] for row in values]
        headers, body = _split_header_rows(values)
        if len(headers) > MAX_COLUMNS:
            raise WorkflowError(f"열 수가 안전 처리 한도({MAX_COLUMNS}열)를 넘습니다: {path.name}")
        rows = [row for row in body if any(_text(cell) for cell in row)]
    elif suffix in {".xlsx", ".xlsm"}:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            ws = book[sheet_name] if sheet_name else book[book.sheetnames[0]]
            if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > MAX_COLUMNS:
                raise WorkflowError(f"표 크기가 안전 처리 한도를 넘습니다: {path.name}")
            raw_rows = [list(row) for row in ws.iter_rows(values_only=True)]
            if not raw_rows:
                raise WorkflowError(f"빈 시트입니다: {path.name}")
            headers, body = _split_header_rows(raw_rows)
            rows = []
            for row_number, row in enumerate(body, 2):
                if row_number > MAX_ROWS + 1:
                    raise WorkflowError(f"행 수가 안전 처리 한도({MAX_ROWS:,}행)를 넘습니다: {path.name}")
                if len(row) > MAX_COLUMNS:
                    raise WorkflowError(f"열 수가 안전 처리 한도({MAX_COLUMNS}열)를 넘습니다: {path.name}")
                if any(_text(cell) for cell in row):
                    rows.append(list(row))
        finally:
            book.close()
    else:
        raise WorkflowError(f"지원하지 않는 파일 형식입니다: {suffix or '확장자 없음'}")
    return TableData(headers, rows, "원본", source_files=[path]).normalized()


def _safe_csv_value(value: str) -> str:
    """CSV 값이 결과 통합문서에서 의도치 않은 수식으로 실행되지 않게 한다."""
    stripped = value.lstrip()
    if stripped.startswith(("=", "+", "@")) or (stripped.startswith("-") and len(stripped) > 1 and not stripped[1].isdigit()):
        return "'" + value
    return value


def merge_files(paths: Iterable[Path], include_source: bool = True) -> TableData:
    paths = [Path(path) for path in paths]
    if not paths:
        raise WorkflowError("취합할 파일을 한 개 이상 선택하세요.")
    tables = [read_table(path) for path in paths]
    union: list[str] = []
    canonical_headers: dict[str, str] = {}
    for table in tables:
        for header in table.headers:
            canonical = _key(header)
            if canonical not in canonical_headers:
                canonical_headers[canonical] = header
                union.append(header)
    warnings: list[str] = []
    first_keys = {_key(header): header for header in tables[0].headers}
    for table, path in zip(tables[1:], paths[1:]):
        table_keys = {_key(header): header for header in table.headers}
        missing = [header for key, header in first_keys.items() if key not in table_keys]
        extra = [header for header in table.headers if _key(header) not in first_keys]
        if missing or extra:
            details = []
            if missing:
                details.append("없는 열: " + ", ".join(missing))
            if extra:
                details.append("추가 열: " + ", ".join(extra))
            warnings.append(f"{path.name} — " + "; ".join(details))
    source_header = "취합 원본 파일명" if "원본 파일명" in union else "원본 파일명"
    headers = ([source_header] if include_source else []) + union
    rows: list[list[Any]] = []
    for table, path in zip(tables, paths):
        positions = {_key(header): index for index, header in enumerate(table.headers)}
        for row in table.rows:
            values = [row[positions[_key(h)]] if _key(h) in positions and positions[_key(h)] < len(row) else None for h in union]
            rows.append(([path.name] if include_source else []) + values)
    return TableData(headers, rows, "파일 취합", warnings, paths).normalized()


def detect_kind(header: str, values: Iterable[Any] = ()) -> str | None:
    compact = _key(header)
    for kind, keywords in HEADER_KINDS.items():
        if any(_key(keyword) in compact for keyword in keywords):
            return kind
    samples = [_text(value) for value in values if _text(value)][:30]
    if not samples:
        return None
    patterns = {
        "주민등록번호": r"^\d{6}\s*-?\s*[1-8]\d{6}$",
        "전화번호": r"^(?:0\d{1,2})[- .]?\d{3,4}[- .]?\d{4}$",
        "이메일": r"^[^@\s]+@[^@\s]+\.[^@\s]+$",
    }
    for kind, pattern in patterns.items():
        if sum(bool(re.fullmatch(pattern, value)) for value in samples) / len(samples) >= 0.6:
            return kind
    return None


def suggest_mask_rules(table: TableData) -> list[MaskRule]:
    result: list[MaskRule] = []
    for index, header in enumerate(table.headers):
        kind = detect_kind(header, (row[index] for row in table.rows))
        if kind:
            result.append(MaskRule(header, "partial", kind))
    return result


def _mask_name(value: str) -> str:
    chars = list(value)
    indices = [i for i, char in enumerate(chars) if char.isalpha()]
    if len(indices) <= 1:
        return "*" * len(value)
    targets = indices[1:-1] or [indices[-1]]
    for index in targets:
        chars[index] = "*"
    return "".join(chars)


def partial_mask(value: Any, kind: str = "자동") -> str:
    text = _text(value)
    if not text:
        return text
    if kind == "이름":
        return _mask_name(text)
    if kind == "이메일" and "@" in text:
        local, domain = text.split("@", 1)
        keep = local[:1]
        return keep + "*" * max(2, len(local) - len(keep)) + "@" + domain
    if kind in {"전화번호", "주민등록번호", "계좌번호", "학번"}:
        digits = [i for i, char in enumerate(text) if char.isdigit()]
        keep_start = 3 if kind == "전화번호" else 2
        keep_end = 2
        mask_indices = set(digits[keep_start:max(keep_start, len(digits) - keep_end)])
        return "".join("*" if i in mask_indices else char for i, char in enumerate(text))
    if kind == "생년월일":
        digits = [i for i, char in enumerate(text) if char.isdigit()]
        mask_indices = set(digits[4:])
        return "".join("*" if i in mask_indices else char for i, char in enumerate(text))
    if kind == "주소":
        parts = text.split()
        if len(parts) > 2:
            return " ".join(parts[:2] + ["***"])
        if len(parts) == 2:
            return parts[0] + " " + partial_mask(parts[1], "자동")
        return partial_mask(text, "자동")
    if len(text) <= 2:
        return text[0] + "*" * (len(text) - 1)
    return text[0] + "*" * (len(text) - 2) + text[-1]


def full_mask(value: Any) -> str:
    text = _text(value)
    return "".join(char if char.isspace() or char in "-@._/" else "*" for char in text)


def apply_masking(table: TableData, rules: Iterable[MaskRule]) -> TableData:
    rules_by_column = {rule.column: rule for rule in rules}
    unknown = [column for column in rules_by_column if column not in table.headers]
    if unknown:
        raise WorkflowError("마스킹할 열을 찾을 수 없습니다: " + ", ".join(unknown))
    keep_headers = [header for header in table.headers if rules_by_column.get(header, MaskRule(header, "keep")).mode != "drop"]
    rows: list[list[Any]] = []
    for row in table.normalized().rows:
        output: list[Any] = []
        for header, value in zip(table.headers, row):
            rule = rules_by_column.get(header)
            if not rule:
                output.append(value)
            elif rule.mode == "drop":
                continue
            elif rule.mode == "full":
                output.append(full_mask(value))
            elif rule.mode == "partial":
                output.append(partial_mask(value, rule.kind))
            else:
                raise WorkflowError(f"알 수 없는 마스킹 방식입니다: {rule.mode}")
        rows.append(output)
    return TableData(keep_headers, rows, "개인정보 보호", list(table.warnings), list(table.source_files))


def verify_masking(original: TableData, masked: TableData, rules: Iterable[MaskRule]) -> None:
    rules = list(rules)
    for rule in rules:
        if rule.mode == "drop":
            if rule.column in masked.headers:
                raise WorkflowError(f"열 삭제 검증 실패: {rule.column}")
            continue
        if rule.column not in original.headers or rule.column not in masked.headers:
            raise WorkflowError(f"마스킹 검증에 필요한 열이 없습니다: {rule.column}")
        before_i = original.headers.index(rule.column)
        after_i = masked.headers.index(rule.column)
        for row_no, (before, after) in enumerate(zip(original.rows, masked.rows), 2):
            raw = _text(before[before_i])
            if raw and raw == _text(after[after_i]):
                raise WorkflowError(f"마스킹 검증 실패: {rule.column} {row_no}행의 값이 바뀌지 않았습니다.")


def _resolve_columns(request: str, headers: list[str]) -> list[str]:
    compact = _key(request)
    matches = [header for header in headers if _key(header) and _key(header) in compact]
    # 사용자가 화면의 열 번호를 그대로 말하는 경우도 실제 헤더로 변환한다.
    for raw_index in re.findall(r"(?:열|컬럼|번째\s*열)\s*(\d+)", request):
        index = int(raw_index) - 1
        if 0 <= index < len(headers) and headers[index] not in matches:
            matches.append(headers[index])
    # 따옴표 안 표현은 띄어쓰기나 기호가 다른 헤더도 비교한다.
    quoted = re.findall(r"['\"‘’“”]([^'\"‘’“”]+)['\"‘’“”]", request)
    for phrase in quoted:
        for header in headers:
            if _key(phrase) == _key(header) and header not in matches:
                matches.append(header)
    return matches


def plan_cleaning(request: str, table: TableData) -> CleanPlan:
    request = request.strip()
    if not request:
        return CleanPlan(request, [], "데이터 정리를 건너뜁니다.", 1.0)
    headers = table.headers
    columns = _resolve_columns(request, headers)
    actions: list[CleanAction] = []
    clauses = [item.strip() for item in re.split(r"(?:그리고|그다음|다음으로|후에|,|\n|;)", request) if item.strip()]
    recognized_clauses = 0
    for clause in clauses:
        clause_columns = _resolve_columns(clause, headers) or columns
        before_count = len(actions)
        if re.search(r"중복|겹치", clause):
            actions.append(CleanAction("deduplicate", tuple(clause_columns)))
        if re.search(r"빈\s*칸|공백\s*(?:행)?\s*(?:제거|삭제)|누락|(?:값|금액|내용)이?\s*없는|없는\s*행", clause):
            operation = "flag_blanks" if re.search(r"표시|찾|확인|점검", clause) else "drop_blanks"
            actions.append(CleanAction(operation, tuple(clause_columns)))
        if re.search(r"합계|합산|더해|총합|평균|개수|건수|세어", clause):
            if len(clause_columns) < 2:
                actions.append(CleanAction("summarize", tuple(clause_columns)))
            else:
                operation = "average" if "평균" in clause else "count" if re.search(r"개수|건수|세어", clause) else "sum"
                actions.append(CleanAction(operation, tuple(clause_columns)))
        if re.search(r"정렬|순서", clause):
            descending = bool(re.search(r"내림|큰\s*순|높은\s*순|많은\s*순|최신", clause))
            # 집계와 정렬이 한 문장에 있으면 집계 결과의 값 열을 정렬한다.
            sort_columns = (clause_columns[-1],) if clause_columns else ()
            actions.append(CleanAction("sort", sort_columns, descending=descending))
        if re.search(r"앞뒤\s*공백|공백\s*정리|띄어쓰기\s*정리", clause):
            actions.append(CleanAction("trim", tuple(clause_columns)))
        if re.search(r"날짜.*(?:통일|정리|변환|형식)", clause):
            actions.append(CleanAction("normalize_date", tuple(clause_columns)))
        if re.search(r"숫자|금액|콤마", clause) and re.search(r"통일|정리|변환|형식|제거", clause):
            actions.append(CleanAction("normalize_number", tuple(clause_columns)))
        if re.search(r"바꿔|변경|치환", clause):
            quoted_values = re.findall(r"['\"‘’“”]([^'\"‘’“”]+)['\"‘’“”]", clause)
            if len(quoted_values) >= 2:
                old, new = quoted_values[-2:]
                actions.append(CleanAction("replace", tuple(clause_columns), f"{old}\0{new}"))
            else:
                pair = re.search(r"(?:에서\s*)?([^\s]+?)(?:을|를)\s*([^\s]+?)(?:으로|로)\s*(?:바꿔|변경|치환)", clause)
                if pair:
                    actions.append(CleanAction("replace", tuple(clause_columns), f"{pair.group(1)}\0{pair.group(2)}"))
        if len(actions) > before_count:
            recognized_clauses += 1
    warnings: list[str] = []
    if not actions:
        warnings.append("요청에서 안전하게 실행할 작업을 결정하지 못했습니다. 열 이름과 원하는 작업을 함께 적어주세요.")
    for action in actions:
        if action.operation not in {"trim"} and not action.columns and action.operation not in {"summarize"}:
            warnings.append(f"'{action.operation}' 작업의 대상 열을 찾지 못했습니다.")
    confidence = 0.0 if not clauses else max(0.0, min(1.0, recognized_clauses / len(clauses)))
    labels = {
        "deduplicate": "중복 행 제거", "drop_blanks": "빈칸 행 제거", "flag_blanks": "빈칸 표시",
        "sort": "정렬", "sum": "그룹별 합계", "average": "그룹별 평균", "count": "그룹별 개수",
        "trim": "공백 정리", "normalize_date": "날짜 통일", "normalize_number": "숫자 통일",
        "replace": "값 바꾸기", "summarize": "요약",
    }
    explanation = " → ".join(labels[action.operation] for action in actions) or "실행할 작업 없음"
    return CleanPlan(request, actions, explanation, confidence, warnings)


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", "").replace("원", "").replace("명", "")
    try:
        return float(text)
    except ValueError:
        return None


def apply_clean_plan(table: TableData, plan: CleanPlan) -> TableData:
    if plan.warnings:
        raise WorkflowError(" ".join(plan.warnings))
    result = table.normalized()
    for action in plan.actions:
        index = {header: i for i, header in enumerate(result.headers)}
        resolved_columns: list[str] = []
        for column in action.columns:
            if column in index:
                resolved_columns.append(column)
                continue
            derived = [header for header in result.headers if header.startswith(column + " ")]
            if len(derived) == 1:
                resolved_columns.append(derived[0])
        if len(resolved_columns) != len(action.columns):
            raise WorkflowError("정리 작업의 대상 열을 찾을 수 없습니다.")
        positions = [index[column] for column in resolved_columns]
        if action.operation == "deduplicate":
            positions = positions or list(range(len(result.headers)))
            seen = set()
            kept = []
            for row in result.rows:
                key = tuple(_text(row[i]).casefold() for i in positions)
                if key not in seen:
                    seen.add(key)
                    kept.append(row)
            result.rows = kept
        elif action.operation in {"drop_blanks", "flag_blanks"}:
            positions = positions or list(range(len(result.headers)))
            if action.operation == "drop_blanks":
                result.rows = [row for row in result.rows if all(_text(row[i]) for i in positions)]
            else:
                result.headers.append("빈칸 점검")
                for row in result.rows:
                    missing = [result.headers[i] for i in positions if not _text(row[i])]
                    row.append(", ".join(missing) if missing else "정상")
        elif action.operation == "sort":
            if not positions:
                raise WorkflowError("정렬할 열을 요청에 포함하세요.")
            pos = positions[0]
            def sort_key(row: list[Any]):
                number = _number(row[pos])
                return (2, "") if not _text(row[pos]) else (0, number) if number is not None else (1, _text(row[pos]).casefold())
            result.rows.sort(key=sort_key, reverse=action.descending)
        elif action.operation in {"sum", "average", "count"}:
            if len(positions) < 2:
                raise WorkflowError("그룹 열과 계산할 값 열을 요청에 정확히 포함하세요.")
            group_indices, value_i = positions[:-1], positions[-1]
            groups: defaultdict[tuple[str, ...], list[float]] = defaultdict(list)
            for row in result.rows:
                value = _number(row[value_i])
                if action.operation == "count" or value is not None:
                    group_key = tuple(_text(row[i]) for i in group_indices)
                    groups[group_key].append(1.0 if action.operation == "count" else value or 0.0)
            label = {"sum": "합계", "average": "평균", "count": "개수"}[action.operation]
            group_headers = [result.headers[i] for i in group_indices]
            result.headers = [*group_headers, f"{result.headers[value_i]} {label}"]
            result.rows = [[*group, sum(values) / len(values) if action.operation == "average" and values else sum(values)] for group, values in groups.items()]
        elif action.operation == "trim":
            positions = positions or list(range(len(result.headers)))
            for row in result.rows:
                for pos in positions:
                    if isinstance(row[pos], str):
                        row[pos] = re.sub(r"\s+", " ", row[pos]).strip()
        elif action.operation == "normalize_number":
            for row in result.rows:
                for pos in positions:
                    number = _number(row[pos])
                    if number is not None:
                        row[pos] = int(number) if number.is_integer() else number
        elif action.operation == "normalize_date":
            for row in result.rows:
                for pos in positions:
                    value = row[pos]
                    if isinstance(value, (datetime, date)):
                        row[pos] = value.strftime("%Y-%m-%d")
                    else:
                        match = re.fullmatch(r"\s*(\d{4})[./년 -]+(\d{1,2})[./월 -]+(\d{1,2})일?\s*", _text(value))
                        if match:
                            row[pos] = f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
        elif action.operation == "replace":
            old, new = (action.value or "\0").split("\0", 1)
            positions = positions or list(range(len(result.headers)))
            for row in result.rows:
                for pos in positions:
                    if isinstance(row[pos], str):
                        row[pos] = row[pos].replace(old, new)
        elif action.operation == "summarize":
            raise WorkflowError("요약하려면 그룹 열과 계산할 값 열을 요청에 포함하세요.")
    result.stage = "데이터 정리"
    return result


def save_table(table: TableData, path: Path, sheet_name: str = "결과") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    ws = book.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "_", sheet_name)[:31] or "결과"
    ws.append(table.headers)
    for row in table.normalized().rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    book.save(path)
    if not path.exists() or path.stat().st_size == 0:
        raise WorkflowError("결과 파일 저장에 실패했습니다.")
    check = read_table(path)
    if check.headers != table.headers or len(check.rows) != len(table.rows):
        path.unlink(missing_ok=True)
        raise WorkflowError("저장 후 검증에서 행 또는 열이 달라졌습니다.")
    return path


def save_clean_workbook(
    source: TableData,
    plan: CleanPlan,
    preview: TableData,
    path: Path,
    use_formulas: bool = True,
    include_source: bool = True,
) -> Path:
    """집계 결과는 미리보기 값 대신 검증 가능한 Excel 수식으로 저장한다."""
    aggregation_index = next(
        (i for i, action in enumerate(plan.actions) if action.operation in {"sum", "average", "count"}),
        None,
    )
    if not use_formulas or aggregation_index is None:
        return save_table(preview, path, "정리결과")

    aggregation = plan.actions[aggregation_index]
    prefix = CleanPlan(plan.request, plan.actions[:aggregation_index], "집계 전 정리", plan.confidence)
    formula_source = apply_clean_plan(source, prefix) if prefix.actions else source.normalized()
    source_index = {header: i for i, header in enumerate(formula_source.headers)}
    if any(column not in source_index for column in aggregation.columns) or len(aggregation.columns) < 2:
        raise WorkflowError("수식 저장에 필요한 그룹 열과 값 열을 찾을 수 없습니다.")

    group_columns = list(aggregation.columns[:-1])
    value_column = aggregation.columns[-1]
    book = Workbook()
    result_sheet = book.active
    result_sheet.title = "정리결과"
    result_sheet.append(preview.headers)
    data_sheet = book.create_sheet("원본데이터")
    data_sheet.append(formula_source.headers)
    for row in formula_source.rows:
        data_sheet.append(row)
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = data_sheet.dimensions

    last_row = max(2, len(formula_source.rows) + 1)
    for output_row, preview_row in enumerate(preview.rows, 2):
        for column, value in enumerate(preview_row[:-1], 1):
            result_sheet.cell(output_row, column, value)
        criteria: list[str] = []
        for group_position, group_column in enumerate(group_columns, 1):
            source_letter = get_column_letter(source_index[group_column] + 1)
            criteria.extend([f"'원본데이터'!${source_letter}$2:${source_letter}${last_row}", f"{get_column_letter(group_position)}{output_row}"])
        if aggregation.operation == "count":
            formula = f"=COUNTIFS({','.join(criteria)})"
        else:
            value_letter = get_column_letter(source_index[value_column] + 1)
            function = "SUMIFS" if aggregation.operation == "sum" else "AVERAGEIFS"
            formula = f"={function}('원본데이터'!${value_letter}$2:${value_letter}${last_row},{','.join(criteria)})"
        result_sheet.cell(output_row, len(preview.headers), formula)

    result_sheet.freeze_panes = "A2"
    result_sheet.auto_filter.ref = result_sheet.dimensions
    for sheet in (result_sheet, data_sheet):
        for cell in sheet[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
    if not include_source:
        data_sheet.sheet_state = "hidden"

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    book.calculation.calcMode = "auto"
    book.save(path)
    check = load_workbook(path, read_only=True, data_only=False)
    try:
        if check["정리결과"].max_row - 1 != len(preview.rows):
            raise WorkflowError("수식 결과 저장 후 행 수 검증에 실패했습니다.")
        formulas = [check["정리결과"].cell(row, len(preview.headers)).value for row in range(2, len(preview.rows) + 2)]
        if not formulas or any(not isinstance(value, str) or not value.startswith("=") for value in formulas):
            raise WorkflowError("수식 결과 저장 후 수식 검증에 실패했습니다.")
    finally:
        check.close()
    return path
