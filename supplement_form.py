"""K-에듀파인 **추경** 세출예산요구내역(입력 후 내려받은 파일)을 읽고 점검한다.

2026년 제1회 추경 입력본 3부로 확인한 서식이다. 본예산 서식과 열이 다르다.

    레벨 | 순번 | * 사업항목 | * 원가통계비목 | * 상세내역 |
    2026년도 기정 (산출기초 · 금액①) | 2026년도 최종요구 (* 산출기초 · * 금액②) |
    추경요구(②-①) | 성립전 확정 후 추경 반영액 | 성립전 진행 금액 | 연계등록 | 목적사업 배분금액

  - 3행이 총계, 4행부터 트리. 계층은 사업항목 들여쓰기(2칸 = 1단)다.
  - 산출기초는 **원 단위**, 금액은 천원. 끝에 '=' 가 붙는다.
  - 추경에서만 나오는 산출기초 모양
        (184,000,000원×3개원)-(9,926,670원×3개원)=         깎은 만큼 빼서 적는다
        5,000,000원×20개원×1회+4,000,000원×19개원×1회=     늘린 만큼 더해 적는다
        0=                                                  전액 감액
  - 새로 넣은 줄은 기정 칸이 비어 있다(①=없음).

오류만 싣는다. 모두 입력한 값이 틀렸다는 뜻이다.

  안에서   ②−① = 추경요구 · 산출기초 = 금액(①·② 둘 다) · 위 줄 = 아래 줄 합(①·②·추경요구)
           · 총계 = 1단 합 · 레벨 = 들여쓰기÷2+1 · 금액이 있는 맨 아래 줄에 비목·상세내역·산출기초
  설명서와 사업 증감 = 설명서 증감 · 금액①/② = 기정/추경안 · 설명서 ① 항목 증감마다
           같은 증감(되도록 같은 기정·추경안)의 줄이 있어야 하고, 그 아래 산출 줄의 증감도 있어야 한다
  UBIS 와  사업 증감 = 추경요구액(B) · 금액① = 기정액(A)
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from crosscheck import SKIPPED, Report
from supplement import (SuppDocument, UbisSupp, _find_row, _fmt, _formula_ok, _same, calc as supp_calc,
                        match_key)

GOAL = "목표2B"
MARKER = re.compile(r"^\s*(?:\d+\.|[가-힣]\.|\d+\)|[가-힣]\)|\(\d+\)|\([가-힣]\))\s*")


@dataclass
class SuppFormRow:
    row: int
    level: Optional[int]
    sequence: Optional[int]
    name: str                 # 번호를 뗀 이름
    label: str                # 원문 그대로 ('      가) 사립유치원지원비')
    indent: int
    code: str = ""
    detail: str = ""
    basis_before: str = ""
    before: Optional[float] = None       # 금액①
    basis_after: str = ""
    after: Optional[float] = None        # 금액②
    change: Optional[float] = None       # 추경요구(②-①)

    @property
    def depth(self) -> int:
        return self.indent // 2 + 1


@dataclass
class SuppForm:
    path: str
    total: Optional[SuppFormRow] = None
    rows: list[SuppFormRow] = field(default_factory=list)

    def children(self, index: int) -> list[int]:
        depth = self.rows[index].depth
        found = []
        for position in range(index + 1, len(self.rows)):
            if self.rows[position].depth <= depth:
                break
            if self.rows[position].depth == depth + 1:
                found.append(position)
        return found

    def subtree(self, index: int) -> list[int]:
        depth = self.rows[index].depth
        end = index + 1
        while end < len(self.rows) and self.rows[end].depth > depth:
            end += 1
        return list(range(index, end))

    def projects(self) -> list[int]:
        return [index for index, row in enumerate(self.rows) if row.depth == 1]


def _flat(value) -> str:
    return re.sub(r"\s+", "", str(value or "")).replace("*", "")


def is_supplement_form(path: str) -> bool:
    """K-에듀파인 다운로드가 추경 서식인가. '추경요구' 열이 있으면 추경이다."""
    try:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for row in book.active.iter_rows(min_row=1, max_row=3, values_only=True):
                if any("추경요구" in _flat(value) for value in row):
                    return True
        finally:
            book.close()
    except Exception:  # noqa: BLE001
        return False
    return False


def _columns(sheet) -> dict[str, int]:
    """머리글 두 줄을 이어 읽어 열을 찾는다. 1행은 두 열에 걸쳐 병합돼 오른쪽이 비어 있다."""
    found: dict[str, int] = {}
    carried = ""
    for column in range(1, sheet.max_column + 1):
        top = _flat(sheet.cell(1, column).value)
        sub = _flat(sheet.cell(2, column).value)
        if top:
            carried = top
        elif sub:
            top = carried
        pairs = (("레벨", "레벨"), ("순번", "순번"), ("사업항목", "사업항목"), ("비목", "원가통계"),
                 ("상세내역", "상세내역"), ("추경요구", "추경요구"))
        for key, needle in pairs:
            if key not in found and needle in top:
                found[key] = column
        if "기정" in top and "산출기초" in sub:
            found.setdefault("기정산출", column)
        elif "기정" in top and "금액" in sub:
            found.setdefault("기정금액", column)
        elif "최종요구" in top and "산출기초" in sub:
            found.setdefault("요구산출", column)
        elif "최종요구" in top and "금액" in sub:
            found.setdefault("요구금액", column)
    return found


def _number(value) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def read_supplement_form(path: str) -> SuppForm:
    from openpyxl import load_workbook

    sheet = load_workbook(path, data_only=True).active
    columns = _columns(sheet)
    needed = ("사업항목", "기정금액", "요구금액", "추경요구")
    missing = [key for key in needed if key not in columns]
    if missing:
        raise ValueError(f"'{Path(path).name}' 은 K-에듀파인 추경 세출예산요구내역이 아닙니다. "
                         f"'{' · '.join(missing)}' 열을 찾지 못했습니다.")
    form = SuppForm(path=path)

    def cell(row: int, key: str):
        return sheet.cell(row, columns[key]).value if key in columns else None

    for row in range(3, sheet.max_row + 1):
        label = str(cell(row, "사업항목") or "")
        if not label.strip():
            continue
        entry = SuppFormRow(
            row=row,
            level=int(_number(cell(row, "레벨"))) if _number(cell(row, "레벨")) is not None else None,
            sequence=int(_number(cell(row, "순번"))) if _number(cell(row, "순번")) is not None else None,
            name=MARKER.sub("", label).strip(), label=label, indent=len(label) - len(label.lstrip(" ")),
            code=str(cell(row, "비목") or "").strip(), detail=str(cell(row, "상세내역") or "").strip(),
            basis_before=str(cell(row, "기정산출") or "").strip(), before=_number(cell(row, "기정금액")),
            basis_after=str(cell(row, "요구산출") or "").strip(), after=_number(cell(row, "요구금액")),
            change=_number(cell(row, "추경요구")))
        if label.strip() == "총계" and form.total is None and not form.rows:
            form.total = entry
            continue
        form.rows.append(entry)
    return form


# ------------------------------------------------------------------ 계산

def basis_value(text: str) -> Optional[float]:
    """K-에듀파인 산출기초(원 단위) → 천원. '0=' 은 0."""
    body = (text or "").split("=")[0].strip()
    if not body:
        return None
    if re.fullmatch(r"0+", body):
        return 0.0
    return supp_calc(body)


def _zero(value: Optional[float]) -> float:
    return value or 0.0


# ------------------------------------------------------------------ 검사

def _where(form: SuppForm, row: SuppFormRow) -> str:
    return f"[{Path(form.path).name} {row.row}행 '{row.name}']"


def check_form(form: SuppForm, report: Optional[Report] = None) -> Report:
    """다운로드본 안에서 맞아야 할 것."""
    report = report or Report()
    rows = form.rows
    broken: set[int] = set()
    for index, row in enumerate(rows):
        # 한 줄에서 나온 문제는 한 건으로 묶는다. 금액 한 칸을 잘못 치면 '②−①≠추경요구'와
        # '산출기초≠금액'이 함께 드러나는데, 두 건으로 세면 고칠 곳이 둘인 줄 안다.
        problems: list[str] = []
        pair = (None, None)
        if row.level is not None and row.level != row.depth:
            problems.append(f"레벨 칸은 {row.level}인데 들여쓰기로는 {row.depth}단입니다(다른 단계에 넣음)")
        for text, amount, word in ((row.basis_before, row.before, "기정"), (row.basis_after, row.after, "최종요구")):
            if not text or amount is None:
                continue
            value = basis_value(text)
            if value is None:
                report.add(GOAL, SKIPPED, row.name, "", f"{_where(form, row)} {word} 산출기초 계산 불가: {text}")
                continue
            if not _formula_ok(value, amount):
                broken.add(index)
                problems.append(f"{word} 산출기초 {text} 을 계산하면 {_fmt(value)}천원인데 금액은 {_fmt(amount)}천원")
                pair = (amount, value)
        if row.change is not None and not _same(_zero(row.after) - _zero(row.before), row.change):
            broken.add(index)
            problems.append(f"금액② {_fmt(row.after)} − 금액① {_fmt(row.before)} = "
                            f"{_fmt(_zero(row.after) - _zero(row.before))}천원인데 추경요구가 {_fmt(row.change)}천원")
            if pair == (None, None):
                pair = (row.change, _zero(row.after) - _zero(row.before))
        if not form.children(index) and (_zero(row.after) or _zero(row.before)):
            missing = [word for word, value in (("원가통계비목", row.code), ("상세내역", row.detail),
                                                ("최종요구 산출기초", row.basis_after)) if not value]
            if missing:
                problems.append(f"금액이 있는 맨 아래 줄인데 {' · '.join(missing)}이(가) 비어 있음")
        if problems:
            report.add(GOAL, "오류", row.name, "", f"{_where(form, row)} " + " / ".join(problems) + ".",
                       pair[0], pair[1])
    # 위 줄 = 바로 아래 줄들의 합. 깊은 곳부터, 아래가 이미 틀렸으면 위는 말하지 않는다.
    for index in range(len(rows) - 1, -1, -1):
        children = form.children(index)
        if not children or any(position in broken for position in form.subtree(index)[1:]):
            if children and any(position in broken for position in form.subtree(index)[1:]):
                broken.add(index)
            continue
        row = rows[index]
        for key, word in (("before", "금액①(기정)"), ("after", "금액②(최종요구)"), ("change", "추경요구")):
            mine = getattr(row, key)
            summed = sum(_zero(getattr(rows[position], key)) for position in children)
            if not _same(_zero(mine), summed):
                broken.add(index)
                report.add(GOAL, "오류", row.name, "",
                           f"{_where(form, row)} {word} {_fmt(mine)}천원인데 바로 아래 {len(children)}줄을 더하면 "
                           f"{_fmt(summed)}천원입니다.", mine, summed)
                break
    if form.total is not None:
        tops = form.projects()
        for key, word in (("before", "금액①"), ("after", "금액②"), ("change", "추경요구")):
            written = getattr(form.total, key)
            summed = sum(_zero(getattr(rows[position], key)) for position in tops)
            if written is not None and not _same(written, summed):
                report.add(GOAL, "오류", "총계", "",
                           f"[{Path(form.path).name} 총계] {word} {_fmt(written)}천원 · 1단 사업 합 {_fmt(summed)}천원.",
                           written, summed)
    return report


def _find_project(document: SuppDocument, name: str):
    key = match_key(name)
    found = [one for one in document.projects if match_key(one.name) == key]
    return found[0] if len(found) == 1 else None


def _node_for(form: SuppForm, span: list[int], after, before, change, used: set[int]) -> Optional[int]:
    """설명서 한 줄(① 항목)에 맞는 다운로드본 줄. 증감이 같고, 되도록 기정·추경안까지 같은 줄."""
    candidates = [position for position in span if position not in used
                  and _same(_zero(form.rows[position].change), change)]
    exact = [position for position in candidates
             if (before is None or _same(_zero(form.rows[position].before), before))
             and (after is None or _same(_zero(form.rows[position].after), after))]
    pool = exact or candidates
    return pool[0] if pool else None


def compare_form(form: SuppForm, document: Optional[SuppDocument] = None,
                 ubis: Optional[UbisSupp] = None, report: Optional[Report] = None) -> Report:
    """다운로드본 ↔ 추경 설명서 · UBIS 추경 검토조서."""
    report = report or Report()
    rows = form.rows
    for top in form.projects():
        row = rows[top]
        span = form.subtree(top)
        name = Path(form.path).name
        if ubis is not None:
            found, _alias = _find_row(ubis.rows, row.name)
            if found is None:
                report.add(GOAL, "오류", row.name, "",
                           f"[{name}] '{row.name}' 이 UBIS 추경 검토조서에 없습니다. 사업명을 확인해 주세요.")
            else:
                if not _same(_zero(row.change), found.change):
                    report.add(GOAL, "오류", row.name, "",
                               f"[{name}] 추경요구 {_fmt(row.change)}천원 · UBIS 추경요구액(B) {_fmt(found.change)}천원 — "
                               f"{_fmt(abs(_zero(row.change) - found.change))}천원 차이.", row.change, found.change)
                if not _same(_zero(row.before), found.before):
                    report.add(GOAL, "오류", row.name, "",
                               f"[{name}] 금액①(기정) {_fmt(row.before)}천원 · UBIS 기정액(A) {_fmt(found.before)}천원 — "
                               f"{_fmt(abs(_zero(row.before) - found.before))}천원 차이.", row.before, found.before)
        if document is None:
            continue
        project = _find_project(document, row.name)
        if project is None:
            report.add(GOAL, "오류", row.name, "",
                       f"[{name}] '{row.name}' 사업 쪽이 추경 설명서에 없습니다. 설명서에 없는 사업을 입력했거나 "
                       "사업명이 다릅니다.")
            continue
        if project.change is not None and not _same(_zero(row.change), project.change):
            report.add(GOAL, "오류", row.name, "",
                       f"[{name}] 추경요구 {_fmt(row.change)}천원 · 설명서 증감 {_fmt(project.change)}천원 — "
                       f"{_fmt(abs(_zero(row.change) - project.change))}천원 차이.", row.change, project.change)
            continue       # 사업 전체가 어긋났으면 아래 줄 대조는 같은 잘못을 여러 번 말할 뿐이다
        partial = ubis is not None and _find_row(ubis.rows, row.name)[0] is not None and \
            _same(_zero(row.before), _find_row(ubis.rows, row.name)[0].before)
        for key, word, theirs in (("before", "금액①(기정)", project.before), ("after", "금액②(최종요구)", project.after)):
            if theirs is None or _same(_zero(getattr(row, key)), theirs) or partial:
                # partial: 다운로드본 기정이 UBIS 와 같다. 설명서 쪽이 UBIS 사업의 일부만 적은
                # 총액배분 사업이라 금액①·②가 설명서와 다른 게 정상이다(증감은 위에서 맞췄다).
                continue
            report.add(GOAL, "오류", row.name, "",
                       f"[{name}] {word} {_fmt(getattr(row, key))}천원 · 설명서 {_fmt(theirs)}천원 — "
                       f"{_fmt(abs(_zero(getattr(row, key)) - theirs))}천원 차이.", getattr(row, key), theirs)
        _compare_groups(form, span[1:], project, report)
    return report


def _compare_groups(form: SuppForm, span: list[int], project, report: Report) -> None:
    """설명서 ① 항목과 그 아래 산출 줄의 증감이 다운로드본에 들어 있는지."""
    used: set[int] = set()
    name = Path(form.path).name
    for table in project.tables:
        if table.kind != "증감":
            continue
        for group in table.groups:
            if group.change is None or _same(group.change, 0):
                continue
            title = group.title or group.business
            node = _node_for(form, span, group.after, group.before, group.change, used)
            if node is None:
                report.add(GOAL, "오류", project.name, title,
                           f"[{name}] 설명서 '{title}' 의 증감 {_fmt(group.change)}천원(기정 {_fmt(group.before)} → "
                           f"추경 {_fmt(group.after)})이 입력본 어느 줄에도 없습니다. 빠뜨렸거나 금액을 다르게 입력했습니다.",
                           group.change, None)
                continue
            used.add(node)
            inside = form.subtree(node)
            taken: set[int] = {node}
            lines = group.lines
            for position, line in enumerate(lines):
                if line.amount is None or _same(line.amount, 0):
                    continue
                parent = position + 1 < len(lines) and lines[position + 1].depth > line.depth
                hit = _node_for(form, inside, None, None, line.amount, taken)
                if hit is None and parent:
                    continue      # 묶음 줄이다. 입력본이 묶지 않고 풀어 적었어도 아래 줄들이 맞으면 된다
                if hit is None:
                    report.add(GOAL, "오류", project.name, line.name,
                               f"[{name}] 설명서 '{title} > {line.name}' 의 증감 {_fmt(line.amount)}천원이 입력본 "
                               f"'{form.rows[node].name}' 아래 어느 줄에도 없습니다.", line.amount, None)
                    continue
                taken.add(hit)
