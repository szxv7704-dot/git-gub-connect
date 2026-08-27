from __future__ import annotations

import re
from pathlib import Path

from PySide6.QtCore import Qt, QRectF
from PySide6.QtGui import QColor, QImage, QPen, QPixmap
from PySide6.QtWidgets import (
    QAbstractItemView, QDialog, QDialogButtonBox, QGraphicsPixmapItem,
    QGraphicsRectItem, QGraphicsScene, QGraphicsView, QGridLayout, QHBoxLayout,
    QLabel, QPushButton, QSpinBox, QSplitter, QTableWidget, QTableWidgetItem,
    QVBoxLayout, QWidget,
)


class MaskRectItem(QGraphicsRectItem):
    def __init__(self, rect: QRectF, candidate=None):
        super().__init__(rect)
        self.candidate = candidate
        self.setPen(QPen(QColor("#E11D48"), 3))
        self.setBrush(QColor(225, 29, 72, 55))
        self.setFlags(
            QGraphicsRectItem.GraphicsItemFlag.ItemIsMovable
            | QGraphicsRectItem.GraphicsItemFlag.ItemIsSelectable
        )


class VisualMaskDialog(QDialog):
    """PDF/이미지의 마스킹 좌표를 저장 전에 직접 확인하고 보정한다."""

    def __init__(self, path: Path, candidates: list, parent=None):
        super().__init__(parent)
        self.path = path
        self.candidates = candidates
        self.scale = 1.0
        self.page_no = self._page_number(candidates[0].location) if candidates else 1
        self.setWindowTitle(f"영역 검토 · {path.name}")
        self.resize(1120, 800)

        root = QVBoxLayout(self)
        self.notice = QLabel("붉은 영역이 실제 개인정보와 정확히 겹치는지 확인하세요. 영역을 끌어 이동하고 아래 숫자로 크기를 조절할 수 있습니다.")
        root.addWidget(self.notice)
        self.scene = QGraphicsScene(self)
        self.view = QGraphicsView(self.scene)
        self.view.setDragMode(QGraphicsView.DragMode.ScrollHandDrag)
        root.addWidget(self.view, 1)

        controls = QHBoxLayout()
        add = QPushButton("＋ 영역 추가")
        add.clicked.connect(self.add_rect)
        delete = QPushButton("선택 영역 삭제")
        delete.clicked.connect(self.delete_rect)
        controls.addWidget(add); controls.addWidget(delete); controls.addSpacing(18)
        self.spins = []
        for label in ("X", "Y", "너비", "높이"):
            controls.addWidget(QLabel(label))
            spin = QSpinBox(); spin.setRange(0, 20000); spin.setSingleStep(2)
            spin.valueChanged.connect(self.apply_geometry)
            controls.addWidget(spin); self.spins.append(spin)
        controls.addStretch(); root.addLayout(controls)

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel)
        buttons.button(QDialogButtonBox.StandardButton.Save).setText("이 영역으로 반영")
        buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject)
        root.addWidget(buttons)
        self.scene.selectionChanged.connect(self.selection_changed)
        self._load_page()

    @staticmethod
    def _page_number(location: str) -> int:
        match = re.search(r"(\d+)페이지", location)
        return int(match.group(1)) if match else 1

    def _load_page(self):
        suffix = self.path.suffix.lower()
        added = 0
        if suffix == ".pdf":
            import fitz
            from privacy_core import _ocr_candidate_rects, _ocr_pdf_page
            doc = fitz.open(self.path)
            page = doc[self.page_no - 1]
            pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
            image = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888).copy()
            self.scale = pix.width / page.rect.width
            self.scene.addItem(QGraphicsPixmapItem(QPixmap.fromImage(image)))
            unresolved = []
            for candidate in self.candidates:
                rect = None
                if candidate.bbox:
                    rect = fitz.Rect(candidate.bbox)
                elif candidate.value and not candidate.manual:
                    found = page.search_for(candidate.value)
                    rect = found[0] if found else None
                if rect:
                    self._add_item(QRectF(rect.x0*self.scale, rect.y0*self.scale, rect.width*self.scale, rect.height*self.scale), candidate)
                    added += 1
                else:
                    unresolved.append(candidate)
            # 스캔 PDF에는 검색 가능한 텍스트층이 없으므로 OCR 좌표를 다시 계산한다.
            seen = set()
            for config in ("--psm 6", "--psm 11"):
                if not unresolved:
                    break
                data, ocr_scale = _ocr_pdf_page(page, config=config)
                for candidate, rect in _ocr_candidate_rects(data, unresolved, ocr_scale):
                    key = (id(candidate),) + tuple(round(value, 1) for value in (rect.x0, rect.y0, rect.x1, rect.y1))
                    if key in seen:
                        continue
                    seen.add(key)
                    self._add_item(QRectF(rect.x0*self.scale, rect.y0*self.scale, rect.width*self.scale, rect.height*self.scale), candidate)
                    added += 1
            doc.close()
        else:
            from PIL import Image
            from privacy_core import _ocr_candidate_rects, _ocr_data
            pixmap = QPixmap(str(self.path))
            self.scene.addItem(QGraphicsPixmapItem(pixmap))
            unresolved = []
            for candidate in self.candidates:
                if candidate.bbox:
                    x0,y0,x1,y1 = candidate.bbox
                    self._add_item(QRectF(x0,y0,x1-x0,y1-y0), candidate)
                    added += 1
                else:
                    unresolved.append(candidate)
            if unresolved:
                with Image.open(self.path) as source:
                    image = source.convert("RGB")
                enlarged = image.resize((image.width * 2, image.height * 2), Image.Resampling.LANCZOS)
                seen = set()
                for ocr_image, ocr_scale, config in ((image, 1.0, "--psm 6"), (enlarged, 2.0, "--psm 11")):
                    data = _ocr_data(ocr_image, config=config)
                    for candidate, rect in _ocr_candidate_rects(data, unresolved, ocr_scale):
                        key = (id(candidate),) + tuple(round(value, 1) for value in (rect.x0, rect.y0, rect.x1, rect.y1))
                        if key in seen:
                            continue
                        seen.add(key)
                        self._add_item(QRectF(rect.x0, rect.y0, rect.width, rect.height), candidate)
                        added += 1
        if not added:
            self.notice.setText("⚠ 선택한 후보의 자동 좌표를 찾지 못했습니다. ‘＋ 영역 추가’로 개인정보 영역을 직접 지정하세요.")
            self.notice.setStyleSheet("color:#B91C1C; font-weight:700;")
        self.scene.setSceneRect(self.scene.itemsBoundingRect())
        self.view.fitInView(self.scene.sceneRect(), Qt.AspectRatioMode.KeepAspectRatio)

    def _add_item(self, rect, candidate=None):
        item = MaskRectItem(rect, candidate)
        self.scene.addItem(item)
        return item

    def add_rect(self):
        center = self.view.mapToScene(self.view.viewport().rect().center())
        item = self._add_item(QRectF(center.x()-70, center.y()-18, 140, 36))
        item.setSelected(True)

    def delete_rect(self):
        for item in list(self.scene.selectedItems()):
            if isinstance(item, MaskRectItem): self.scene.removeItem(item)

    def selection_changed(self):
        items = [x for x in self.scene.selectedItems() if isinstance(x, MaskRectItem)]
        if not items: return
        rect = items[0].mapRectToScene(items[0].rect())
        values = (rect.x(), rect.y(), rect.width(), rect.height())
        for spin, value in zip(self.spins, values):
            spin.blockSignals(True); spin.setValue(round(value)); spin.blockSignals(False)

    def apply_geometry(self):
        items = [x for x in self.scene.selectedItems() if isinstance(x, MaskRectItem)]
        if not items: return
        x,y,w,h = [spin.value() for spin in self.spins]
        item = items[0]
        item.setPos(0,0); item.setRect(QRectF(x,y,max(2,w),max(2,h)))

    def areas(self):
        result = []
        for item in self.scene.items():
            if not isinstance(item, MaskRectItem): continue
            rect = item.mapRectToScene(item.rect())
            result.append((item.candidate, (rect.x()/self.scale, rect.y()/self.scale, rect.right()/self.scale, rect.bottom()/self.scale)))
        return result


class ExcelPreviewDialog(QDialog):
    def __init__(self, path: Path, candidate, replacement: str, parent=None):
        super().__init__(parent)
        self.setWindowTitle(f"셀 비교 · {path.name}")
        self.resize(1050, 560)
        root = QVBoxLayout(self)
        root.addWidget(QLabel(f"위치: {candidate.location}  |  선택한 셀은 노란색으로 표시됩니다."))
        split = QSplitter(); root.addWidget(split, 1)
        left = self._panel("원본 시트"); right = self._panel("처리 예정")
        split.addWidget(left[0]); split.addWidget(right[0])
        self._fill(path, candidate, left[1], right[1], replacement)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Close); buttons.rejected.connect(self.reject); root.addWidget(buttons)

    def _panel(self, title):
        box=QWidget(); layout=QVBoxLayout(box); layout.addWidget(QLabel(title)); table=QTableWidget(); table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers); layout.addWidget(table); return box,table

    def _fill(self, path, candidate, original, result, replacement):
        if path.suffix.lower()==".xls":
            import xlrd
            book=xlrd.open_workbook(path); sheet_name,coordinate=candidate.location.rsplit("!",1); ws=book.sheet_by_name(sheet_name)
            match=re.fullmatch(r"R(\d+)C(\d+)",coordinate)
            if not match:return
            row=int(match.group(1)); col=int(match.group(2)); r0=max(1,row-5); r1=min(ws.nrows,row+5); c0=max(1,col-4); c1=min(ws.ncols,col+4)
            for table in (original,result):
                table.setRowCount(r1-r0+1); table.setColumnCount(c1-c0+1)
                table.setHorizontalHeaderLabels([f"C{c}" for c in range(c0,c1+1)]); table.setVerticalHeaderLabels([str(r) for r in range(r0,r1+1)])
            for rr,r in enumerate(range(r0,r1+1)):
                for cc,c in enumerate(range(c0,c1+1)):
                    value=str(ws.cell_value(r-1,c-1)); a=QTableWidgetItem(value); b=QTableWidgetItem(replacement if (r,c)==(row,col) else value)
                    if (r,c)==(row,col): a.setBackground(QColor("#FEF3C7")); b.setBackground(QColor("#DCFCE7"))
                    original.setItem(rr,cc,a); result.setItem(rr,cc,b)
            return
        from openpyxl import load_workbook
        book=load_workbook(path, data_only=False, read_only=True)
        sheet_name, coordinate=candidate.location.rsplit("!",1)
        ws=book[sheet_name]
        match=re.fullmatch(r"([A-Z]+)(\d+)", coordinate)
        if not match: book.close(); return
        from openpyxl.utils.cell import column_index_from_string
        col=column_index_from_string(match.group(1)); row=int(match.group(2))
        r0=max(1,row-5); r1=row+5; c0=max(1,col-4); c1=col+4
        for table in (original,result):
            table.setRowCount(r1-r0+1); table.setColumnCount(c1-c0+1)
            table.setHorizontalHeaderLabels([ws.cell(1,c).column_letter for c in range(c0,c1+1)])
            table.setVerticalHeaderLabels([str(r) for r in range(r0,r1+1)])
        for rr,r in enumerate(range(r0,r1+1)):
            for cc,c in enumerate(range(c0,c1+1)):
                value="" if ws.cell(r,c).value is None else str(ws.cell(r,c).value)
                a=QTableWidgetItem(value); b=QTableWidgetItem(replacement if (r,c)==(row,col) else value)
                if (r,c)==(row,col): a.setBackground(QColor("#FEF3C7")); b.setBackground(QColor("#DCFCE7"))
                original.setItem(rr,cc,a); result.setItem(rr,cc,b)
        book.close()
