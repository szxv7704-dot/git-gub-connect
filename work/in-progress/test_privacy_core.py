import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from privacy_core import (
    Candidate,
    _ocr_account_candidates,
    _ocr_name_candidates,
    business_number_valid,
    collect_files,
    find_in_text,
    flexible_pattern,
    looks_like_name,
    luhn_valid,
    process_file,
    replacement_for,
    rrn_checksum_valid,
    scan_file,
    unsupported_files,
)


class PrivacyCoreTests(unittest.TestCase):
    def test_custom_value_ignores_case_and_whitespace(self):
        self.assertIsNotNone(flexible_pattern("Test User").search("t E s T     u s e r"))

    def test_rule_detection(self):
        found = find_in_text("이름: 홍길동 / 010-1234-5678", Path("a.xlsx"), "S!A1", [])
        self.assertEqual({x.kind for x in found}, {"이름", "전화번호"})

    def test_xlsx_literal_filter_and_alias_output(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "input.xlsx"
            output = folder / "out"
            output.mkdir()
            book = Workbook()
            book.active["A1"] = "담당자: 홍길동 / TEST user"
            book.save(source)
            candidates, warnings = scan_file(source, ["testuser"])
            self.assertFalse(warnings)
            self.assertEqual(len(candidates), 1)
            self.assertEqual(candidates[0].kind, "사용자 지정")
            result = process_file(source, candidates, "alias", output, {})
            self.assertEqual(result.status, "성공")
            value = load_workbook(result.output).active["A1"].value
            self.assertIn("홍길동", value)
            self.assertNotIn("TEST user", value)
            self.assertIn("[사용자 지정_0001]", value)

    def test_xlsx_column_context_detects_name_and_account_and_black_masks(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "홍길동_지급내역.xlsx"
            output = folder / "out"
            output.mkdir()
            book = Workbook()
            sheet = book.active
            sheet.append(["성명", "계좌번호"])
            sheet.append(["홍길동", "999-01-110001"])
            book.save(source)

            candidates, warnings = scan_file(source, [])
            self.assertFalse(warnings)
            self.assertEqual({c.kind for c in candidates}, {"이름", "계좌번호"})
            result = process_file(source, candidates, "delete", output, {})
            self.assertEqual(result.status, "성공", result.message)
            values = load_workbook(result.output).active
            self.assertEqual(values["A2"].value, "■" * len("홍길동"))
            self.assertEqual(values["B2"].value, "■" * len("999-01-110001"))
            self.assertNotIn("홍길동", result.output.name)

    def test_xlsx_merged_title_and_late_header_are_supported(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "merged.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.merge_cells("A1:F1")
            sheet["A1"] = "교육 이수 결과"
            sheet["C6"] = "성명"
            sheet["C7"] = "홍길동"
            sheet["C8"] = "김   영"
            book.save(source)

            candidates, warnings = scan_file(source, ["성명"])
            self.assertFalse(warnings)
            self.assertEqual([candidate.location for candidate in candidates], [f"{sheet.title}!C7", f"{sheet.title}!C8"])

    def test_explicit_name_header_accepts_rare_surname_but_does_not_mask_signature_column(self):
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "register.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["연번", "성함", "서명"])
            sheet.append([1, "가경화", "가경화"])
            sheet.append([2, "남궁민수", "남궁민수"])
            book.save(source)

            candidates, warnings = scan_file(source, ["성명"])
            self.assertFalse(warnings)
            self.assertEqual([candidate.location for candidate in candidates], [f"{sheet.title}!B2", f"{sheet.title}!B3"])

    def test_header_driven_address_and_phone_detection(self):
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "contacts.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["주소", "휴대폰 번호"])
            sheet.append(["서울 어딘가 101호", "010 1234 5678"])
            book.save(source)

            candidates, warnings = scan_file(source, ["주소", "전화번호"])
            self.assertFalse(warnings)
            self.assertEqual({candidate.kind for candidate in candidates}, {"주소", "전화번호"})

    def test_collect_files_ignores_excel_temporary_lock_files(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            (folder / "normal.xlsx").touch()
            (folder / "~$normal.xlsx").touch()
            self.assertEqual([path.name for path in collect_files([folder])], ["normal.xlsx"])

    def test_account_and_depositor_detection_in_text(self):
        found = find_in_text("예금주: 홍길동 / 계좌번호: 9999-07-770007", Path("a.jpg"), "이미지 전체", [])
        self.assertEqual({x.kind for x in found}, {"이름", "계좌번호"})

    def test_entered_field_names_filter_out_other_detected_types(self):
        text = "성명: 홍길동 / 주민등록번호: 850314-1163517 / 계좌번호: 999-01-110001"
        found = find_in_text(text, Path("a.xlsx"), "S!A1", ["성명", "계좌번호"])
        self.assertEqual({x.kind for x in found}, {"이름", "계좌번호"})

    def test_entered_literal_shows_only_that_value(self):
        text = "성명: 홍길동 / 주민등록번호: 850314-1163517 / 별칭: TARGET VALUE"
        found = find_in_text(text, Path("a.xlsx"), "S!A1", ["targetvalue"])
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0].kind, "사용자 지정")
        self.assertEqual(found[0].value, "TARGET VALUE")

    def test_ocr_table_name_after_account_number_joins_split_syllables(self):
        data = {
            "text": ["1002-548-731946", "김", "서", "윤", "연락처"],
            "conf": ["95", "91", "92", "93", "90"],
            "left": [100, 300, 315, 330, 400], "top": [10] * 5,
            "width": [120, 12, 12, 12, 40], "height": [14] * 5,
            "block_num": [1] * 5, "par_num": [1] * 5, "line_num": [1] * 5,
        }
        found = _ocr_name_candidates(data, Path("scan.pdf"), "1페이지(OCR)")
        self.assertEqual([candidate.value for candidate in found], ["김서윤"])

    def test_ocr_account_rejoins_spaced_and_hyphenated_digits(self):
        data = {
            "text": ["계좌번호", "625-019015-01-012"], "conf": ["70", "92"],
            "left": [10, 100], "top": [10, 10], "width": [70, 180], "height": [20, 20],
            "block_num": [1, 1], "par_num": [1, 1], "line_num": [1, 1],
        }
        found = _ocr_account_candidates(data, Path("bank.jpg"), require_label=True)
        self.assertEqual([candidate.value for candidate in found], ["625-019015-01-012"])

    def test_hwpx_table_header_detects_only_name_column(self):
        import zipfile
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "table.hwpx"
            xml = """<?xml version='1.0' encoding='UTF-8'?><section><tbl>
            <tr><tc><t>행사 담당자 등록부</t></tc></tr>
            <tr><tc><t>연번</t></tc><tc><t>성함</t></tc><tc><t>서명</t></tc></tr>
            <tr><tc><t>1</t></tc><tc><t>홍은성</t></tc><tc><t>홍은성</t></tc></tr>
            </tbl></section>"""
            with zipfile.ZipFile(source, "w") as archive:
                archive.writestr("Contents/section0.xml", xml)
            candidates, warnings = scan_file(source, ["이름"])
            self.assertFalse(warnings)
            self.assertEqual([(candidate.value, candidate.location) for candidate in candidates], [("홍은성", "Contents/section0.xml 표1 행3 열2")])

    def test_pdf_permanent_redaction_removes_text(self):
        import fitz
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "input.pdf"
            output = folder / "out"
            output.mkdir()
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "Email: private@example.com")
            doc.save(source)
            doc.close()
            candidates, warnings = scan_file(source, [])
            self.assertFalse(warnings)
            result = process_file(source, candidates, "delete", output, {})
            self.assertEqual(result.status, "성공")
            sanitized = fitz.open(result.output)
            self.assertNotIn("private@example.com", "".join(p.get_text() for p in sanitized))
            pixel = sanitized[0].get_pixmap(clip=fitz.Rect(70, 60, 220, 80), alpha=False)
            self.assertLess(min(pixel.samples), 20)
            sanitized.close()

    def test_scanned_pdf_ocr_and_pixel_redaction(self):
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "scan.pdf"
            output = folder / "out"
            output.mkdir()
            image = Image.new("RGB", (1400, 300), "white")
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 72)
            draw.text((40, 90), "PHONE 010-1234-5678", fill="black", font=font)
            image_path = folder / "scan.png"
            image.save(image_path)
            doc = fitz.open()
            page = doc.new_page(width=700, height=150)
            page.insert_image(page.rect, filename=str(image_path))
            doc.save(source)
            doc.close()
            candidates, warnings = scan_file(source, [])
            phones = [c for c in candidates if c.kind == "전화번호"]
            self.assertTrue(phones, warnings)
            self.assertTrue(phones[0].location.endswith("(OCR)"))
            result = process_file(source, phones, "delete", output, {})
            self.assertEqual(result.status, "성공", result.message)
            remaining, _ = scan_file(result.output, [])
            self.assertFalse([c for c in remaining if c.kind == "전화번호"])

    def test_mixed_pdf_runs_ocr_and_is_rebuilt_without_original_objects(self):
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "mixed.pdf"
            output = folder / "out"
            output.mkdir()
            image = Image.new("RGB", (1400, 300), "white")
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 72)
            draw.text((40, 90), "PHONE 010-1234-5678", fill="black", font=font)
            image_path = folder / "scan.png"
            image.save(image_path)
            doc = fitz.open()
            page = doc.new_page(width=700, height=150)
            page.insert_image(page.rect, filename=str(image_path))
            page.insert_text((10, 15), "PAGE 1")
            page.add_text_annot((650, 20), "hidden note")
            doc.embfile_add("secret.txt", b"original attachment")
            doc.save(source)
            doc.close()

            candidates, warnings = scan_file(source, [])
            phones = [candidate for candidate in candidates if candidate.kind == "전화번호"]
            self.assertTrue(phones, warnings)
            result = process_file(source, phones, "delete", output, {})
            self.assertEqual(result.status, "성공", result.message)

            sanitized = fitz.open(result.output)
            self.assertEqual(sanitized.embfile_count(), 0)
            self.assertEqual(sanitized[0].get_text("text"), "")
            self.assertIsNone(sanitized[0].first_annot)
            self.assertEqual(len(sanitized[0].get_images(full=True)), 1)
            sanitized.close()

    def test_image_delete_replaces_pixels_and_survives_ocr_recheck(self):
        from PIL import Image, ImageDraw, ImageFont, ImageStat
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "phone.png"
            output = folder / "out"
            output.mkdir()
            image = Image.new("RGB", (1000, 220), "white")
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 64)
            draw.text((30, 70), "010-1234-5678", fill="black", font=font)
            image.save(source)
            candidates, warnings = scan_file(source, [])
            phones = [candidate for candidate in candidates if candidate.kind == "전화번호"]
            self.assertTrue(phones, warnings)
            result = process_file(source, phones, "delete", output, {})
            self.assertEqual(result.status, "성공", result.message)
            with Image.open(result.output) as sanitized:
                self.assertLess(min(ImageStat.Stat(sanitized).extrema[0]), 10)


class DetectionRuleTests(unittest.TestCase):
    def test_compact_birth_date_is_not_a_service_phone_number(self):
        found = find_in_text("생년월일: 19760913", Path("form.pdf"), "1페이지(OCR)", [])
        self.assertEqual([(candidate.kind, candidate.value) for candidate in found], [("생년월일", "19760913")])

    def test_foreign_registration_number_is_detected(self):
        found = find_in_text("등록번호 900101-5123456", Path("a.xlsx"), "S!A1", [])
        self.assertEqual([c.kind for c in found], ["주민등록번호"])

    def test_invalid_birth_date_is_not_a_registration_number(self):
        self.assertFalse(find_in_text("정산코드 991399-1234567", Path("a.xlsx"), "S!A1", []))

    def test_internet_and_toll_free_phone_numbers_are_detected(self):
        for number in ("070-1234-5678", "0505-123-4567", "080-123-4567", "1588-1234"):
            with self.subTest(number=number):
                found = find_in_text(f"연락 {number}", Path("a.xlsx"), "S!A1", [])
                self.assertEqual([c.kind for c in found], ["전화번호"], number)

    def test_order_number_is_not_reported_as_a_card_number(self):
        self.assertFalse([c for c in find_in_text("주문번호 2026082512345678", Path("a.xlsx"), "S!A1", []) if c.kind == "카드번호"])

    def test_card_number_passing_luhn_is_reported(self):
        found = [c for c in find_in_text("4111-1111-1111-1111", Path("a.xlsx"), "S!A1", []) if c.kind == "카드번호"]
        self.assertEqual(len(found), 1)
        self.assertGreater(found[0].confidence, 0.9)

    def test_passport_and_licence_and_business_numbers(self):
        found = find_in_text("여권번호: M12345678 / 면허 11-22-334455-66 / 사업자등록번호 220-81-62517", Path("a.xlsx"), "S!A1", [])
        self.assertEqual({c.kind for c in found}, {"여권번호", "운전면허번호", "사업자등록번호"})
        self.assertEqual(next(c.value for c in found if c.kind == "여권번호"), "M12345678")

    def test_checksums(self):
        self.assertTrue(luhn_valid("4111111111111111"))
        self.assertFalse(luhn_valid("4111111111111112"))
        self.assertTrue(rrn_checksum_valid("850314-1163511"))
        self.assertFalse(rrn_checksum_valid("850314-1163517"))
        self.assertTrue(business_number_valid("220-81-62517"))
        self.assertFalse(business_number_valid("220-81-62518"))

    def test_compound_surname_is_recognised_without_an_explicit_header(self):
        self.assertTrue(looks_like_name("남궁민수"))
        self.assertFalse(looks_like_name("결과보고"))

    def test_unsupported_files_are_reported_to_the_user(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            (folder / "ok.xlsx").touch()
            (folder / "old.hwp").touch()
            (folder / "메모.docx").touch()
            self.assertEqual([p.name for p in unsupported_files([folder])], ["old.hwp", "메모.docx"])


class OutputSafetyTests(unittest.TestCase):
    def test_pdf_alias_text_keeps_korean_characters(self):
        import fitz
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            output = folder / "out"
            output.mkdir()
            source = folder / "input.pdf"
            doc = fitz.open()
            doc.new_page().insert_text((72, 72), "Email: private@example.com")
            doc.save(source)
            doc.close()
            candidates, _ = scan_file(source, [])
            result = process_file(source, candidates, "alias", output, {})
            self.assertEqual(result.status, "성공", result.message)
            sanitized = fitz.open(result.output)
            text = "".join(page.get_text() for page in sanitized)
            sanitized.close()
            self.assertIn("[이메일_0001]", text)
            self.assertNotIn("?", text)
            self.assertNotIn("private@example.com", text)

    def test_pdf_scan_does_not_report_the_same_value_twice(self):
        import fitz
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "input.pdf"
            doc = fitz.open()
            doc.new_page().insert_text((72, 72), "Email: private@example.com")
            doc.save(source)
            doc.close()
            candidates, _ = scan_file(source, [])
            self.assertEqual(len(candidates), 1)

    def test_pdf_output_is_rejected_when_the_original_text_survives(self):
        import fitz
        from privacy_core import Candidate, _verify_pdf_output
        with tempfile.TemporaryDirectory() as raw:
            output = Path(raw) / "kept.pdf"
            doc = fitz.open()
            doc.new_page().insert_text((72, 72), "private@example.com")
            doc.save(output)
            doc.close()
            candidate = Candidate("이메일", "private@example.com", Path("a.pdf"), "1페이지", 0.98)
            with self.assertRaises(RuntimeError):
                _verify_pdf_output(output, [candidate])
            self.assertFalse(output.exists())

    def test_hwpx_preview_text_and_author_are_sanitized(self):
        import zipfile
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            output = folder / "out"
            output.mkdir()
            source = folder / "table.hwpx"
            xml = """<?xml version='1.0' encoding='UTF-8'?><section><tbl>
            <tr><tc><t>연번</t></tc><tc><t>성함</t></tc></tr>
            <tr><tc><t>1</t></tc><tc><t>홍은성</t></tc></tr>
            </tbl></section>"""
            hpf = "<?xml version='1.0' encoding='UTF-8'?><package xmlns:dc='http://purl.org/dc/elements/1.1/'><dc:creator>홍은성</dc:creator></package>"
            with zipfile.ZipFile(source, "w") as archive:
                archive.writestr("Contents/section0.xml", xml)
                archive.writestr("Contents/content.hpf", hpf)
                archive.writestr("Preview/PrvText.txt", "담당자 홍은성 확인")
                archive.writestr("Preview/PrvImage.png", b"\x89PNG-original-preview-bytes")
            candidates, _ = scan_file(source, ["이름"])
            result = process_file(source, candidates, "delete", output, {})
            self.assertEqual(result.status, "성공", result.message)
            with zipfile.ZipFile(result.output) as archive:
                preview = archive.read("Preview/PrvText.txt").decode("utf-8")
                self.assertNotIn("홍은성", preview)
                self.assertNotIn(b"original-preview-bytes", archive.read("Preview/PrvImage.png"))
                self.assertNotIn("홍은성", archive.read("Contents/content.hpf").decode("utf-8"))

    def test_jpeg_output_carries_no_exif(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            output = folder / "out"
            output.mkdir()
            source = folder / "photo.jpg"
            exif = Image.Exif()
            exif[271] = "SECRET-MAKER"
            Image.new("RGB", (240, 80), "white").save(source, exif=exif)
            with Image.open(source) as before:
                self.assertEqual(before.getexif()[271], "SECRET-MAKER")
            result = process_file(source, [], "delete", output, {})
            self.assertEqual(result.status, "성공", result.message)
            with Image.open(result.output) as after:
                self.assertEqual(dict(after.getexif()), {})

    def test_excel_comment_author_and_text_are_masked(self):
        from openpyxl.comments import Comment
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            output = folder / "out"
            output.mkdir()
            source = folder / "note.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["성명"])
            sheet.append(["홍길동"])
            sheet["A2"].comment = Comment("홍길동 재확인 필요", "홍길동")
            book.save(source)
            candidates, _ = scan_file(source, [])
            result = process_file(source, candidates, "delete", output, {})
            self.assertEqual(result.status, "성공", result.message)
            saved = load_workbook(result.output).active
            self.assertNotIn("홍길동", saved["A2"].comment.text)
            self.assertNotIn("홍길동", saved["A2"].comment.author or "")


class ReviewWorkflowTests(unittest.TestCase):
    def test_four_output_styles_have_distinct_replacements(self):
        candidate = Candidate("이름", "홍길동", Path("a.xlsx"), "S!A1", 1.0)
        self.assertEqual(replacement_for(candidate, "delete", {}), "■■■")
        self.assertEqual(replacement_for(candidate, "label_ko", {}), "[이름]")
        self.assertEqual(replacement_for(candidate, "label_en", {}), "[NAME]")
        self.assertEqual(replacement_for(candidate, "alias", {}), "[이름_0001]")

    def test_manual_pdf_rectangle_is_applied_without_text_search(self):
        import fitz
        with tempfile.TemporaryDirectory() as raw:
            folder=Path(raw); source=folder/"manual.pdf"; output=folder/"out"; output.mkdir()
            doc=fitz.open(); page=doc.new_page(width=220,height=100); page.insert_text((30,55),"MANUAL SECRET"); doc.save(source); doc.close()
            candidate=Candidate("수동 지정","수동 영역",source,"1페이지",1.0,True,(20,30,150,70),True,True)
            result=process_file(source,[candidate],"delete",output,{})
            self.assertEqual(result.status,"성공",result.message)
            checked=fitz.open(result.output)
            self.assertNotIn("MANUAL SECRET","".join(p.get_text() for p in checked))
            checked.close()


if __name__ == "__main__":
    unittest.main()
