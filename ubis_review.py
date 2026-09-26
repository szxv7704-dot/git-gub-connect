"""UBIS 검토조서(서식1) = 본예산 세출요구 자료를 읽는다.

산출식도 비목도 없는 요약 자료다. 사업별 요구액만 뽑아 금액 대조에 쓴다.
`사업항목 및 산출기초` 열은 '◎ 항목(재원) 금액천원' 형태의 요약문이지 산출식이 아니다.
파싱해서 입력원으로 쓰지 않는다.

열은 위치가 아니라 헤더 이름으로 찾는다.
"""

from __future__ import annotations

import re
from typing import Optional

from openpyxl import load_workbook

HEADER_SCAN = 10


def _flat(value) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def _columns(sheet) -> tuple[int, dict[str, int]]:
    """헤더 행 번호와 열 위치를 찾는다."""
    for row in range(1, min(HEADER_SCAN, sheet.max_row) + 1):
        labels = {column: _flat(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1)}
        if not any("사업명" in label for label in labels.values()):
            continue
        found: dict[str, int] = {}
        for column, label in labels.items():
            if not label:
                continue
            if "사업명" in label and "사업명" not in found:
                found["사업명"] = column
            elif "사업항목" in label and "세세부사업" in label and "사업항목" not in found:
                found["사업항목"] = column
            elif "세부사업" in label and "세부사업" not in found:
                found["세부사업"] = column
            elif "기관명" in label and "기관명" not in found:
                found["기관명"] = column
            elif "본예산요구액" in label and "요구액" not in found:
                found["요구액"] = column
            elif "조정액" in label and "조정액" not in found:
                found["조정액"] = column
        if "사업명" in found and "요구액" in found:
            return row, found
    for row in range(1, min(HEADER_SCAN, sheet.max_row) + 1):
        if any("추경요구액" in _flat(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1)):
            raise ValueError("추경 검토조서입니다. 본예산 설명서와는 맞춰 볼 수 없습니다.\n\n"
                             "추경 사업별 설명서(HWPX)와 함께 넣거나, 본예산 검토조서를 넣어 주세요.")
    raise ValueError("UBIS 검토조서(서식1)가 아닙니다. '사업명'과 '본예산요구액' 열을 찾지 못했습니다.")


def read_requests(path: str) -> dict[str, float]:
    """사업명 → 요구액(천원). 사업명이 비었으면 사업항목(세세부사업)으로 대신한다."""
    sheet = load_workbook(path, data_only=True).active
    header, columns = _columns(sheet)
    requests: dict[str, float] = {}
    for row in range(header + 1, sheet.max_row + 1):
        name = str(sheet.cell(row, columns["사업명"]).value or "").strip()
        if not name and "사업항목" in columns:
            name = str(sheet.cell(row, columns["사업항목"]).value or "").strip()
        amount = sheet.cell(row, columns["요구액"]).value
        if not name or not isinstance(amount, (int, float)):
            continue
        # 첫 행은 기관 합계인 경우가 있다. 단위·세부사업이 비어 있으면 건너뛴다.
        if "세부사업" in columns and not str(sheet.cell(row, columns["세부사업"]).value or "").strip():
            continue
        requests[name] = requests.get(name, 0) + float(amount)
    if not requests:
        raise ValueError("UBIS 검토조서에서 사업별 요구액을 읽지 못했습니다.")
    return requests


def read_adjustments(path: str) -> dict[str, float]:
    """사업명 → 조정액(D). 0 인 사업은 싣지 않는다.

    2025·2026 본예산 설명서의 금액은 UBIS 요구액이 아니라 **요구액 + 조정액**과 같았다
    (2026년 17개 사업, 2025년 3개 사업 — 차액이 조정액과 1천원도 다르지 않다).
    이걸 모르면 조정이 있었던 사업이 전부 'UBIS 와 다르다'는 오류로 나온다.
    """
    sheet = load_workbook(path, data_only=True).active
    header, columns = _columns(sheet)
    if "조정액" not in columns:
        return {}
    found: dict[str, float] = {}
    for row in range(header + 1, sheet.max_row + 1):
        name = str(sheet.cell(row, columns["사업명"]).value or "").strip()
        if not name and "사업항목" in columns:
            name = str(sheet.cell(row, columns["사업항목"]).value or "").strip()
        value = sheet.cell(row, columns["조정액"]).value
        if not name or not isinstance(value, (int, float)) or not value:
            continue
        if "세부사업" in columns and not str(sheet.cell(row, columns["세부사업"]).value or "").strip():
            continue
        found[name] = found.get(name, 0) + float(value)
    return found


def total(path: str) -> Optional[float]:
    values = read_requests(path)
    return sum(values.values()) if values else None
