"""단위과제카드 목록(선택). 넣으면 '내가 입력할 사업'만 남긴다.

K-에듀파인은 자기가 가진 과제카드만 보여 주고, 그 카드에 딸린 사업만 입력할 수 있다.
과 전체 설명서에는 69개 사업이 있지만 한 담당자가 입력하는 것은 그중 몇 개뿐이다.
전부 늘어놓으면 자기 것을 찾는 데만 시간이 든다.

카드 목록은 `유비스 → 예산요구 → 과제카드세출예산요구목록` 에서 받는다. 첫 열이
`유보통합정책운영관리[P10AAA24101510398569]` 처럼 이름과 카드번호를 함께 담고 있다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Optional

from openpyxl import load_workbook

CODE = re.compile(r"\[([A-Z0-9]+)\]\s*$")
TAIL = re.compile(r"\([^()]*\)\s*$")
NOISE = re.compile(r"[\s·․ㆍ∙・_\-~]")


def squeeze(name: str) -> str:
    """이름 대조용. 설명서와 카드는 띄어쓰기·가운뎃점·붙임표가 제멋대로 다르다.

        설명서 '유-보-초 이음교육 운영'  ↔  카드 '유보초이음교육운영(특교포함)'
        설명서 '영유아교육내실화 지원'    ↔  카드 '영유아교육내실화지원(특교)'

    꼬리 괄호는 카드 쪽에만 붙는 표시라 떼고 본다.
    """
    cleaned = TAIL.sub("", (name or "").strip()).strip()
    return NOISE.sub("", cleaned)


@dataclass
class Card:
    name: str
    code: str = ""
    program: str = ""        # 세부사업

    @property
    def key(self) -> str:
        return squeeze(self.name)


@dataclass
class CardList:
    cards: list = field(default_factory=list)
    source: str = ""

    def __len__(self) -> int:
        return len(self.cards)

    def find(self, project_name: str) -> Optional[Card]:
        key = squeeze(project_name)
        if not key:
            return None
        for card in self.cards:
            if card.key == key:
                return card
        # 카드 이름이 더 길게 적힌 경우가 있다. ('특수학교 교육과정 운영 지원(고교학점제…)')
        for card in self.cards:
            if card.key.startswith(key) or key.startswith(card.key):
                return card
        return None


def load_cards(path: str) -> CardList:
    """과제카드세출예산요구목록 엑셀을 읽는다."""
    sheet = load_workbook(path, data_only=True).active
    column = None
    header = 1
    for row in range(1, min(6, sheet.max_row + 1)):
        for index in range(1, sheet.max_column + 1):
            if "단위과제카드" in str(sheet.cell(row, index).value or "").replace(" ", ""):
                column, header = index, row
                break
        if column:
            break
    if column is None:
        raise ValueError("과제카드세출예산요구목록이 아닙니다. '단위과제카드' 열을 찾지 못했습니다.")

    found = CardList(source=path)
    for row in range(header + 1, sheet.max_row + 1):
        raw = str(sheet.cell(row, column).value or "").strip()
        if not raw or raw in ("합계", "소계", "계"):
            continue
        code = CODE.search(raw)
        name = CODE.sub("", raw).strip()
        if not name:
            continue
        found.cards.append(Card(name=name, code=code.group(1) if code else "",
                                program=str(sheet.cell(row, column + 1).value or "").strip()))
    if not found.cards:
        raise ValueError("과제카드 목록에서 카드를 한 장도 찾지 못했습니다.")
    return found


def check_cards(projects, cards: CardList):
    """설명서와 카드 목록을 맞춰 본다.

    카드에 없는 사업은 **오류가 아니다.** 다른 담당자의 사업일 뿐이다. 다만 카드에는
    있는데 설명서에 없는 사업은 요구를 빠뜨렸을 수 있으니 짚어 준다.
    """
    from crosscheck import Report

    report = Report()
    matched = set()
    for project in projects:
        card = cards.find(project.name)
        if card is not None:
            matched.add(card.key)
        elif not project.handover:
            report.add("목표4", "안내", project.name, "",
                       "내 과제카드에 없는 사업입니다. 다른 담당자가 입력합니다.")
    for card in cards.cards:
        if card.key not in matched:
            report.add("목표4", "확인 필요", card.name, "",
                       "과제카드에는 있는데 설명서에서 같은 이름의 사업을 찾지 못했습니다. "
                       "올해 요구가 없는 카드인지, 이름이 달라 못 찾은 것인지 확인해 주세요.")
    return report
